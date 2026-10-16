# -*- coding: utf-8 -*-
"""output 副本写入快照（写入侧人工改动感知 · 阶段一：快照 + 行标识重定位）。

背景（《2026-10-15 写入侧人工改动感知实施计划》决策 3/5/7）：
- 写入侧以原件为需求基准、消费侧以 output 副本为基准，而副本可被人工
  自由改动（改值/删行/插行），系统此前完全无感知（MX1 复活事故）；
- 本模块在每次系统写副本后全量重扫副本建快照（绝不增量记账），
  落 output/{batch_id}/output_copy_snapshot.json（备份机制 output/ 全量
  覆盖，零额外接线；rerun 归档把旧快照随 run 归档，零特判）；
- 行标识 row_key = (工厂 MAKER_MEI_KJ, SKU SHOHIN_CD, 出现序号 seq)：
  KANRI_NO 实测重复严重不可作定位依据（同一票同 SKU 拆两行是真实业务
  形态），仅作辅助校验列记入快照；同 (工厂, SKU) 组内按行号升序取 seq
  （从 1 起），与 row_map[工厂][SKU] 列表一一对应。

阶段一边界：快照 version=2 字段一次到位（qty/protected/protected_fields/
tombstones 全部预留），但 protected/tombstones 恒为默认值——结转合并
逻辑随阶段二 diff 一起落地；本阶段不含任何冲突语义。

容错哲学（与 _write_audit 同）：快照是辅助设施，refresh_snapshot 任何
失败只记 error 日志并退化为「无快照」，绝不阻塞写入主流程。
"""
from __future__ import annotations

import json
import logging
import os
import threading
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.config import get_settings

logger = logging.getLogger(__name__)

SNAPSHOT_VERSION = 2
SNAPSHOT_FILENAME = "output_copy_snapshot.json"

# row_key 的 JSON 字符串键分隔符：字面 "\0" 两个字符（反斜杠 + 数字零，
# 不是 NUL 控制字符）——工厂名/SKU 不可能包含它；v1 曾误用控制字符踩坑
KEY_SEP = "\\0"

# 辅助校验列（不参与定位）：管理番号 / MX 人工行识别列
COL_KANRI_NO = "KANRI_NO"
COL_PURCHASE_ORDER = "PURCHASE_ORDER"

# 并发兜底：快照读写是单进程内「读-改-写」，per-batch 字典锁
# （照 factory_watch._state_lock 先例；_locks_guard 保护字典本身的创建）
_locks: dict[str, threading.Lock] = {}
_locks_guard = threading.Lock()


def _batch_lock(batch_id: str) -> threading.Lock:
    with _locks_guard:
        lock = _locks.get(batch_id)
        if lock is None:
            lock = threading.Lock()
            _locks[batch_id] = lock
        return lock


def snapshot_path(batch_id: str) -> Path:
    """快照文件路径：output/{batch_id}/output_copy_snapshot.json。"""
    return get_settings().batch_output_dir(batch_id) / SNAPSHOT_FILENAME


def _find_copy(batch_id: str) -> Path | None:
    """定位批次副本：containers/ 下 mtime 最新的 *_filled.xlsx（与 XD 透视/
    分体积统计同口径）；无副本返回 None。"""
    containers_dir = get_settings().batch_containers_dir(batch_id)
    if not containers_dir.is_dir():
        return None
    filled = [p for p in containers_dir.glob("*_filled.xlsx") if p.is_file()]
    if not filled:
        return None
    return max(filled, key=lambda p: p.stat().st_mtime)


def norm_key_text(v) -> str:
    """行标识键成分规整：None/空白 → ""；数值型单元格（Excel 把 13 位条码
    存成数字时）转整数字符串，与 parse_downstream 的 pandas str 口径对齐。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v).strip()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def round2(v) -> float | None:
    """数值规整：net/gross/qty 统一 round(float(v), 2) 或 None
    （与 writer 写入 round(..., 2) 对齐，防浮点尾差误判；非数值 → None）。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    try:
        return round(float(str(v).strip()), 2)
    except (TypeError, ValueError):
        return None


def _text_or_none(v) -> str | None:
    s = norm_key_text(v)
    return s or None


def _is_mx(v) -> bool:
    """MX 人工行判定（PO 值以 "MX" 开头）。与 writer._is_mx_po 同口径；
    此处不复用是为了让本模块只依赖 app.config（writer 反向依赖本模块）。"""
    if v is None:
        return False
    return str(v).strip().upper().startswith("MX")


def row_key_str(key: tuple[str, str, int]) -> str:
    """row_key 元组 → JSON 字符串键（字面 \\0 分隔，见 KEY_SEP 注释）。"""
    factory, sku, seq = key
    return f"{factory}{KEY_SEP}{sku}{KEY_SEP}{seq}"


def parse_row_key(s: str) -> tuple[str, str, int]:
    """row_key_str 逆运算；格式不符抛 ValueError。"""
    parts = s.split(KEY_SEP)
    if len(parts) != 3:
        raise ValueError(f"非法 row_key 字符串键: {s!r}")
    return parts[0], parts[1], int(parts[2])


def build_row_index(ws) -> dict[tuple[str, str, int], int]:
    """全量扫 ws 建 row_key → 当前行号索引（决策 3 重定位的共享扫描）。

    - row_key = (工厂, SKU, seq)：同 (工厂, SKU) 组内按行号升序取第 N 次
      出现（seq 从 1 起），与 row_map[工厂][SKU] 列表一一对应；
    - 工厂/SKU 任一空值的行（排版填充行）不进索引；
    - 表头缺 MAKER_MEI_KJ/SHOHIN_CD 属异常布局，零容错直接 ValueError。
    """
    settings = get_settings()
    it = ws.iter_rows()
    header_row = next(it, None)
    if header_row is None:
        return {}
    header = [c.value for c in header_row]
    missing = [n for n in (settings.col_factory, settings.col_sku)
               if n not in header]
    if missing:
        raise ValueError(f"下游表缺少关键列 {missing}，无法建立行标识索引")
    ci_f = header.index(settings.col_factory)
    ci_s = header.index(settings.col_sku)

    index: dict[tuple[str, str, int], int] = {}
    seqs: dict[tuple[str, str], int] = {}
    for row_no, row in enumerate(it, start=2):
        factory = norm_key_text(row[ci_f].value if ci_f < len(row) else None)
        sku = norm_key_text(row[ci_s].value if ci_s < len(row) else None)
        if not factory or not sku:
            continue
        group = (factory, sku)
        seqs[group] = seqs.get(group, 0) + 1
        index[(factory, sku, seqs[group])] = row_no
    return index


def build_snapshot(batch_id: str, origin: str = "system") -> dict | None:
    """全量扫副本建快照（version=2 格式一次到位）；无副本返回 None。

    origin 应用于全部行（阶段一无逐行结转）："system" 常规重建 /
    "baseline" 无快照迁移基线（决策 7，由 refresh_snapshot 判定传入）；
    "manual_added" 留待阶段五人工新增行纳入时使用。
    """
    copy_path = _find_copy(batch_id)
    if copy_path is None:
        return None
    settings = get_settings()
    st = copy_path.stat()

    wb = load_workbook(copy_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows()
        header_row = next(it, None)
        header = [c.value for c in header_row] if header_row else []
        missing = [n for n in (settings.col_factory, settings.col_sku)
                   if n not in header]
        if missing:
            raise ValueError(f"副本缺少关键列 {missing}，无法建快照: {copy_path}")

        def _col(name: str) -> int | None:
            return header.index(name) if name in header else None

        ci_f = header.index(settings.col_factory)
        ci_s = header.index(settings.col_sku)
        # 三列/管理番号/数量/PO 均可缺：新建副本基线时三列尚未插入，
        # 缺失列一律记 None（kanri_no 仅辅助校验，缺列不参与定位）
        ci_kanri = _col(COL_KANRI_NO)
        ci_cn = _col(settings.col_name_cn)
        ci_net = _col(settings.col_net)
        ci_gross = _col(settings.col_gross)
        ci_qty = _col(settings.col_qty)
        ci_po = _col(COL_PURCHASE_ORDER)

        def val(row, ci: int | None):
            return row[ci].value if ci is not None and ci < len(row) else None

        rows: list[dict] = []
        seqs: dict[tuple[str, str], int] = {}
        data_row_count = 0  # 副本数据行总数（含工厂/SKU 空值的填充行）
        for row_no, row in enumerate(it, start=2):
            data_row_count += 1
            factory = norm_key_text(val(row, ci_f))
            sku = norm_key_text(val(row, ci_s))
            if not factory or not sku:
                continue  # 工厂/SKU 皆空/任一空的行不进快照（已确认 §七.6）
            group = (factory, sku)
            seqs[group] = seqs.get(group, 0) + 1
            rows.append({
                "key": {"factory": factory, "sku": sku, "seq": seqs[group]},
                "row_no": row_no,
                "kanri_no": _text_or_none(val(row, ci_kanri)),
                "name_cn": _text_or_none(val(row, ci_cn)),
                "net": round2(val(row, ci_net)),
                "gross": round2(val(row, ci_gross)),
                "qty": round2(val(row, ci_qty)),
                "is_mx": _is_mx(val(row, ci_po)),
                "origin": origin,
                # 阶段一恒为默认；protected 结转逻辑随阶段二 diff 落地
                "protected": False,
                "protected_fields": [],
            })
    finally:
        wb.close()

    return {
        "version": SNAPSHOT_VERSION,
        "batch_id": batch_id,
        "copy_file": copy_path.name,
        "taken_at": datetime.now().isoformat(timespec="seconds"),
        "copy_mtime_ns": st.st_mtime_ns,
        "copy_size": st.st_size,
        "data_row_count": data_row_count,
        # 阶段一恒为空名单；墓碑结转（只增不减）随阶段二 diff 落地
        "tombstones": [],
        "rows": rows,
    }


def save_snapshot(batch_id: str, snap: dict) -> None:
    """原子写快照：先写 .tmp 再 os.replace（跨平台），绝不写半截文件。"""
    path = snapshot_path(batch_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps(snap, ensure_ascii=False, indent=2),
                   encoding="utf-8")
    os.replace(tmp, path)


def load_snapshot(batch_id: str) -> dict | None:
    """读快照；文件不存在 / JSON 损坏 / 版本不符一律返回 None
    （损坏与版本不符记 warning，按无快照处理）。"""
    path = snapshot_path(batch_id)
    if not path.exists():
        return None
    try:
        snap = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as e:
        logger.warning("[快照] 批次 %s 快照文件损坏，按无快照处理：%s",
                       batch_id, e)
        return None
    if not isinstance(snap, dict) or snap.get("version") != SNAPSHOT_VERSION:
        logger.warning("[快照] 批次 %s 快照版本不符（%r ≠ %d），按无快照处理",
                       batch_id,
                       snap.get("version") if isinstance(snap, dict)
                       else type(snap).__name__,
                       SNAPSHOT_VERSION)
        return None
    return snap


def refresh_snapshot(batch_id: str) -> dict | None:
    """build + save 一体入口（快照三更新点统一调它）。

    无快照迁移（决策 7）：副本存在但无快照文件 → 以当前副本现状为基线，
    全行 origin="baseline" + warning（此前人工改动不可追溯，一次性取舍）。
    辅助设施包死：任何失败只记 error 日志返回 None，绝不阻塞写入主流程。
    """
    batch_id = str(batch_id)
    with _batch_lock(batch_id):
        try:
            existed = snapshot_path(batch_id).exists()
            snap = build_snapshot(
                batch_id, origin="system" if existed else "baseline")
            if snap is None:
                return None
            if not existed:
                logger.warning(
                    "[快照] 批次 %s 无写入快照，以副本现状为基线（baseline），"
                    "此前人工改动不可追溯", batch_id)
            save_snapshot(batch_id, snap)
            return snap
        except Exception:  # noqa: BLE001 辅助设施失败静默降级，见模块 docstring
            logger.error("[快照] 批次 %s 快照刷新失败（写入主流程不受影响）",
                         batch_id, exc_info=True)
            return None
