# -*- coding: utf-8 -*-
"""F2 报关单统计导出（截单信息-分体积）：按最新分票提案统计每票×管理号
的 件/PSC/净重/毛重/体积，输出 Excel。

数据源：
1. declarations 表该 split_thread_id 最新 version 的全部票
   （pending/confirmed 均可；无记录抛 ValueError("分票提案不存在")）；
2. output/{batch_id}/containers/ 下 mtime 最新的 *_filled.xlsx
   （没有抛 FileNotFoundError("装箱单不存在")），batch_id = split_thread_id
   去掉 "split-" 前缀。

票内行过滤复用 declare/aggregator.py 的 rows_for_ticket（与报关单同口径），
归一化 + SKU 级商检标注镜像 declare/service.py 的调用方式。

体积算法（纯 Python，零 LLM）：
- 票内含 MX 行（is_mx）→ 该票体积列一律留空；
- 管理号只出现在一张票 → 体积 = 该管理号 M3 列唯一值
  （多个不同非空 M3 → 取第一个 + 中文警告）；
- 管理号拆到多张票 → 按 ticket_no 排序，票号最大者 = 柜总 M3 − 其余票
  分体积之和（余额法）；其余票按公式
  Σ(SOTOBAKO_D_HACCHU_SU × 外箱宽×深×高 ÷ 10⁶)（cm→m³），三位小数；
- 长宽高缺失或为 0 的行 → 跳过不计 + 中文警告（含管理号和票号）；
- 余额 ≤ 0 → 中文警告，仍输出数值。

输出：output/{batch_id}/stats/截单信息-分体积_{safe(batch_id)}.xlsx，
覆盖同名文件。版面：每港口一个 sheet（sheet 名=港口名，顺序按票号首现
顺序）；单栏顺序 block：票号标题行 → 表头（管理号|柜型|件|净重|毛重|
体积|PSC）→ 明细行 → 空行；票按 ticket_no 排序。
全程 openpyxl，禁止 pandas 写入；格式复用 writer._apply_write_format。
"""
from __future__ import annotations

import logging
import re
from pathlib import Path

from openpyxl import Workbook

from app.config import batch_stats_dir, get_settings
from app.db.models import Declaration, ProductMapping
from app.db.session import get_session
from app.declare.aggregator import rows_for_ticket
from app.declare.mapping import build_mapping_index
from app.factory_match import load_excel_normalize_map
from app.nodes.writer import (
    _apply_write_format,
    _format_file_busy_msg,
    _is_file_lock_error,
    _probe_writable,
)
from app.split.loader import load_filled_excel
from app.split.normalize import (
    load_sku_inspection_map,
    normalize_maker,
    resolve_inspection,
)
from app.split.schemas import RawItem, Ticket, TicketItem

logger = logging.getLogger(__name__)

HEADERS = ["管理号", "柜型", "件", "净重", "毛重", "体积", "PSC"]
_COL_WIDTHS = [16, 10, 8, 10, 10, 10, 8]

# Excel sheet 名非法字符（[]:*?/\），替换为下划线；长度上限 31
_SHEET_NAME_UNSAFE = re.compile(r"[\[\]:*?/\\]")


# ---------------------------------------------------------------------------
# 数据加载
# ---------------------------------------------------------------------------

def _load_declarations(split_thread_id: str) -> list[Declaration]:
    """取该 split_thread_id 最新 version 的全部票（pending/confirmed）。

    按 ticket_no 排序返回；无记录抛 ValueError("分票提案不存在")。
    """
    with get_session() as sess:
        decls = (
            sess.query(Declaration)
            .filter(
                Declaration.split_thread_id == split_thread_id,
                Declaration.status.in_(["pending", "confirmed"]),
            )
            .all()
        )
        if not decls:
            raise ValueError("分票提案不存在")
        max_version = max(d.version for d in decls)
        decls = [d for d in decls if d.version == max_version]
        # expire_on_commit=False，session 关闭后字段仍可用
        decls.sort(key=lambda d: d.ticket_no)
        return decls


def _resolve_filled_excel(batch_id: str) -> Path:
    """output/{batch_id}/containers/ 下 mtime 最新的 *_filled.xlsx。"""
    containers_dir = get_settings().batch_containers_dir(batch_id)
    if containers_dir.is_dir():
        filled = [p for p in containers_dir.glob("*_filled.xlsx") if p.is_file()]
        if filled:
            latest = max(filled, key=lambda p: p.stat().st_mtime)
            logger.info("[分体积统计] 数据源取 containers 最新 filled：%s", latest)
            return latest
    raise FileNotFoundError("装箱单不存在")


def _load_items(split_thread_id: str, batch_id: str) -> list[RawItem]:
    """读 filled Excel + 归一化 + SKU 级商检标注（镜像 declare/service.py）。"""
    source = _resolve_filled_excel(batch_id)
    raw_items = load_filled_excel(source)
    if not raw_items:
        raise ValueError(f"filled Excel 无数据行: {source}")

    with get_session() as sess:
        from sqlalchemy.orm import selectinload
        mapping_index = build_mapping_index(
            sess.query(ProductMapping)
            .options(selectinload(ProductMapping.sku_links))
            .all()
        )
        sku_map = load_sku_inspection_map(sess)

    normalize_map = load_excel_normalize_map()
    for r in raw_items:
        r.maker = normalize_maker(r.maker, normalize_map)
    for r in raw_items:
        r.inspection = resolve_inspection(
            r.maker, r.sku, r.name_cn, sku_map, mapping_index
        )
    logger.info(
        "[分体积统计] split_thread_id=%s 数据源 %s，共 %d 行",
        split_thread_id, source, len(raw_items),
    )
    return raw_items


# ---------------------------------------------------------------------------
# 票内聚合
# ---------------------------------------------------------------------------

class _KanriAgg:
    """一票×一管理号的聚合容器。"""

    __slots__ = ("kanri_no", "container_type", "pcs", "pieces",
                 "net", "gross", "has_net", "has_gross", "rows")

    def __init__(self, kanri_no: str):
        self.kanri_no = kanri_no
        self.container_type = ""
        self.pcs = 0
        self.pieces = 0
        self.net = 0.0
        self.gross = 0.0
        self.has_net = False
        self.has_gross = False
        self.rows: list[RawItem] = []


def _aggregate_ticket_rows(rows: list[RawItem]) -> list[_KanriAgg]:
    """票内按管理号聚合（保持管理号首现顺序）。

    件 = Σ SOTOBAKO_D_HACCHU_SU（pcs）；PSC = Σ D_HACCHU_SU（qty_pieces）；
    净重/毛重 = 过滤行求和（空值按 0 计，整组皆空则 has_* =False → 输出空单元格）；
    柜型 = 该管理号行 CONTAINER_MEI 首个非空值。
    """
    order: list[str] = []
    aggs: dict[str, _KanriAgg] = {}
    for r in rows:
        a = aggs.get(r.kanri_no)
        if a is None:
            a = aggs[r.kanri_no] = _KanriAgg(r.kanri_no)
            order.append(r.kanri_no)
        a.rows.append(r)
        if not a.container_type and r.container_type:
            a.container_type = r.container_type
        a.pcs += r.pcs or 0
        a.pieces += r.qty_pieces or 0
        if r.net_weight is not None:
            a.has_net = True
            a.net += r.net_weight
        if r.gross_weight is not None:
            a.has_gross = True
            a.gross += r.gross_weight
    return [aggs[k] for k in order]


# ---------------------------------------------------------------------------
# 体积算法
# ---------------------------------------------------------------------------

def _unique_m3(rows: list[RawItem], kanri_no: str,
               warnings: list[str]) -> float | None:
    """管理号 M3 列唯一非空值；多个不同值取第一个 + 警告；全空返回 None。"""
    vals: list[float] = []
    for r in rows:
        if r.m3 is not None and r.m3 not in vals:
            vals.append(r.m3)
    if not vals:
        return None
    if len(vals) > 1:
        warnings.append(
            f"管理号 {kanri_no} 存在多个不同 M3 值"
            f"（{'、'.join(str(v) for v in vals)}），取第一个 {vals[0]}"
        )
    return vals[0]


def _formula_volume(agg: _KanriAgg, ticket_no: str,
                    warnings: list[str]) -> float:
    """公式分体积：Σ(件 × 宽 × 深 × 高 ÷ 10⁶)，三位小数。

    长宽高缺失或为 0 的行跳过不计并记中文警告（含管理号和票号）。
    """
    total = 0.0
    skipped = 0
    for r in agg.rows:
        dims = (r.carton_width, r.carton_depth, r.carton_height)
        if any(d is None or d == 0 for d in dims):
            skipped += 1
            continue
        total += (r.pcs or 0) * dims[0] * dims[1] * dims[2] / 1e6
    if skipped:
        warnings.append(
            f"管理号 {agg.kanri_no}（票 {ticket_no}）：{skipped} 行长宽高"
            "缺失或为 0，已跳过不计体积"
        )
    return round(total, 3)


def _compute_volumes(
    ticket_aggs: dict[str, list[_KanriAgg]],
    mx_tickets: set[str],
    warnings: list[str],
) -> dict[tuple[str, str], float | None]:
    """计算每 (票号, 管理号) 的体积；MX 票一律 None（留空）。

    Args:
        ticket_aggs: {票号: 票内管理号聚合（首现顺序）}
        mx_tickets: 含 MX 行的票号集合
    """
    volumes: dict[tuple[str, str], float | None] = {}

    # 管理号 → 出现的票号列表
    kanri_tickets: dict[str, list[str]] = {}
    for ticket_no, aggs in ticket_aggs.items():
        for a in aggs:
            kanri_tickets.setdefault(a.kanri_no, []).append(ticket_no)

    for kanri_no, tnos in kanri_tickets.items():
        tnos_sorted = sorted(tnos)
        agg_by_ticket = {
            tno: next(a for a in ticket_aggs[tno] if a.kanri_no == kanri_no)
            for tno in tnos_sorted
        }
        all_rows = [r for tno in tnos_sorted for r in agg_by_ticket[tno].rows]

        if len(tnos_sorted) == 1:
            # 整柜票：M3 列原值；MX 票留空
            tno = tnos_sorted[0]
            if tno in mx_tickets:
                volumes[(tno, kanri_no)] = None
                continue
            m3 = _unique_m3(all_rows, kanri_no, warnings)
            if m3 is None:
                warnings.append(f"管理号 {kanri_no} 无 M3 值，体积留空")
            volumes[(tno, kanri_no)] = m3
            continue

        # 拆票：票号最大者走余额法，其余按公式
        total_m3 = _unique_m3(all_rows, kanri_no, warnings)
        formula: dict[str, float | None] = {}
        for tno in tnos_sorted[:-1]:
            if tno in mx_tickets:
                formula[tno] = None
            else:
                formula[tno] = _formula_volume(agg_by_ticket[tno], tno, warnings)
        last = tnos_sorted[-1]
        if last in mx_tickets:
            formula[last] = None
        elif total_m3 is None:
            warnings.append(
                f"管理号 {kanri_no} 拆 {len(tnos_sorted)} 票但柜总 M3 为空，"
                f"余额票 {last} 体积留空"
            )
            formula[last] = None
        else:
            others = sum(v for v in formula.values() if v is not None)
            balance = total_m3 - others
            if balance <= 0:
                warnings.append(
                    f"管理号 {kanri_no} 余额票 {last} 体积余额 "
                    f"{round(balance, 3)} ≤ 0（柜总 M3={total_m3}，"
                    f"其余票合计={round(others, 3)}），请人工核对"
                )
            formula[last] = balance
        for tno, v in formula.items():
            volumes[(tno, kanri_no)] = v

    return volumes


# ---------------------------------------------------------------------------
# Excel 输出
# ---------------------------------------------------------------------------

def _safe_sheet_name(port: str, used: set[str]) -> str:
    """sheet 名=港口名，替换非法字符、截断 31、撞名加序号。"""
    base = _SHEET_NAME_UNSAFE.sub("_", port).strip() or "未命名港口"
    base = base[:31]
    name = base
    i = 2
    while name in used:
        suffix = f"_{i}"
        name = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def _write_workbook(
    decls: list[Declaration],
    ticket_aggs: dict[str, list[_KanriAgg]],
    volumes: dict[tuple[str, str], float | None],
    out_path: Path,
) -> None:
    """按港口分 sheet 写 block 版面，覆盖同名文件。"""
    wb = Workbook()
    used_names: set[str] = set()

    # 港口分组，sheet 顺序按票号（已排序）首现顺序
    ports_order: list[str] = []
    by_port: dict[str, list[Declaration]] = {}
    for d in decls:
        if d.port not in by_port:
            by_port[d.port] = []
            ports_order.append(d.port)
        by_port[d.port].append(d)

    for idx, port in enumerate(ports_order):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = _safe_sheet_name(port, used_names)
        for j, w in enumerate(_COL_WIDTHS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w

        row_idx = 1
        for d in by_port[port]:  # decls 已按 ticket_no 排序
            # 票号标题行
            cell = ws.cell(row=row_idx, column=1, value=d.ticket_no)
            _apply_write_format(cell)
            row_idx += 1
            # 表头行
            for j, h in enumerate(HEADERS, start=1):
                _apply_write_format(ws.cell(row=row_idx, column=j, value=h))
            row_idx += 1
            # 明细行
            for a in ticket_aggs.get(d.ticket_no, []):
                volume = volumes.get((d.ticket_no, a.kanri_no))
                if volume is not None:
                    volume = round(volume, 3)
                values = [
                    a.kanri_no,
                    a.container_type or d.container_type,
                    a.pcs,
                    round(a.net, 2) if a.has_net else None,
                    round(a.gross, 2) if a.has_gross else None,
                    volume,
                    a.pieces,
                ]
                for j, v in enumerate(values, start=1):
                    _apply_write_format(ws.cell(row=row_idx, column=j, value=v))
                row_idx += 1
            # block 末尾空行
            row_idx += 1

    _probe_writable(out_path)
    try:
        wb.save(out_path)
    except (PermissionError, OSError) as e:
        if _is_file_lock_error(e):
            raise RuntimeError(_format_file_busy_msg(out_path)) from e
        raise


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------

def generate_split_stats(split_thread_id: str) -> dict:
    """生成「截单信息-分体积」统计 Excel。

    Returns:
        {"file_path": str, "warnings": list[str]}

    Raises:
        ValueError: 分票提案不存在 / filled Excel 无数据行；
        FileNotFoundError: 装箱单不存在。
    """
    batch_id = split_thread_id.removeprefix("split-")
    warnings: list[str] = []

    decls = _load_declarations(split_thread_id)
    raw_items = _load_items(split_thread_id, batch_id)

    # sj_map 派生量（rows_for_ticket 兼容签名用，新口径读 RawItem.inspection）
    sj_map: dict[str, bool] = {}
    for r in raw_items:
        if r.maker:
            sj_map[r.maker] = sj_map.get(r.maker, False) or r.inspection

    # 每票过滤行 → 票内管理号聚合；MX 票集合
    ticket_aggs: dict[str, list[_KanriAgg]] = {}
    mx_tickets: set[str] = set()
    for d in decls:
        ticket = Ticket(
            ticket_no=d.ticket_no,
            port=d.port,
            container_type=d.container_type,
            items=[TicketItem(**it) for it in d.items],
        )
        rows = rows_for_ticket(ticket, raw_items, sj_map)
        if any(r.is_mx for r in rows):
            mx_tickets.add(d.ticket_no)
        ticket_aggs[d.ticket_no] = _aggregate_ticket_rows(rows)

    volumes = _compute_volumes(ticket_aggs, mx_tickets, warnings)

    out_dir = batch_stats_dir(batch_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"截单信息-分体积_{get_settings().safe_path_tag(batch_id)}.xlsx"
    _write_workbook(decls, ticket_aggs, volumes, out_path)

    logger.info(
        "[分体积统计] split_thread_id=%s → %s（%d 票，%d 条警告）",
        split_thread_id, out_path, len(decls), len(warnings),
    )
    return {"file_path": str(out_path), "warnings": warnings}
