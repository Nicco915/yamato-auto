# -*- coding: utf-8 -*-
"""F1 XD 透视导出：装箱单按 港口 → 管理号 做类 Excel 透视聚合，输出 Excel。

数据源（每次导出实时重读，不缓存）：
1. 优先 output/{batch_id}/containers/ 下 mtime 最新的 *_filled.xlsx；
2. 不存在则回退批次登记的下游原件（batch_store.get_batch 的
   downstream_file_path）；批次不存在抛 ValueError("批次不存在")；
3. 两者都没有抛 FileNotFoundError("装箱单不存在")（API 层映射 409）。

聚合（纯 Python，openpyxl 只读读取 + 写入，禁止 pandas 写入）：
- 按表头定位列：MINATO_MEI_KJ（港口）/ KANRI_NO（管理号）/
  CONTAINER_MEI（柜型）/ M3 / SOTOBAKO_D_HACCHU_SU（件数）；
- 忽略 KANRI_NO 为空的行；
- 按 港口（首现顺序）→ 管理号（首现顺序）分组，件数求和；
- M3 取该管理号下唯一非空值，多个不同值取第一个并产生中文警告；
- 柜型取该管理号下首个非空值。

输出：output/{batch_id}/pivot/XD透视_{safe(batch_id)}.xlsx，重复导出覆盖同名文件。
版面：表头 港口|管理号|柜型|M3|汇总；港口名只在每组第一行显示；每港口末尾
「{港口} 汇总」行（件数合计在 E 列）；最后一行「总计」。
"""
from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.config import batch_pivot_dir, get_settings
from app.db.batch_store import get_batch
from app.nodes.writer import _apply_write_format, _probe_writable

logger = logging.getLogger(__name__)

# 下游装箱单关键列（按表头文本定位，不按列序号，模板列序变动也能工作）
COL_PORT = "MINATO_MEI_KJ"      # 港口
COL_KANRI = "KANRI_NO"          # 管理号
COL_CONTAINER = "CONTAINER_MEI"  # 柜型
COL_M3 = "M3"                   # 体积
COL_QTY = "SOTOBAKO_D_HACCHU_SU"  # 件数（外箱发注数量）

REQUIRED_COLS = (COL_PORT, COL_KANRI, COL_CONTAINER, COL_M3, COL_QTY)

# 表头扫描上限：装箱单表头一般在第 1 行，留余量防前置说明行
_HEADER_SCAN_ROWS = 20


def _resolve_source(batch_id: str) -> Path:
    """解析数据源文件路径：优先 containers 下最新 *_filled.xlsx，回退批次原件。"""
    containers_dir = get_settings().batch_containers_dir(batch_id)
    if containers_dir.is_dir():
        filled = [p for p in containers_dir.glob("*_filled.xlsx") if p.is_file()]
        if filled:
            latest = max(filled, key=lambda p: p.stat().st_mtime)
            logger.info("[XD透视] 数据源取 containers 最新 filled：%s", latest)
            return latest
    batch = get_batch(batch_id)
    if batch is None:
        raise ValueError("批次不存在")
    downstream = batch.get("downstream_file_path")
    if downstream:
        path = Path(downstream)
        if path.is_file():
            logger.info("[XD透视] 数据源回退批次登记原件：%s", path)
            return path
    raise FileNotFoundError("装箱单不存在")


def _locate_header(ws) -> tuple[int, dict[str, int]]:
    """在前 _HEADER_SCAN_ROWS 行内找包含全部必需列的表头行。

    返回 (表头行号, {列名: 列序号(1-based)})；找不到抛 ValueError。
    """
    for row in ws.iter_rows(min_row=1, max_row=_HEADER_SCAN_ROWS):
        col_map: dict[str, int] = {}
        for cell in row:
            if cell.value is None:
                continue
            name = str(cell.value).strip()
            if name in REQUIRED_COLS and name not in col_map:
                col_map[name] = cell.column
        if all(c in col_map for c in REQUIRED_COLS):
            return row[0].row, col_map
    raise ValueError("装箱单缺少必要表头列（港口/管理号/柜型/M3/件数）")


def _to_number(value) -> float:
    """件数/M3 单元格转 float；空值或非数值按 0。"""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return 0.0


def _clean_number(value: float):
    """整数值写 int，避免 Excel 里显示 150.0。"""
    return int(value) if float(value).is_integer() else value


def _aggregate(source: Path) -> tuple[list[dict], list[str]]:
    """读取装箱单并聚合：返回 ([{港口, 管理号: [{柜型, M3, 件数}]...}], warnings)。

    返回结构保持首现顺序：list of {"port": str, "orders": list of
    {"kanri": str, "container": str|None, "m3": float|None, "qty": float}}。
    """
    warnings: list[str] = []
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        header_row, col = _locate_header(ws)
        # 管理号 → 聚合状态；港口 → 管理号有序列表（dict 保插入序）
        ports: dict[str, dict[str, dict]] = {}
        for row in ws.iter_rows(min_row=header_row + 1):
            def val(name):  # noqa: B023 闭包读取当前行
                idx = col[name]
                # openpyxl read_only 下 iter_rows 返回从 A 列开始的元组
                return row[idx - 1].value if idx - 1 < len(row) else None

            kanri_raw = val(COL_KANRI)
            kanri = str(kanri_raw).strip() if kanri_raw is not None else ""
            if not kanri:
                continue  # 忽略管理号为空的行
            port_raw = val(COL_PORT)
            port = str(port_raw).strip() if port_raw is not None else ""
            container_raw = val(COL_CONTAINER)
            container = (
                str(container_raw).strip() if container_raw is not None else ""
            )
            m3_raw = val(COL_M3)
            m3 = _to_number(m3_raw) if m3_raw is not None else None
            qty = _to_number(val(COL_QTY))

            port_orders = ports.setdefault(port, {})
            order = port_orders.get(kanri)
            if order is None:
                order = {
                    "kanri": kanri,
                    "container": container or None,
                    "m3": m3,
                    "m3_values": [m3] if m3 is not None else [],
                    "qty": qty,
                }
                port_orders[kanri] = order
            else:
                order["qty"] += qty
                if order["container"] is None and container:
                    order["container"] = container
                if m3 is not None:
                    # M3 取唯一值；多个不同非空值取第一个 + 中文警告
                    if order["m3"] is None:
                        order["m3"] = m3
                        order["m3_values"].append(m3)
                    elif m3 not in order["m3_values"]:
                        order["m3_values"].append(m3)
                        warnings.append(
                            f"管理号 {kanri} 存在多个不同的 M3 值"
                            f"（{order['m3_values']}），已取第一个 {order['m3']}"
                        )
                        logger.warning("[XD透视] %s", warnings[-1])
    finally:
        wb.close()

    result = [
        {"port": port, "orders": list(orders.values())}
        for port, orders in ports.items()
    ]
    return result, warnings


def _write_pivot(groups: list[dict], out_path: Path) -> None:
    """写透视 Excel：表头 + 分组明细 + 港口小计 + 总计。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    def put(row: int, col_idx: int, value) -> None:
        cell = ws.cell(row=row, column=col_idx, value=value)
        _apply_write_format(cell)

    for c, name in enumerate(("港口", "管理号", "柜型", "M3", "汇总"), start=1):
        put(1, c, name)

    row_no = 2
    grand_total = 0.0
    for group in groups:
        port = group["port"]
        port_total = 0.0
        first = True
        for order in group["orders"]:
            put(row_no, 1, port if first else None)  # 港口名只在每组第一行显示
            put(row_no, 2, order["kanri"])
            put(row_no, 3, order["container"])
            put(row_no, 4, _clean_number(order["m3"]) if order["m3"] is not None else None)
            put(row_no, 5, _clean_number(order["qty"]))
            port_total += order["qty"]
            first = False
            row_no += 1
        put(row_no, 1, f"{port} 汇总")
        for c in (2, 3, 4):
            put(row_no, c, None)
        put(row_no, 5, _clean_number(port_total))
        grand_total += port_total
        row_no += 1
    put(row_no, 1, "总计")
    for c in (2, 3, 4):
        put(row_no, c, None)
    put(row_no, 5, _clean_number(grand_total))

    _probe_writable(out_path)  # Windows 文件锁探测：占用时给中文提示
    wb.save(out_path)
    logger.info("[XD透视] 已写出 %s（%d 个港口组）", out_path, len(groups))


def generate_pivot(batch_id: str) -> dict:
    """生成 XD 透视导出 Excel。

    返回 {"file_path": str, "warnings": list[str]}。
    异常：ValueError("批次不存在") / FileNotFoundError("装箱单不存在") /
    ValueError("装箱单缺少必要表头列...")。
    """
    source = _resolve_source(batch_id)
    groups, warnings = _aggregate(source)

    out_dir = batch_pivot_dir(batch_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = get_settings().safe_path_tag(batch_id)
    out_path = out_dir / f"XD透视_{safe}.xlsx"
    _write_pivot(groups, out_path)
    return {"file_path": str(out_path), "warnings": warnings}
