# -*- coding: utf-8 -*-
"""原件差异检测三入口 + 完全重跑警告测试
（写入侧人工改动感知 · 阶段三 §3.6，决策 9/10）。

覆盖（对应计划 §3.6 五项 + 纯函数单测）：
1. 指纹未变 → 跳过解析（spy 断言 parse_source_rows 未调用）；
   指纹变 → 解析并给出按厂差异（数量改）；
2. 补充工厂入口：原件给已完成工厂加一行 → 未确认 SourceChangesError
   未入图 + 路由 409 结构体（code=source_changes）；确认后 state 吸收
   （row_map 合并）+ 副本重写（新行追加 + 三列通道 A 写入）+ 指纹重记
   + 审计 source_absorbed；
3. 单厂重试入口：指纹变 + 该厂数量改 → 未确认 SourceChangesError /
   dispatcher needs_confirm；确认后吸收进 state 并继续重试；
4. reopen 入口：指纹变 → payload 带 source_changes_warning 且不阻断；
   指纹未变 → None；
5. 完全重跑警告：副本有人工改动痕迹 → 未确认 ManualDiscardError +
   路由 409（code=discard_manual_changes）且未归档；确认后放行归档；
   dispatcher update_downstream needs_confirm → confirm 后 applied。

运行：
    cd app && PYTHONPATH=. python3 -m pytest tests/test_source_diff.py -v

隔离红线（2026-08-11 事故教训，骨架照抄 test_write_diff_conflicts.py）：
- import app 模块前先设 YAMATO_TEST_MODE=1 + YAMATO_DOTENV_PATH=临时空 .env；
- import 后调 validation/_test_isolation.isolate_to_tmp 把
  checkpoint/master/output/sessions 全部指向临时目录。
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

import pytest

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

# ---- 隔离门①：import 任何 app 模块前的双门（血泪红线）----
_TMP = Path(tempfile.mkdtemp(prefix="yamato_source_diff_test_"))
(_TMP / ".env").write_text("# isolated .env\n", encoding="utf-8")
os.environ["YAMATO_TEST_MODE"] = "1"
os.environ["YAMATO_DOTENV_PATH"] = str(_TMP / ".env")
os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from app.api import service  # noqa: E402
from app.api.main import app  # noqa: E402  确保全部 app 模块先于隔离 import
from app.config import get_settings  # noqa: E402
from app.db import batch_store  # noqa: E402
from app.db.models import ReviewAudit  # noqa: E402
from app.db.session import get_session  # noqa: E402
from app.dispatcher import tools as dispatcher_tools  # noqa: E402
from app.orchestrator import copy_snapshot, source_diff  # noqa: E402

# ---- 隔离门②：import 后把 db/output/sessions 指向临时目录 ----
from _test_isolation import isolate_to_tmp  # noqa: E402

TMP = isolate_to_tmp("yamato_source_diff_test_")

client = TestClient(app)


@pytest.fixture(autouse=True)
def _settings_pin():
    """settings 单例钉根：全量混跑时其他测试文件的 isolate_to_tmp 会把
    output_dir 重钉到它们自己的临时目录——每个用例前把 output_dir 钉回
    本模块的 TMP，结束后恢复原值（照 test_write_diff_conflicts 先例）。"""
    s = get_settings()
    original = s.output_dir
    s.output_dir = str(TMP / "output")
    try:
        yield
    finally:
        s.output_dir = original


@pytest.fixture(autouse=True)
def _no_auto_backup(monkeypatch):
    """结案自动备份打空（照 test_awaiting_factories 先例，防全量混跑撞锁）。"""
    from app.orchestrator import backup
    monkeypatch.setattr(backup, "trigger_async", lambda reason, **kw: True)
    yield


FA = "山東中地"          # 已完成/在批工厂
SKU1 = "4550596211831"
SKU2 = "4900000009201"

HEADERS = ["PURCHASE_ORDER", "KANRI_NO", "MAKER_MEI_KJ", "SHOHIN_CD",
           "SHOHIN_MEI_E", "SOTOBAKO_D_HACCHU_SU", "中文品名", "净重", "毛重"]
COL = {h: HEADERS.index(h) + 1 for h in HEADERS}


def _make_xlsx(path: Path, rows: list[dict]) -> None:
    """按 HEADERS 造 xlsx；rows 为 {列名: 值}，缺列留空。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    wb.save(path)


def _make_copy(batch_id: str, rows: list[dict],
               stem: str | None = None) -> Path:
    """在批次 containers/ 下造 *_filled.xlsx 副本（口径同阶段二测试）。"""
    cdir = get_settings().batch_containers_dir(batch_id)
    cdir.mkdir(parents=True, exist_ok=True)
    stem = stem or f"ContentsOfTheContainer_{batch_id}"
    path = cdir / f"{stem}_filled.xlsx"
    _make_xlsx(path, rows)
    return path


def _ws(path: Path):
    return load_workbook(path).active


def _row(factory=FA, sku=SKU1, qty=10, name_cn="钳子", net=15.0, gross=25.0,
         po="272752", kanri="XC268452"):
    """标准数据行（缺省为「系统已写入」状态：三列有值）。"""
    return {"PURCHASE_ORDER": po, "KANRI_NO": kanri, "MAKER_MEI_KJ": factory,
            "SHOHIN_CD": sku, "SHOHIN_MEI_E": "PLIERS",
            "SOTOBAKO_D_HACCHU_SU": qty, "中文品名": name_cn,
            "净重": net, "毛重": gross}


def _cell(path: Path, row: int, field: str):
    return _ws(path).cell(row=row, column=COL[field]).value


def _set_cell(path: Path, row: int, field: str, value) -> None:
    """模拟人工改单元格。"""
    wb = load_workbook(path)
    wb.active.cell(row=row, column=COL[field]).value = value
    wb.save(path)


class _FakeSnap:
    def __init__(self, values=None, next_nodes=(), tasks=()):
        self.values = values if values is not None else {}
        self.next = tuple(next_nodes)
        self.tasks = tuple(tasks)


class _FakeGraph:
    """记录 update_state / stream 调用次数的假 graph（照阶段二先例）。"""

    def __init__(self, snap: _FakeSnap):
        self._snap = snap
        self.updates: list[dict] = []
        self.stream_calls = 0

    def get_state(self, _cfg):
        return self._snap

    def update_state(self, _cfg, update, **_kw):
        self.updates.append(update)
        self._snap.values.update(update)

    def stream(self, *_a, **_kw):
        self.stream_calls += 1
        return iter([])     # 无事件 → completed 兜底分支


def _setup_batch(batch: str, src_rows: list[dict]):
    """造批次现场：原件 + 副本 + 快照 + 批次配置 + 指纹基线。

    返回 (src, copy_path)。原件 stem 与副本同 stem（sync 定位口径）。
    """
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    _make_xlsx(src, src_rows)
    copy_path = _make_copy(batch, src_rows, stem=stem)
    copy_snapshot.refresh_snapshot(batch)
    batch_store.upsert_batch(batch, status="completed")
    service._write_batch_config(batch, {
        "thread_id": batch,
        "downstream_file_path": str(src),
        "upstream_root": str(TMP / "upstream"),
    })
    source_diff.record_fingerprint(batch, str(src))
    return src, copy_path


# ---------------------------------------------------------------------------
# 0. 纯函数：diff_source_requirements 增/删/数量改 + for_factories/describe
# ---------------------------------------------------------------------------

def test_diff_source_requirements_pure():
    old_row_map = {FA: {SKU1: [2, 3], SKU2: [4]}}
    # SKU1 删一行 + 数量改一处；SKU2 加一行；FB 整厂新增
    fb = "Ｃ．正達工芸品"
    new_rows = {FA: {SKU1: [10.0], SKU2: [5.0, 6.0]},
                fb: {SKU1: [1.0]}}
    old_qty_map = {FA: {SKU1: [7.0, 8.0], SKU2: [5.0]}}  # 行序与旧行对应

    sdiff = source_diff.diff_source_requirements(
        old_row_map, new_rows, old_qty_map=old_qty_map)
    by_factory = {c["factory"]: c for c in sdiff.changes}
    assert by_factory[FA] == {"factory": FA, "added": 1, "removed": 1,
                              "qty_changed": 1}
    assert by_factory[fb] == {"factory": fb, "added": 1, "removed": 0,
                              "qty_changed": 0}
    assert sdiff.any()
    # 工厂名排序确定 + describe 人读
    assert [c["factory"] for c in sdiff.changes] == sorted(by_factory)
    desc = sdiff.describe()
    assert "+1行" in desc and "-1行" in desc and "数量改1处" in desc

    # for_factories 子集过滤
    only_fa = sdiff.for_factories([FA])
    assert [c["factory"] for c in only_fa.changes] == [FA]

    # 无数量基线 → qty 维度跳过（结构维度不受影响）
    sdiff2 = source_diff.diff_source_requirements(old_row_map, new_rows)
    assert all(c["qty_changed"] == 0 for c in sdiff2.changes)
    assert sdiff2.any()

    # 无差异 → any() False / changes 空
    same = source_diff.diff_source_requirements(
        {FA: {SKU1: [2]}}, {FA: {SKU1: [10.0]}},
        old_qty_map={FA: {SKU1: [10.0]}})
    assert not same.any() and same.changes == ()

    # 数量相等口径：None==None，float 容差（10 vs 10.0 相等）
    assert source_diff._qty_equal(None, None)
    assert not source_diff._qty_equal(None, 1.0)
    assert source_diff._qty_equal(10, 10.0)
    assert not source_diff._qty_equal(10, 10.5)


# ---------------------------------------------------------------------------
# 1. 指纹预检：未变 → 跳过解析（零成本短路）；变 → 解析给出差异
# ---------------------------------------------------------------------------

def test_fingerprint_shortcircuit_skips_parse(monkeypatch):
    batch = "T-SD-FP"
    src, _ = _setup_batch(batch, [_row()])

    # 指纹未变：直接短路，不解析原件（spy 断言零调用）
    calls: list[str] = []
    monkeypatch.setattr(source_diff, "parse_source_rows",
                        lambda p: calls.append(str(p)) or {})
    values = {"downstream_row_map": {FA: {SKU1: [2]}}}
    assert service._diff_factory_source_changes(batch, values, [FA]) == []
    assert calls == []

    # 指纹变（改数量 → size/mtime 变）：解析被调用，给出 qty_changed
    monkeypatch.undo()
    _set_cell(src, 2, "SOTOBAKO_D_HACCHU_SU", 20)
    changes = service._diff_factory_source_changes(batch, values, [FA])
    assert changes == [{"factory": FA, "added": 0, "removed": 0,
                        "qty_changed": 1}]

    # 历史批次（无指纹记录）→ 懒建基线返回 False（决策 7 同哲学）
    batch2 = "T-SD-FP-LEGACY"
    stem = f"ContentsOfTheContainer_{batch2}"
    src2 = TMP / "src" / f"{stem}.xlsx"
    _make_xlsx(src2, [_row()])
    service._write_batch_config(batch2, {
        "thread_id": batch2,
        "downstream_file_path": str(src2),
        "upstream_root": str(TMP / "upstream"),
    })
    assert source_diff.source_fingerprint_changed(batch2) is False
    assert source_diff.read_fingerprint(batch2) is not None  # 基线已落盘


# ---------------------------------------------------------------------------
# 2. 入口一·补充工厂：已完成工厂原件加行 → 409 确认门 → 确认后吸收+副本重写
# ---------------------------------------------------------------------------

def test_add_factories_source_changes_gate(monkeypatch):
    batch = "T-SD-ADD"
    src, copy_path = _setup_batch(batch, [_row()])

    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {FA: {
            "factory_name": FA,
            "calculated_items": [
                {"sku": SKU1, "name_cn": "钳子",
                 "calculation": {"calculated_unit_net": 1.5,
                                 "calculated_unit_gross": 2.5}},
                {"sku": SKU2, "name_cn": "扳手",
                 "calculation": {"calculated_unit_net": 2.0,
                                 "calculated_unit_gross": 3.0}},
            ],
        }},
        "pending_factories": [],
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)

    # 批次完成后人工改原件：给已完成工厂 FA 加一行（SKU2，qty=3）
    wb = load_workbook(src)
    wb.active.append(["272900", "XC268453", FA, SKU2, "WRENCH", 3,
                      None, None, None])
    wb.save(src)

    # 未确认：SourceChangesError，未入图（update_state/stream 均未发生）
    with pytest.raises(service.SourceChangesError) as ei:
        service.add_factories_to_batch(batch)
    assert ei.value.changes == [{"factory": FA, "added": 1, "removed": 0,
                                 "qty_changed": 0}]
    assert "是否吸收重新写入" in str(ei.value)
    assert fake.updates == [] and fake.stream_calls == 0

    # 路由层契约：409 + 结构体 detail（code=source_changes，含 changes/count）
    resp = client.post(f"/api/v1/batches/{batch}/add-factories", json={})
    assert resp.status_code == 409
    detail = resp.json()["detail"]
    assert detail["code"] == "source_changes"
    assert detail["count"] == 1
    assert detail["changes"][0]["factory"] == FA
    assert detail["changes"][0]["added"] == 1
    assert fake.stream_calls == 0          # 路由层同样未入图

    # 确认后（前端契约：confirm_source_changes 布尔字段必须被接受）：放行
    resp2 = client.post(f"/api/v1/batches/{batch}/add-factories",
                        json={"confirm_source_changes": True})
    assert resp2.status_code == 200
    result = resp2.json()
    assert result["status"] == "completed"
    assert result["added"] == 0            # 无新工厂，纯吸收
    assert result["source_changes_absorbed"] == [FA]
    assert "已吸收原件变更" in result["message"]

    # state 吸收：row_map 合并原件新行（SKU2），requirements 同步
    assert fake.updates, "确认后应合并 state"
    merged_rm = fake.updates[0]["downstream_row_map"]
    assert merged_rm[FA][SKU1] == [2]
    assert merged_rm[FA][SKU2] == [3]
    merged_req = fake.updates[0]["downstream_requirements"]
    assert merged_req[FA] == [SKU1, SKU2]

    # 副本重写：原件新增行追加副本尾部 + 通道 A 语义写三列
    # （SKU2 单重 2.0/3.0 × qty 3 → 净重 6.0 / 毛重 9.0）
    assert _ws(copy_path).max_row == 3
    assert _cell(copy_path, 3, "MAKER_MEI_KJ") == FA
    assert _cell(copy_path, 3, "SHOHIN_CD") == SKU2
    assert _cell(copy_path, 3, "中文品名") == "扳手"
    assert _cell(copy_path, 3, "净重") == 6.0
    assert _cell(copy_path, 3, "毛重") == 9.0
    # 既有行不受影响
    assert _cell(copy_path, 2, "净重") == 15.0

    # 审计留痕：source_absorbed
    with get_session() as session:
        audits = [a for a in session.query(ReviewAudit)
                  .filter(ReviewAudit.thread_id == batch,
                          ReviewAudit.result_status == "source_absorbed")
                  .all()]
    assert len(audits) == 1

    # 指纹基线已重记：再次调用不再报 source_changes
    assert source_diff.source_fingerprint_changed(batch) is False
    resp3 = client.post(f"/api/v1/batches/{batch}/add-factories", json={})
    assert resp3.status_code == 200
    assert resp3.json()["added"] == 0


def test_add_factories_manual_conflicts_gate_still_works(monkeypatch):
    """门槛顺序回归：source_changes 先于 manual_conflicts（一次 409 只报一种）。

    原件无变化时，既有通道 B 人工冲突门行为不变（阶段二契约回归）。
    """
    batch = "T-SD-ADD-ORDER"
    fb = "Ｃ．正達工芸品"
    src_rows = [
        _row(),
        _row(factory=fb, sku=SKU2, qty=5, name_cn=None, net=None, gross=None,
             po="272900", kanri="XC268453"),
    ]
    src, copy_path = _setup_batch(batch, src_rows)
    _set_cell(copy_path, 3, "净重", 66.0)   # 人工抢先填待补充厂 FB 行净重

    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {FA: {"factory_name": FA, "calculated_items": []}},
        "pending_factories": [],
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)

    # 原件未动（指纹未变）→ source_changes 门不触发；人工冲突门照旧 409
    resp = client.post(f"/api/v1/batches/{batch}/add-factories", json={})
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "manual_conflicts"


# ---------------------------------------------------------------------------
# 3. 入口二·单厂重试：指纹变 + 该厂数量改 → 确认门；确认后吸收继续重试
# ---------------------------------------------------------------------------

def _retry_setup(batch: str):
    """造单厂重试现场：批次挂起在 Node5，当前工厂 FA，原件 qty 10。"""
    src, copy_path = _setup_batch(batch, [_row()])
    fake = _FakeGraph(_FakeSnap(
        values={
            "batch_id": batch,
            "current_factory_data": {"factory_name": FA},
            "downstream_requirements": {FA: [SKU1]},
            "downstream_row_map": {FA: {SKU1: [2]}},
            "downstream_file_path": str(src),
        },
        next_nodes=(service.NODE5,),
    ))
    return src, copy_path, fake


def test_retry_factory_source_changes_gate(monkeypatch):
    batch = "T-SD-RETRY"
    src, _, fake = _retry_setup(batch)
    monkeypatch.setattr(service, "get_graph", lambda: fake)

    # 批次挂起期间人工改原件：FA 数量 10 → 20
    _set_cell(src, 2, "SOTOBAKO_D_HACCHU_SU", 20)

    # 未确认：SourceChangesError（qty_changed=1），未重跑
    with pytest.raises(service.SourceChangesError) as ei:
        service.retry_factory_extraction(batch)
    assert ei.value.changes == [{"factory": FA, "added": 0, "removed": 0,
                                 "qty_changed": 1}]
    assert fake.stream_calls == 0

    # dispatcher 确认门接线：needs_confirm + code + changes（引擎回喂 LLM）
    out = dispatcher_tools._exec_retry_factory({"thread_id": batch})
    assert out["status"] == "needs_confirm"
    assert out["code"] == "source_changes"
    assert "confirm_source_changes=true" in out["message"]
    assert out["changes"][0]["qty_changed"] == 1
    assert fake.stream_calls == 0

    # dispatcher preview 警告条：透出原件变更摘要（不阻断）
    monkeypatch.setattr(service, "get_review_payload",
                        lambda tid: {"factory_name": FA, "items": [{}]})
    preview = dispatcher_tools._preview_retry_factory({"thread_id": batch})
    assert any("confirm_source_changes=true" in w
               for w in preview["warnings"])

    # 确认后：吸收该厂最新需求进 state + 继续重试（completed 兜底分支）
    result = dispatcher_tools._exec_retry_factory(
        {"thread_id": batch, "confirm_source_changes": True})
    assert result["status"] == "completed"
    assert result["source_changes_absorbed"] == [
        {"factory": FA, "added": 0, "removed": 0, "qty_changed": 1}]
    assert fake.stream_calls == 1
    assert fake.updates and fake.updates[0]["force_reextract"] is True
    # state 吸收：该厂 row_map 已按原件重解析并入
    assert fake.updates[0]["downstream_row_map"][FA] == {SKU1: [2]}

    # 指纹基线不在此更新（单厂吸收不代表 state 反映整份原件）——
    # 其余工厂差异仍应由入口一检出
    assert source_diff.source_fingerprint_changed(batch) is True


# ---------------------------------------------------------------------------
# 4. 入口三·reopen：指纹变 → payload 带警告条数据且不阻断；未变 → None
# ---------------------------------------------------------------------------

def _reopen_graph(batch: str) -> _FakeGraph:
    return _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {FA: [{
            "sku": SKU1,
            "extracted_data": {"total_quantity": 10,
                               "total_net_weight": 15.0,
                               "total_gross_weight": 25.0},
            "calculation": {"calculated_unit_net": 1.5},
        }]},
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
    }))


def test_reopen_source_changes_warning(monkeypatch):
    batch = "T-SD-REOPEN"
    src, _ = _setup_batch(batch, [_row()])
    fake = _reopen_graph(batch)
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(service, "_latest_approved_audit",
                        lambda tid, fn: None)

    # 指纹未变：无警告条
    payload = service.reopen_factory_for_edit(batch, FA)
    assert payload is not None and payload["reopen_mode"] is True
    assert payload["source_changes_warning"] is None

    # 人工改原件（FA 数量 10 → 20）→ 警告条数据透出且不阻断 reopen
    _set_cell(src, 2, "SOTOBAKO_D_HACCHU_SU", 20)
    payload2 = service.reopen_factory_for_edit(batch, FA)
    assert payload2 is not None and payload2["reopen_mode"] is True
    warning = payload2["source_changes_warning"]
    assert warning is not None
    assert "原件已在批次运行后被修改" in warning["message"]
    assert warning["changes"] == [{"factory": FA, "added": 0, "removed": 0,
                                   "qty_changed": 1}]

    # 路由层契约：GET reopen 响应透传可选字段（不阻断，200）
    resp = client.get(f"/api/v1/review/{batch}/reopen/{FA}")
    assert resp.status_code == 200
    assert resp.json()["source_changes_warning"]["changes"][0][
        "qty_changed"] == 1


# ---------------------------------------------------------------------------
# 5. 决策 10 完全重跑警告：人工改动痕迹未确认拒绝（未归档）；确认后放行
# ---------------------------------------------------------------------------

def test_rerun_discard_manual_gate(monkeypatch):
    batch = "T-SD-RERUN"
    src, copy_path = _setup_batch(batch, [_row()])
    _set_cell(copy_path, 2, "净重", 99.0)   # 人工改副本 → 痕迹

    settings = get_settings()
    out_dir = settings.batch_output_dir(batch)

    # 痕迹检测：diff value_conflicts 非空
    traces = service.detect_manual_traces(batch)
    assert traces is not None and traces["value_conflicts"] == 1

    # 未确认：ManualDiscardError 拒绝，且无任何副作用（未归档）
    with pytest.raises(service.ManualDiscardError) as ei:
        service.rerun_batch(batch)
    assert "副本上的全部人工改动将丢失" in str(ei.value)
    assert "_history/" in str(ei.value)
    assert ei.value.traces["value_conflicts"] == 1
    assert out_dir.exists() and (out_dir / "batch_config.json").exists()
    assert not settings.history_output_dir(batch).exists()

    # 路由层契约：409 结构体（code=discard_manual_changes，带 traces）
    resp = client.post(f"/api/v1/batches/{batch}/rerun")
    assert resp.status_code == 409
    detail = resp.json()["detail"]
    assert detail["code"] == "discard_manual_changes"
    assert detail["traces"]["value_conflicts"] == 1
    assert out_dir.exists()                  # 路由层同样未归档

    # 确认后：放行——output 归档 _history/、checkpoint 重置、指纹重记
    fake = _FakeGraph(_FakeSnap(values={"batch_id": batch}))
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    resp2 = client.post(f"/api/v1/batches/{batch}/rerun",
                        json={"confirm_discard_manual": True})
    assert resp2.status_code == 200
    assert resp2.json()["status"] == "completed"
    assert fake.updates, "checkpoint 应被重置（as_node=START 合并）"
    assert settings.history_output_dir(batch).exists()
    assert out_dir.exists()                  # 新 run 目录重建
    cfg = json.loads((out_dir / "batch_config.json").read_text(
        encoding="utf-8"))
    assert cfg["run_count"] == 2
    fp = source_diff.read_fingerprint(batch)
    assert fp is not None and fp["path"] == str(src) and fp.get("qty_map")


def test_rerun_no_manual_traces_passes_without_confirm(monkeypatch):
    """无人工改动痕迹时 rerun 不需确认（决策 10 只拦有痕迹场景）。"""
    batch = "T-SD-RERUN-CLEAN"
    _setup_batch(batch, [_row()])
    fake = _FakeGraph(_FakeSnap(values={"batch_id": batch}))
    monkeypatch.setattr(service, "get_graph", lambda: fake)

    result = service.rerun_batch(batch)      # 未确认也放行
    assert result["status"] == "completed"
    assert get_settings().history_output_dir(batch).exists()


def test_update_downstream_discard_manual_gate(monkeypatch):
    """dispatcher update_downstream：full 分支未确认 needs_confirm 且
    不改批次配置；确认后全量重跑放行。"""
    batch = "T-SD-UD"
    src, copy_path = _setup_batch(batch, [_row()])
    _set_cell(copy_path, 2, "净重", 99.0)   # 人工改副本 → 痕迹
    src2 = TMP / "src" / "new_downstream.xlsx"
    _make_xlsx(src2, [_row()])               # 新下游表（内容无所谓，diff mock）

    fake_diff = {
        "error": None, "recommendation": "full", "reason": "列结构变化",
        "structure_changed": True, "added_columns": ["X"],
        "removed_columns": [], "added_factories": [],
        "removed_factories": [], "changed_factories": [],
        "unchanged_factories": [], "factory_count": 1,
    }
    monkeypatch.setattr(
        "app.orchestrator.downstream_diff.compare_downstream",
        lambda _o, _n: dict(fake_diff))
    fake = _FakeGraph(_FakeSnap(values={"batch_id": batch}))
    monkeypatch.setattr(service, "get_graph", lambda: fake)

    # preview：决策 10 警告条透出
    preview = dispatcher_tools._preview_update_downstream(
        {"thread_id": batch, "downstream_file_path": str(src2)}, None)
    assert any("全部人工改动将丢失" in w for w in preview["warnings"])
    assert any("confirm_discard_manual=true" in w
               for w in preview["warnings"])

    # 未确认：needs_confirm，且批次配置未被改动（拦在路径更新之前）
    out = dispatcher_tools._execute_update_downstream(
        {"thread_id": batch, "downstream_file_path": str(src2)}, None)
    assert out["status"] == "needs_confirm"
    assert out["code"] == "discard_manual_changes"
    assert out["traces"]["value_conflicts"] == 1
    assert service._read_batch_config(batch)[
        "downstream_file_path"] == str(src)
    assert fake.stream_calls == 0

    # 确认后：更新路径 + 全量重跑放行
    out2 = dispatcher_tools._execute_update_downstream(
        {"thread_id": batch, "downstream_file_path": str(src2),
         "confirm_discard_manual": True}, None)
    assert out2["status"] == "applied"
    assert out2["message"] == "已全量重识别"
    assert service._read_batch_config(batch)[
        "downstream_file_path"] == str(src2)
    assert fake.stream_calls == 1


if __name__ == "__main__":
    import pytest as _pytest
    _pytest.main([__file__, "-v"])
