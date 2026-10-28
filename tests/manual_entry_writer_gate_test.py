# -*- coding: utf-8 -*-
"""第七阶段：writer 遇值确认门（§1.3/§2.4）测试。

覆盖：
1. baseline 人工预填 × 系统值不同 → interrupt payload type=overwrite_confirm，
   conflicts 含 SKU/Excel 行号/old/new；
2. resume keep → 人工值保留、manual_kept_rows 落账，后续 _write_excel /
   clear_sku_rows 整行跳过；
3. resume overwrite → 写系统值，manual_values_overwritten 与逐格 old/new
   留痕；
4. 缺陷回归：本厂已写入过（factory_outputs 有快照）时，origin=
   manual_added 的人工新增行仍必须进确认门；origin=baseline 则保持
   「仅首次写入采信」的防误判口径。

隔离（血泪红线 2026-08-11）：import app 前设 YAMATO_TEST_MODE +
YAMATO_DOTENV_PATH 临时空 .env；import 后 isolate_to_tmp 把 output/db/
sessions 全部指向临时目录，绝不读写真实生产目录/库。
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
from pathlib import Path

import pytest

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

_TMP = Path(tempfile.mkdtemp(prefix="yamato_manual_gate_test_"))
(_TMP / ".env").write_text("# isolated .env\n", encoding="utf-8")
os.environ["YAMATO_TEST_MODE"] = "1"
os.environ["YAMATO_DOTENV_PATH"] = str(_TMP / ".env")
os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

from openpyxl import Workbook, load_workbook  # noqa: E402

from app.config import get_settings  # noqa: E402
from app.nodes import writer  # noqa: E402
from app.orchestrator import copy_snapshot  # noqa: E402

from _test_isolation import isolate_to_tmp  # noqa: E402

TMP = isolate_to_tmp("yamato_manual_gate_test_")

FACTORY = "第七阶段测试厂"
SKU_GATE = "4900000009301"
SKU_NORMAL = "4900000009302"
OLD_NAME, OLD_NET, OLD_GROSS = "人工品名", 99.0, 88.0

HEADERS = ["MAKER_MEI_KJ", "SHOHIN_CD", "PURCHASE_ORDER", "SHOHIN_MEI_E",
           "SOTOBAKO_D_HACCHU_SU", "中文品名", "净重", "毛重"]
COL = {h: HEADERS.index(h) + 1 for h in HEADERS}


@pytest.fixture(autouse=True)
def _pin_output_dir():
    """全量混跑时其他测试可能重钉 settings.output_dir——每例钉回本模块。"""
    s = get_settings()
    original = s.output_dir
    s.output_dir = str(TMP / "output")
    try:
        yield
    finally:
        s.output_dir = original


def _write_sheet(path: Path, *, prefilled: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)

    def put(row: int, col: str, value) -> None:
        ws.cell(row=row, column=COL[col], value=value)

    put(2, "MAKER_MEI_KJ", FACTORY)
    put(2, "SHOHIN_CD", SKU_GATE)
    put(2, "PURCHASE_ORDER", "272752")
    put(2, "SOTOBAKO_D_HACCHU_SU", 10)
    put(3, "MAKER_MEI_KJ", FACTORY)
    put(3, "SHOHIN_CD", SKU_NORMAL)
    put(3, "PURCHASE_ORDER", "272753")
    put(3, "SOTOBAKO_D_HACCHU_SU", 5)
    if prefilled:
        put(2, "中文品名", OLD_NAME)
        put(2, "净重", OLD_NET)
        put(2, "毛重", OLD_GROSS)
    wb.save(path)


def _make_copy(batch: str) -> tuple[Path, Path]:
    """原件 + output 副本 + 基线快照（副本与原件的三列一致）。"""
    src = TMP / "src" / f"{batch}.xlsx"
    src.parent.mkdir(parents=True, exist_ok=True)
    _write_sheet(src, prefilled=True)
    dst = get_settings().batch_containers_dir(batch) / f"{src.stem}_filled.xlsx"
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)
    assert copy_snapshot.refresh_snapshot(batch) is not None
    return src, dst


def _state(batch: str, src: Path, **extra) -> dict:
    return {
        "batch_id": batch,
        "downstream_file_path": str(src),
        "validation_status": "Approved",
        "current_factory_data": {
            "factory_name": FACTORY,
            "calculated_items": [
                {"sku": SKU_GATE, "name_cn": "系统品名",
                 "calculation": {"calculated_unit_net": 1.0,
                                 "calculated_unit_gross": 2.0}},
                {"sku": SKU_NORMAL, "name_cn": "系统普通品名",
                 "calculation": {"calculated_unit_net": 3.0,
                                 "calculated_unit_gross": 4.0}},
            ],
        },
        "downstream_row_map": {FACTORY: {SKU_GATE: [2], SKU_NORMAL: [3]}},
        **extra,
    }


def _install_interrupt(monkeypatch, decision: str) -> list[dict]:
    captured: list[dict] = []

    def fake_interrupt(payload):
        captured.append(payload)
        return {"decision": decision}

    monkeypatch.setattr(writer, "interrupt", fake_interrupt)
    monkeypatch.setattr(writer, "_upsert_db", lambda state: (0, 0))
    return captured


def _cell(ws, row: int, col: str):
    return ws.cell(row=row, column=COL[col]).value


def _set_snapshot_origin(batch: str, origin: str) -> None:
    snap = copy_snapshot.load_snapshot(batch)
    assert snap is not None
    for row in snap["rows"]:
        row["origin"] = origin
    copy_snapshot.save_snapshot(batch, snap)


def test_gate_keep_retains_manual_values_and_persists_rows(monkeypatch):
    batch = "MANUAL-GATE-KEEP"
    src, dst = _make_copy(batch)
    captured = _install_interrupt(monkeypatch, "keep")

    result = writer.writer(_state(batch, src))

    assert len(captured) == 1
    payload = captured[0]
    assert payload["type"] == "overwrite_confirm"
    assert payload["factory"] == FACTORY
    conflict = next(c for c in payload["conflicts"] if c["sku"] == SKU_GATE)
    assert conflict["excel_row"] == 2
    assert conflict["old"]["中文品名"] == OLD_NAME
    assert conflict["old"]["净重"] == OLD_NET
    assert conflict["new"]["中文品名"] == "系统品名"
    assert conflict["new"]["净重"] == 10.0
    assert conflict["new"]["毛重"] == 20.0

    ws = load_workbook(dst).active
    assert _cell(ws, 2, "中文品名") == OLD_NAME
    assert _cell(ws, 2, "净重") == OLD_NET
    assert _cell(ws, 2, "毛重") == OLD_GROSS

    kept = result["manual_kept_rows"][FACTORY]
    expected_key = copy_snapshot.row_key_str((
        copy_snapshot.norm_key_text(FACTORY),
        copy_snapshot.norm_key_text(SKU_GATE), 1))
    assert kept == [expected_key]
    assert result["factory_outputs"][FACTORY]["manual_override_kept"] == 3

    # 后续写回：保留行整行跳过（即使系统值再次变化）
    state2 = _state(batch, src, manual_kept_rows=result["manual_kept_rows"])
    state2["current_factory_data"]["calculated_items"][0]["calculation"] = {
        "calculated_unit_net": 7.0, "calculated_unit_gross": 8.0}
    writer._write_excel(state2, dst)
    ws = load_workbook(dst).active
    assert _cell(ws, 2, "中文品名") == OLD_NAME
    assert _cell(ws, 2, "净重") == OLD_NET
    assert _cell(ws, 2, "毛重") == OLD_GROSS

    # reopen 清空：保留行同样跳过
    cleared = writer.clear_sku_rows(state2, dst, [SKU_GATE])
    assert cleared == 0
    ws = load_workbook(dst).active
    assert _cell(ws, 2, "中文品名") == OLD_NAME
    assert _cell(ws, 2, "净重") == OLD_NET
    assert _cell(ws, 2, "毛重") == OLD_GROSS


def test_gate_overwrite_writes_system_values_and_traces(monkeypatch):
    batch = "MANUAL-GATE-OVERWRITE"
    src, dst = _make_copy(batch)
    captured = _install_interrupt(monkeypatch, "overwrite")

    result = writer.writer(_state(batch, src))

    assert len(captured) == 1
    assert captured[0]["type"] == "overwrite_confirm"
    ws = load_workbook(dst).active
    assert _cell(ws, 2, "中文品名") == "系统品名"
    assert _cell(ws, 2, "净重") == 10.0
    assert _cell(ws, 2, "毛重") == 20.0

    cur = result["factory_outputs"][FACTORY]
    assert cur["manual_values_overwritten"] == 3
    traces = {(c["field"]): c for c in cur["manual_overwritten_cells"]}
    assert traces["name_cn"]["old"] == OLD_NAME
    assert traces["name_cn"]["new"] == "系统品名"
    assert traces["net"]["old"] == OLD_NET
    assert traces["net"]["new"] == 10.0
    assert traces["gross"]["old"] == OLD_GROSS
    assert traces["gross"]["new"] == 20.0
    assert "manual_kept_rows" not in result


def test_manual_added_row_gates_even_after_factory_written(monkeypatch):
    """缺陷回归：已写入工厂中 origin=manual_added 的人工新增行不得漏门。"""
    batch = "MANUAL-ADDED-GATE"
    src, dst = _make_copy(batch)
    _set_snapshot_origin(batch, "manual_added")
    state = _state(batch, src,
                   factory_outputs={FACTORY: {"calculated_items": []}})

    conflicts, gate_cells = writer._collect_overwrite_conflicts(state, dst)

    assert conflicts, "manual_added 行在已写入工厂中也必须进确认门"
    conflict = next(c for c in conflicts if c["sku"] == SKU_GATE)
    assert conflict["excel_row"] == 2
    assert conflict["old"]["净重"] == OLD_NET
    assert conflict["new"]["净重"] == 10.0
    assert gate_cells


def test_baseline_row_not_gated_after_factory_written(monkeypatch):
    """防误判口径：已写入工厂的 baseline 行不再按人工预填采信。"""
    batch = "BASELINE-NO-GATE"
    src, dst = _make_copy(batch)
    _set_snapshot_origin(batch, "baseline")
    state = _state(batch, src,
                   factory_outputs={FACTORY: {"calculated_items": []}})

    conflicts, gate_cells = writer._collect_overwrite_conflicts(state, dst)

    assert conflicts == []
    assert gate_cells == set()