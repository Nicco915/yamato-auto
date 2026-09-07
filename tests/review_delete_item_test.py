# -*- coding: utf-8 -*-
"""审核页「删除条目」能力回归测试（识别错误的垃圾卡整条剔除）。

背景：AI 把非 SKU 文本（如 cainz-001）误识为 SKU 时，该卡既不在装箱单
该工厂名下、SKU 又过不了 13 位刚性校验 → 卡死提交。删除条目让人工把
识别错误的卡片显式剔除：不写 Excel、不落主库，审计留「人工删除」痕迹。

覆盖：
- _merge_human_items：deleted 项被剔除、其余项正常合并、未返回项原样保留
- _prepare_audit：deleted 项生成扁平结构「人工删除」审计条目并计入 edited_count
- clear_sku_rows：reopen 路径清空指定 SKU 已写入的三列单元格（真实 xlsx）
- apply_reopen_payload：deleted 项不进写盘 items，clear_sku_rows 收到被删 SKU，
  审计 edited_count 含删除条目

运行方式：
    cd app && PYTHONPATH=. python3 -m pytest tests/review_delete_item_test.py -v

隔离：validation/_test_isolation.isolate_to_tmp（血泪红线，绝不碰真实库）。
"""
from __future__ import annotations

import os
import sys
from pathlib import Path
from types import SimpleNamespace

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

import pytest  # noqa: E402

from app.api import service  # noqa: E402
from app.nodes import human_review as hr  # noqa: E402
from app.nodes import writer as writer_mod  # noqa: E402

from _test_isolation import isolate_to_tmp  # noqa: E402

# 隔离必须在 import 全部 app 模块之后（llm_client 的 load_dotenv override 红线）
TMP = isolate_to_tmp("yamato_delete_item_test_")

FACTORY = "荣汇"
SKU_GOOD = "4901234567890"
SKU_BAD = "cainz-001"  # 真实事故：非 SKU 文本被误识为 SKU


def _orig(sku, **ext_over):
    ext = {"total_quantity": 10, "total_net_weight": 100.0,
           "total_gross_weight": 120.0, "source_file": "/up/荣汇/a.xlsx"}
    ext.update(ext_over)
    return {"sku": sku, "extracted_data": ext,
            "calculation": {"calculated_unit_net": 10.0,
                            "calculated_unit_gross": 12.0},
            "status": "Normal"}


# ---------------------------------------------------------------------------
# 1. _merge_human_items：deleted 项被剔除
# ---------------------------------------------------------------------------

def test_merge_deleted_item_dropped():
    original = [_orig(SKU_GOOD), _orig(SKU_BAD, total_net_weight=5.0)]
    human = [
        {"sku": SKU_GOOD, "extracted_data": {"total_net_weight": 100.0}},
        {"sku": SKU_BAD, "deleted": True},
    ]
    merged = hr._merge_human_items(human, original, approved=True,
                                   factory_name=FACTORY)
    skus = [i.get("sku") for i in merged]
    assert SKU_BAD not in skus, f"被删条目不应进 merged: {skus}"
    assert skus == [SKU_GOOD]
    # 其余项正常走合并逻辑（数值改动被收）
    good = merged[0]
    assert good["extracted_data"]["total_net_weight"] == 100.0


def test_merge_deleted_with_orig_sku():
    """删除卡带 orig_sku（改名后又删）也能被剔除。"""
    original = [_orig(SKU_GOOD)]
    human = [{"sku": "9999999999999", "orig_sku": SKU_GOOD, "deleted": True}]
    merged = hr._merge_human_items(human, original, approved=True,
                                   factory_name=FACTORY)
    assert merged == [], f"唯一条目被删后 merged 应为空: {merged}"


def test_merge_unreturned_items_preserved_not_deleted():
    """「未返回原样保留」语义不变：只有显式 deleted 才剔除（防前端漏传误删）。"""
    original = [_orig(SKU_GOOD), _orig("4901234567891")]
    human = [{"sku": SKU_GOOD, "extracted_data": {}}]  # 第二张卡没返回
    merged = hr._merge_human_items(human, original, approved=True,
                                   factory_name=FACTORY)
    assert len(merged) == 2, "未返回的条目必须原样保留"


# ---------------------------------------------------------------------------
# 2. _prepare_audit：deleted 项生成「人工删除」审计条目
# ---------------------------------------------------------------------------

def test_prepare_audit_deleted_entry(monkeypatch):
    monkeypatch.setattr(service, "get_review_payload", lambda tid: {
        "factory_name": FACTORY,
        "items": [_orig(SKU_GOOD), _orig(SKU_BAD)],
    })
    resume = {"approved": True, "items": [
        {"sku": SKU_GOOD, "extracted_data": {"total_net_weight": 100.0}},
        {"sku": SKU_BAD, "deleted": True},
    ]}
    prepared = service._prepare_audit("TID-DEL", resume)
    assert prepared is not None
    del_changes = [c for c in prepared["changes"] if c["field"] == "条目"]
    assert len(del_changes) == 1
    assert del_changes[0]["sku"] == SKU_BAD
    assert del_changes[0]["old"] == "识别条目"
    assert del_changes[0]["new"] == "人工删除"
    # 扁平结构契约：batch.html 审计区按 c.field/c.old/c.new 渲染
    assert set(del_changes[0]) == {"sku", "field", "old", "new"}
    assert prepared["edited_count"] == 1


def test_prepare_audit_deleted_uses_orig_sku(monkeypatch):
    """改名后删除：审计 SKU 记 orig_sku（人认得的原始识别值）。"""
    monkeypatch.setattr(service, "get_review_payload", lambda tid: {
        "factory_name": FACTORY, "items": [_orig(SKU_GOOD)],
    })
    resume = {"approved": True, "items": [
        {"sku": "9999999999999", "orig_sku": SKU_GOOD, "deleted": True},
    ]}
    prepared = service._prepare_audit("TID-DEL2", resume)
    del_changes = [c for c in prepared["changes"] if c["field"] == "条目"]
    assert del_changes[0]["sku"] == SKU_GOOD


# ---------------------------------------------------------------------------
# 3. clear_sku_rows：清空已写入的三列单元格（真实 xlsx）
# ---------------------------------------------------------------------------

def _make_filled_xlsx(path: Path):
    """造一份「已写入三列」的输出表：表头含三列，两行 SKU 均已填值。"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["SHOHIN_CD", "SHOHIN_MEI_E", "SOTOBAKO_D_HACCHU_SU",
               "中文品名", "净重", "毛重"])
    ws.append([SKU_GOOD, "ITEM-A", 10, "好货", 168.0, 180.0])
    ws.append([SKU_BAD, "ITEM-B", 5, "垃圾卡", 50.0, 55.0])
    wb.save(path)


def _fake_state_for_clear():
    return {
        "current_factory_data": {"factory_name": FACTORY},
        "downstream_row_map": {FACTORY: {SKU_GOOD: [2], SKU_BAD: [3]}},
    }


def test_clear_sku_rows_clears_three_columns():
    from openpyxl import load_workbook
    xlsx = TMP / "clear_rows.xlsx"
    _make_filled_xlsx(xlsx)

    cleared = writer_mod.clear_sku_rows(_fake_state_for_clear(), xlsx, [SKU_BAD])
    assert cleared == 1

    ws = load_workbook(xlsx).active
    # 被删 SKU 行：三列全空
    assert [ws.cell(row=3, column=c).value for c in (4, 5, 6)] == [None, None, None]
    # 其余列不动（SKU/品名/件数原样）
    assert ws.cell(row=3, column=1).value == SKU_BAD
    # 未删 SKU 行不受影响
    assert [ws.cell(row=2, column=c).value for c in (4, 5, 6)] == ["好货", 168.0, 180.0]


def test_clear_sku_rows_empty_skus_noop():
    assert writer_mod.clear_sku_rows(_fake_state_for_clear(),
                                     TMP / "nonexistent.xlsx", []) == 0


def test_clear_sku_rows_unknown_sku_noop():
    from openpyxl import load_workbook
    xlsx = TMP / "clear_unknown.xlsx"
    _make_filled_xlsx(xlsx)
    # 装箱单里没有的 SKU（row_map 匹配不到）：天然无处可清，静默跳过
    cleared = writer_mod.clear_sku_rows(_fake_state_for_clear(), xlsx,
                                        ["0000000000000"])
    assert cleared == 0
    ws = load_workbook(xlsx).active
    assert ws.cell(row=3, column=5).value == 50.0  # 原值未动


# ---------------------------------------------------------------------------
# 4. apply_reopen_payload：deleted 不进写盘、clear_sku_rows 收到被删 SKU
# ---------------------------------------------------------------------------

TID = "TEST-REOPEN-DELETE"


class _FakeGraph:
    def __init__(self, values, next_nodes=()):
        self._values = values
        self._next = tuple(next_nodes)
        self.update_calls: list[dict] = []

    def get_state(self, _config):
        return SimpleNamespace(values=self._values, next=self._next, tasks=[])

    def update_state(self, _config, values, as_node=None):
        self.update_calls.append({"values": values, "as_node": as_node})
        self._values.update(values)


@pytest.fixture
def patched_reopen(monkeypatch):
    """打桩 writer 三件套 + 审计；capture 写盘 items 与 clear_sku_rows 入参。"""
    captured: dict = {}
    monkeypatch.setattr(writer_mod, "_ensure_output_copy",
                        lambda state: TMP / "out.xlsx")

    def fake_write(state, path):
        captured["written_items"] = (
            state["current_factory_data"]["calculated_items"])
        return 1
    monkeypatch.setattr(writer_mod, "_write_excel", fake_write)
    monkeypatch.setattr(writer_mod, "_upsert_db", lambda state: (0, 0))

    def fake_clear(state, path, skus):
        captured["cleared_skus"] = list(skus)
        return len(skus)
    monkeypatch.setattr(writer_mod, "clear_sku_rows", fake_clear)
    monkeypatch.setattr(service, "_write_audit",
                        lambda prepared, source: captured.update(
                            audit=prepared))

    values = {
        "batch_id": TID,
        "downstream_file_path": "/tmp/downstream.xlsx",
        "downstream_row_map": {},
        "factory_outputs": {FACTORY: {
            "factory_name": FACTORY,
            "calculated_items": [_orig(SKU_GOOD), _orig(SKU_BAD)],
        }},
        "current_factory_data": {},
    }
    fake = _FakeGraph(values)
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    captured["graph"] = fake
    return captured


def test_reopen_deleted_item_excluded_and_cleared(patched_reopen):
    items = [
        {**_orig(SKU_GOOD), "is_human_edited": True},
        {"sku": SKU_BAD, "deleted": True},
    ]
    out = service.apply_reopen_payload(
        TID, FACTORY, {"approved": True, "items": items})
    assert out["status"] == "success"

    # 写盘 items 不含被删条目
    written_skus = [i.get("sku") for i in patched_reopen["written_items"]]
    assert written_skus == [SKU_GOOD]
    # clear_sku_rows 收到被删 SKU（清已写单元格）
    assert patched_reopen["cleared_skus"] == [SKU_BAD]
    # 快照回写同样不含被删条目
    snap_items = (patched_reopen["graph"].update_calls[0]["values"]
                  ["factory_outputs"][FACTORY]["calculated_items"])
    assert [i.get("sku") for i in snap_items] == [SKU_GOOD]
    # 审计：edited_count 含删除条目，changes 有「人工删除」
    audit = patched_reopen["audit"]
    assert audit["edited_count"] == 2  # 1 个编辑 + 1 个删除
    del_changes = [c for c in audit["changes"] if c["field"] == "条目"]
    assert len(del_changes) == 1 and del_changes[0]["sku"] == SKU_BAD
    assert "删除 1 行" in out["message"]


def test_reopen_no_deleted_clear_not_called_with_skus(patched_reopen):
    items = [{**_orig(SKU_GOOD), "is_human_edited": True}]
    out = service.apply_reopen_payload(
        TID, FACTORY, {"approved": True, "items": items})
    assert out["status"] == "success"
    assert patched_reopen.get("cleared_skus") == []
    assert "删除" not in out["message"]
