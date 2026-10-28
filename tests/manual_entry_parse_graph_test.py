# -*- coding: utf-8 -*-
"""第七阶段：Node1 手动跳过剔除 / 副本提前生成 / 全手动路由 / 导出口径。

覆盖：
1. Node1 把 manual_skip_factories 从 pending 剔除，但保留在
   requirements/row_map；未命中装箱单的名字被规范化丢弃；
2. Node1 解析后无条件生成 output 副本 + 基线快照（含人工已填三列，
   原件不被修改）；
3. graph：全手动批次 Node1 后直走 Node7（不出现「未知工厂」）；混合批次
   仍进 Node2；
4. export_node：手动录入厂不计 factories_failed，
   manual_entry_factories 透出。

隔离（血泪红线 2026-08-11）：import app 前设 YAMATO_TEST_MODE +
YAMATO_DOTENV_PATH 临时空 .env；import 后 isolate_to_tmp 指向临时目录。
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

import pytest

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

_TMP = Path(tempfile.mkdtemp(prefix="yamato_manual_node1_test_"))
(_TMP / ".env").write_text("# isolated .env\n", encoding="utf-8")
os.environ["YAMATO_TEST_MODE"] = "1"
os.environ["YAMATO_DOTENV_PATH"] = str(_TMP / ".env")
os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

from langgraph.checkpoint.memory import MemorySaver  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from app.config import get_settings  # noqa: E402
from app.graph import NODE1, NODE2, NODE7, build_graph  # noqa: E402
from app.nodes.export_node import export_node  # noqa: E402
from app.nodes.parse_downstream import parse_downstream  # noqa: E402
from app.orchestrator import copy_snapshot  # noqa: E402

from _test_isolation import isolate_to_tmp  # noqa: E402

TMP = isolate_to_tmp("yamato_manual_node1_test_")

HEADERS = ["MAKER_MEI_KJ", "SHOHIN_CD", "PURCHASE_ORDER",
           "SOTOBAKO_D_HACCHU_SU", "中文品名", "净重", "毛重"]
COL = {h: HEADERS.index(h) + 1 for h in HEADERS}

MANUAL_FACTORY = "手动录入厂"
AUTO_FACTORY = "自动识别厂"
MANUAL_NAME, MANUAL_NET, MANUAL_GROSS = "人工品名", 12.3, 23.4


@pytest.fixture(autouse=True)
def _pin_output_dir():
    s = get_settings()
    original = s.output_dir
    s.output_dir = str(TMP / "output")
    try:
        yield
    finally:
        s.output_dir = original


def _write_manifest(path: Path, rows: list[tuple[str, str, int, bool]]) -> Path:
    """rows = [(工厂, SKU, 数量, 是否预填三列)]。"""
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for factory, sku, qty, prefilled in rows:
        ws.append([factory, sku, "272752", qty, None, None, None])
        row = ws.max_row
        if prefilled:
            ws.cell(row=row, column=COL["中文品名"], value=MANUAL_NAME)
            ws.cell(row=row, column=COL["净重"], value=MANUAL_NET)
            ws.cell(row=row, column=COL["毛重"], value=MANUAL_GROSS)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def test_node1_filters_manual_from_pending_but_keeps_requirements():
    batch = "MANUAL-NODE1-FILTER"
    xlsx = _write_manifest(TMP / f"{batch}.xlsx", [
        (MANUAL_FACTORY, "4900000009401", 10, True),
        (MANUAL_FACTORY, "4900000009402", 20, True),
        (AUTO_FACTORY, "4900000009403", 30, False),
    ])

    result = parse_downstream({
        "batch_id": batch,
        "downstream_file_path": str(xlsx),
        "manual_skip_factories": [MANUAL_FACTORY, "装箱单里没有的厂"],
    })

    assert result["pending_factories"] == [AUTO_FACTORY]
    assert set(result["downstream_requirements"]) == {MANUAL_FACTORY, AUTO_FACTORY}
    assert result["downstream_row_map"][MANUAL_FACTORY]["4900000009401"] == [2]
    assert result["downstream_row_map"][MANUAL_FACTORY]["4900000009402"] == [3]
    # normalize：未命中装箱单的名字直接丢弃，不污染 state
    assert result["manual_skip_factories"] == [MANUAL_FACTORY]


def test_node1_generates_copy_and_baseline_snapshot():
    batch = "MANUAL-NODE1-COPY"
    xlsx = _write_manifest(TMP / f"{batch}_src.xlsx", [
        (MANUAL_FACTORY, "4900000009411", 10, True),
        (AUTO_FACTORY, "4900000009412", 30, False),
    ])
    original_bytes = xlsx.read_bytes()

    result = parse_downstream({
        "batch_id": batch,
        "downstream_file_path": str(xlsx),
        "manual_skip_factories": [MANUAL_FACTORY],
    })

    copy_path = Path(result["final_output_path"])
    assert copy_path.exists()
    assert copy_path != xlsx
    assert copy_path.parent == get_settings().batch_containers_dir(batch)
    assert copy_path.name == f"{xlsx.stem}_filled{xlsx.suffix}"
    # 绝不写原件
    assert xlsx.read_bytes() == original_bytes

    ws = load_workbook(copy_path).active
    assert ws.cell(row=2, column=COL["中文品名"]).value == MANUAL_NAME
    assert ws.cell(row=2, column=COL["净重"]).value == MANUAL_NET
    assert ws.cell(row=2, column=COL["毛重"]).value == MANUAL_GROSS

    snap = copy_snapshot.load_snapshot(batch)
    assert snap is not None
    assert snap["version"] == copy_snapshot.SNAPSHOT_VERSION
    assert snap["rows"], "基线快照必须记录副本数据行"
    assert {r["origin"] for r in snap["rows"]} == {"baseline"}
    manual_row = next(
        r for r in snap["rows"]
        if (r.get("key") or {}).get("sku") == "4900000009411")
    assert manual_row["name_cn"] == MANUAL_NAME
    assert manual_row["net"] == MANUAL_NET
    assert manual_row["gross"] == MANUAL_GROSS


def _manifest_for_graph(batch: str, rows) -> Path:
    return _write_manifest(TMP / f"{batch}.xlsx", rows)


def test_graph_all_manual_goes_straight_to_node7():
    batch = "MANUAL-GRAPH-ALL"
    xlsx = _manifest_for_graph(batch, [
        (MANUAL_FACTORY, "4900000009421", 10, True),
        ("手动录入厂二", "4900000009422", 20, True),
    ])
    graph = build_graph(MemorySaver())
    cfg = {"configurable": {"thread_id": batch}}
    events: list[str] = []
    for event in graph.stream({
        "batch_id": batch,
        "downstream_file_path": str(xlsx),
        "manual_skip_factories": [MANUAL_FACTORY, "手动录入厂二"],
    }, cfg, stream_mode="updates"):
        events.extend(event.keys())

    assert NODE1 in events
    assert NODE7 in events
    assert NODE2 not in events, "全手动批次不得进入 Node2 空队列守卫"
    state = graph.get_state(cfg)
    assert state.values["pending_factories"] == []
    assert "未知工厂" not in json.dumps(
        state.values.get("downstream_requirements") or {}, ensure_ascii=False)


def test_graph_mixed_still_enters_node2():
    batch = "MANUAL-GRAPH-MIXED"
    xlsx = _manifest_for_graph(batch, [
        (MANUAL_FACTORY, "4900000009431", 10, True),
        (AUTO_FACTORY, "4900000009432", 20, False),
    ])
    graph = build_graph(MemorySaver())
    cfg = {"configurable": {"thread_id": batch}}
    stream = graph.stream({
        "batch_id": batch,
        "downstream_file_path": str(xlsx),
        "manual_skip_factories": [MANUAL_FACTORY],
    }, cfg, stream_mode="updates")
    first = next(stream)
    stream.close()

    assert NODE1 in first
    assert first[NODE1]["pending_factories"] == [AUTO_FACTORY]
    assert list(graph.get_state(cfg).next) == [NODE2]


def test_export_node_excludes_manual_from_failed():
    batch = "MANUAL-EXPORT"
    state = {
        "batch_id": batch,
        "downstream_file_path": str(TMP / "manifest.xlsx"),
        "final_output_path": str(TMP / "out.xlsx"),
        "downstream_requirements": {"完成厂": ["1"], "手动厂": ["2"], "失败厂": ["3"]},
        "factory_outputs": {"完成厂": {"calculated_items": []}},
        "manual_skip_factories": ["手动厂"],
        "validation_status": "Approved",
    }

    export_node(state)

    summary = json.loads(
        (get_settings().batch_output_dir(batch) / "export_summary.json")
        .read_text(encoding="utf-8"))
    assert summary["factories_failed"] == 1  # 仅「失败厂」
    assert summary["manual_entry_factories"] == ["手动厂"]