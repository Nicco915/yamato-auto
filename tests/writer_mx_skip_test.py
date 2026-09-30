# -*- coding: utf-8 -*-
"""MX 人工行写回防覆盖回归测试（writer._write_excel / clear_sku_rows）。

背景：MX 货物（下游装箱单 PURCHASE_ORDER 列形如 "MX2-268510-001" 的行）
是人工手动加进装箱单的，中文品名/净重/毛重三列已由人工填好并核对。
writer 按行号写回时不得覆盖；reopen 删除条目清空三列时也不得清掉人工值。

覆盖（单元级直调 writer._write_excel / clear_sku_rows，不跑整图）：
1. _write_excel：MX 行三列原样未动、普通行被写入、返回值只计普通行；
2. clear_sku_rows：MX 行三列保留、普通行被清空、返回值只计普通行；
3. 无 PURCHASE_ORDER 列的表：视为无 MX，全部照常写，不报错。

隔离（血泪红线 2026-08-11，照抄 tests/writer_relink_test.py 头部模板）：
先 import 全部 app 模块，再调 _test_isolation.isolate_to_tmp——
全部路径指向临时目录，绝不碰 app/data/ 真实库。

用法（在 app 根目录下）：
  python3 tests/writer_mx_skip_test.py
  或 PYTHONPATH=. python3 -m pytest tests/writer_mx_skip_test.py -v
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

# ---- env 前置（EXTRACTION_MOCK 需在 import app 之前；db 路径在 import 后隔离）----
os.environ["EXTRACTION_MOCK"] = "1"                      # 提取走 mock，不调 LLM
os.environ["DISPATCHER_MOCK"] = "1"

APP_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

from openpyxl import Workbook, load_workbook  # noqa: E402

from app.nodes import writer  # noqa: E402

from _test_isolation import isolate_to_tmp  # noqa: E402

# 隔离必须在 import app 模块之后（load_dotenv override 红线）；
# 本测试不碰数据库，但 settings 单例仍需钉到临时目录
TMP = isolate_to_tmp("yamato_writer_mx_test_")

FACTORY = "MX测试厂"
SKU_MX = "MXSKU001"
SKU_NORMAL = "4900000009201"

MX_PO = "MX2-268510-001"
NORMAL_PO = "272752"

# MX 行人工已填好的三列值（断言写回后原样未动）
MX_NAME_CN, MX_NET, MX_GROSS = "人工品名MX", 12.34, 23.45

HEADERS = ["PURCHASE_ORDER", "SHOHIN_MEI_E", "中文品名", "净重", "毛重",
           "SOTOBAKO_D_HACCHU_SU"]


def _make_xlsx(path: Path, *, with_po_col: bool = True) -> None:
    """造两行事单：row2 = MX 人工行（三列已填），row3 = 普通行（三列待写）。"""
    wb = Workbook()
    ws = wb.active
    headers = HEADERS if with_po_col else [h for h in HEADERS
                                           if h != "PURCHASE_ORDER"]
    ws.append(headers)

    def put(row: int, name: str, value) -> None:
        ws.cell(row=row, column=headers.index(name) + 1, value=value)

    if with_po_col:
        put(2, "PURCHASE_ORDER", MX_PO)
        put(3, "PURCHASE_ORDER", NORMAL_PO)
    # MX 行：人工已填三列
    put(2, "中文品名", MX_NAME_CN)
    put(2, "净重", MX_NET)
    put(2, "毛重", MX_GROSS)
    # 普通行：数量 10，三列留空
    put(3, "SOTOBAKO_D_HACCHU_SU", 10)
    wb.save(path)


def _state() -> dict:
    """最小 state：MX SKU 映射 row2、普通 SKU 映射 row3。"""
    return {
        "current_factory_data": {
            "factory_name": FACTORY,
            "calculated_items": [
                {"sku": SKU_MX, "name_cn": "系统识别品名MX",
                 "calculation": {"calculated_unit_net": 9.99,
                                 "calculated_unit_gross": 19.99}},
                {"sku": SKU_NORMAL, "name_cn": "系统识别品名",
                 "calculation": {"calculated_unit_net": 1.5,
                                 "calculated_unit_gross": 2.5}},
            ],
        },
        "downstream_row_map": {FACTORY: {SKU_MX: [2], SKU_NORMAL: [3]}},
    }


def _cell(ws, row: int, name: str):
    return ws.cell(row=row, column=HEADERS.index(name) + 1).value


def test_write_excel_skips_mx_rows():
    path = TMP / "mx_write.xlsx"
    _make_xlsx(path)

    written, skipped_placeholder = writer._write_excel(_state(), path)
    assert written == 1, f"返回值应只计普通行: {written}"
    assert skipped_placeholder == 0  # 无占位行（提取失败告警闭环 T2 新契约）

    ws = load_workbook(path).active
    # MX 行：三列原样未动（人工值不被覆盖）
    assert _cell(ws, 2, "中文品名") == MX_NAME_CN
    assert _cell(ws, 2, "净重") == MX_NET
    assert _cell(ws, 2, "毛重") == MX_GROSS
    # 普通行：被写入（净重 = 1.5 × 10，毛重 = 2.5 × 10）
    assert _cell(ws, 3, "中文品名") == "系统识别品名"
    assert _cell(ws, 3, "净重") == 15.0
    assert _cell(ws, 3, "毛重") == 25.0
    print("[断言通过] 场景1：_write_excel 跳过 MX 行——"
          "MX 三列原样、普通行写入、返回值只计普通行")


def test_clear_sku_rows_keeps_mx_rows():
    path = TMP / "mx_clear.xlsx"
    _make_xlsx(path)
    # 普通行先填上系统写过的值（模拟此前已写入）
    ws = load_workbook(path).active
    for name, v in (("中文品名", "系统识别品名"), ("净重", 15.0), ("毛重", 25.0)):
        ws.cell(row=3, column=HEADERS.index(name) + 1, value=v)
    wb = ws.parent
    wb.save(path)

    cleared = writer.clear_sku_rows(_state(), path, [SKU_MX, SKU_NORMAL])
    assert cleared == 1, f"返回值应只计普通行: {cleared}"

    ws = load_workbook(path).active
    # MX 行：人工值保留
    assert _cell(ws, 2, "中文品名") == MX_NAME_CN
    assert _cell(ws, 2, "净重") == MX_NET
    assert _cell(ws, 2, "毛重") == MX_GROSS
    # 普通行：三列被清空
    assert _cell(ws, 3, "中文品名") is None
    assert _cell(ws, 3, "净重") is None
    assert _cell(ws, 3, "毛重") is None
    print("[断言通过] 场景2：clear_sku_rows 跳过 MX 行——"
          "MX 三列保留、普通行清空、返回值只计普通行")


def test_write_excel_without_po_column():
    path = TMP / "no_po_col.xlsx"
    _make_xlsx(path, with_po_col=False)

    written, skipped_placeholder = writer._write_excel(_state(), path)
    assert written == 2, f"无 PO 列应全部照常写: {written}"
    assert skipped_placeholder == 0  # 无占位行（提取失败告警闭环 T2 新契约）

    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "PURCHASE_ORDER" not in headers
    col_cn = headers.index("中文品名") + 1
    # 两行都被写入（含原 MX 行，无 PO 列无法识别视为普通行）
    assert ws.cell(row=2, column=col_cn).value == "系统识别品名MX"
    assert ws.cell(row=3, column=col_cn).value == "系统识别品名"
    print("[断言通过] 场景3：无 PURCHASE_ORDER 列——全部照常写，不报错")


def main():
    test_write_excel_skips_mx_rows()
    test_clear_sku_rows_keeps_mx_rows()
    test_write_excel_without_po_column()
    print("\nwriter_mx_skip_test: PASS")


if __name__ == "__main__":
    main()
