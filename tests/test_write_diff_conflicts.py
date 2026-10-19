# -*- coding: utf-8 -*-
"""写前 diff + 通道分流 + sync 复活修复 + 快照结转 + 数量告警测试
（写入侧人工改动感知 · 阶段二 §2.8）。

覆盖（对应任务 10 项）：
1. 改空覆盖：系统写 V → 人工清空 → 通道 C 重写 → 副本=系统新值；
2. 通道 C 改非空跳过：副本保留人工值 W + 审计 changes_json 含
   manual_override_kept + export_summary 摘要字段 + 快照 protected 打标；
3. 快照空+副本非空（决策 1 定性表第三行）：通道 C 跳过保留 / 通道 A 覆写
   + manual_value_overwritten 留痕（old=人工值）；
4. MX1 回归（删行不复活）：人工删行后 _sync_output_copy_rows 不复活、
   无「错位」报错、墓碑含该行键（快照键与墓碑双口径）；
5. 原件尾部真新增：追加副本尾部（表头名映射整行拷贝、三列留空）、
   快照重建含新行；
6. clear_sku_rows 通道 A：人工非空值随 reopen 删除被清空 + 留痕；
7. 补充工厂确认门（通道 B）：未确认 → ManualConflictError 未入图 +
   路由 409 结构体契约；确认后放行覆写（overwrite_manual_factories 入
   state）；响应恒带 manual_overrides_kept；
8. 快照结转（遗忘漏洞回归）：二次写入不覆盖 protected、人工清空
   protected 单元格保护解除后下次写入覆盖；
9. 数量列告警：人工改数量 → 按新数量写入 + warning + 摘要透出；
10. overwrite-protected 端点：补写系统值 + 审计 manual_value_overwritten
    + 保护解除；无 protected 行 → 400。

运行：
    cd app && PYTHONPATH=. python3 -m pytest tests/test_write_diff_conflicts.py -v

隔离红线（2026-08-11 事故教训，骨架照抄 test_copy_snapshot.py）：
- import app 模块前先设 YAMATO_TEST_MODE=1 + YAMATO_DOTENV_PATH=临时空 .env；
- import 后调 validation/_test_isolation.isolate_to_tmp 把
  checkpoint/master/output/sessions 全部指向临时目录。
"""
from __future__ import annotations

import json
import logging
import os
import sys
import tempfile
from pathlib import Path

import pytest

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

# ---- 隔离门①：import 任何 app 模块前的双门（血泪红线）----
_TMP = Path(tempfile.mkdtemp(prefix="yamato_write_diff_test_"))
(_TMP / ".env").write_text("# isolated .env\n", encoding="utf-8")
os.environ["YAMATO_TEST_MODE"] = "1"
os.environ["YAMATO_DOTENV_PATH"] = str(_TMP / ".env")
os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from app.api import service  # noqa: E402
from app.api.main import app  # noqa: E402  确保全部 app 模块先于隔离 import
from app.config import get_settings  # noqa: E402
from app.db import batch_store  # noqa: E402
from app.db.models import ReviewAudit  # noqa: E402
from app.db.session import get_session  # noqa: E402
from app.nodes import writer  # noqa: E402
from app.nodes.export_node import export_node  # noqa: E402
from app.orchestrator import copy_snapshot  # noqa: E402

# ---- 隔离门②：import 后把 db/output/sessions 指向临时目录 ----
from _test_isolation import isolate_to_tmp  # noqa: E402

TMP = isolate_to_tmp("yamato_write_diff_test_")

client = TestClient(app)


@pytest.fixture(autouse=True)
def _settings_pin():
    """settings 单例钉根：全量混跑时其他测试文件的 isolate_to_tmp 会把
    output_dir 重钉到它们自己的临时目录——每个用例前把 output_dir 钉回
    本模块的 TMP，结束后恢复原值（照 test_copy_snapshot 先例）。"""
    s = get_settings()
    original = s.output_dir
    s.output_dir = str(TMP / "output")
    try:
        yield
    finally:
        s.output_dir = original


@pytest.fixture(autouse=True)
def _no_auto_backup(monkeypatch):
    """结案自动备份打空（照 test_awaiting_factories 先例，防全量混跑撞锁）。"""
    from app.orchestrator import backup
    monkeypatch.setattr(backup, "trigger_async", lambda reason, **kw: True)
    yield


FA = "山東中地"          # 已完成/在批工厂
FB = "Ｃ．正達工芸品"     # 待补充工厂
FMX = "ＭＸ第一工厂"      # MX1 回归用被人工删行的工厂
SKU1 = "4550596211831"
SKU2 = "4900000009201"

HEADERS = ["PURCHASE_ORDER", "KANRI_NO", "MAKER_MEI_KJ", "SHOHIN_CD",
           "SHOHIN_MEI_E", "SOTOBAKO_D_HACCHU_SU", "中文品名", "净重", "毛重"]
COL = {h: HEADERS.index(h) + 1 for h in HEADERS}


def _make_xlsx(path: Path, rows: list[dict]) -> None:
    """按 HEADERS 造 xlsx；rows 为 {列名: 值}，缺列留空。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    wb.save(path)


def _make_copy(batch_id: str, rows: list[dict],
               stem: str | None = None) -> Path:
    """在批次 containers/ 下造 *_filled.xlsx 副本（快照发现口径同 XD 透视）。

    stem 缺省为 ContentsOfTheContainer_{batch_id}——sync 用例的原件必须是
    同 stem 无 _filled 后缀的 xlsx（_sync_output_copy_rows 按
    {src.stem}_filled{suffix} 定位副本）。
    """
    cdir = get_settings().batch_containers_dir(batch_id)
    cdir.mkdir(parents=True, exist_ok=True)
    stem = stem or f"ContentsOfTheContainer_{batch_id}"
    path = cdir / f"{stem}_filled.xlsx"
    _make_xlsx(path, rows)
    return path


def _ws(path: Path):
    return load_workbook(path).active


def _row(factory=FA, sku=SKU1, qty=10, name_cn="钳子", net=15.0, gross=25.0,
         po="272752", kanri="XC268452"):
    """标准数据行（缺省为「系统已写入」状态：三列有值）。"""
    return {"PURCHASE_ORDER": po, "KANRI_NO": kanri, "MAKER_MEI_KJ": factory,
            "SHOHIN_CD": sku, "SHOHIN_MEI_E": "PLIERS",
            "SOTOBAKO_D_HACCHU_SU": qty, "中文品名": name_cn,
            "净重": net, "毛重": gross}


def _write_state(batch_id: str, factory: str = FA, sku: str = SKU1,
                 unit_net=1.5, unit_gross=2.5, name_cn="钳子",
                 rows_need: int = 1) -> dict:
    """Node6 写入 state：item 单重 × 副本当前数量 = 系统写入值
    （缺省 qty=10 → 净重 15.0 / 毛重 25.0）。"""
    return {
        "batch_id": batch_id,
        "current_factory_data": {
            "factory_name": factory,
            "calculated_items": [
                {"sku": sku, "name_cn": name_cn,
                 "calculation": {"calculated_unit_net": unit_net,
                                 "calculated_unit_gross": unit_gross}},
            ],
        },
        # row_map 只承载「行数需求」（len），行号值不消费
        "downstream_row_map": {factory: {sku: [2] * rows_need}},
    }


def _cell(path: Path, row: int, field: str):
    return _ws(path).cell(row=row, column=COL[field]).value


def _set_cell(path: Path, row: int, field: str, value) -> None:
    """模拟人工改副本单元格。"""
    wb = load_workbook(path)
    wb.active.cell(row=row, column=COL[field]).value = value
    wb.save(path)


def _snapshot_row(batch_id: str, factory: str, sku: str, seq: int) -> dict:
    snap = copy_snapshot.load_snapshot(batch_id)
    key = copy_snapshot.row_key_str((factory, sku, seq))
    for r in snap["rows"]:
        if copy_snapshot.row_key_str(
                (r["key"]["factory"], r["key"]["sku"], r["key"]["seq"])) == key:
            return r
    raise AssertionError(f"快照中无行 {key}")


# ---------------------------------------------------------------------------
# 1. 改空覆盖：快照非空 × 副本空 = 等重新录入，通道 C 一律正常覆盖
# ---------------------------------------------------------------------------

def test_cleared_cell_overwritten_channel_c():
    batch = "T-WDC-CLEAR-WRITE"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)   # 快照：净重 15.0

    _set_cell(copy_path, 2, "净重", None)   # 人工改空

    state = _write_state(batch)
    written, _ = writer._write_excel(state, copy_path, channel="first_write")
    assert written == 1
    assert _cell(copy_path, 2, "净重") == 15.0      # 系统值覆盖
    cur = state["current_factory_data"]
    assert cur["manual_override_kept"] == 0         # 改空不算冲突
    assert cur["manual_conflicts"] == []
    # 快照重建后无保护标记
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is False


# ---------------------------------------------------------------------------
# 2. 通道 C 改非空：跳过保留 + 审计 manual_override_kept + 摘要 + protected
# ---------------------------------------------------------------------------

def test_channel_c_keep_audit_summary(caplog):
    batch = "T-WDC-CHC-KEEP"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)   # 快照：净重 15.0

    _set_cell(copy_path, 2, "净重", 99.0)   # 人工改非空 W

    state = _write_state(batch)
    with caplog.at_level(logging.WARNING, logger="app.nodes.writer"):
        written, _ = writer._write_excel(state, copy_path, channel="first_write")

    # 副本保留人工值；留痕字段齐备（行内其余列无冲突幂等写入，行计数为 1）
    assert written == 1
    assert _cell(copy_path, 2, "净重") == 99.0
    cur = state["current_factory_data"]
    assert cur["manual_override_kept"] == 1
    assert cur["manual_conflicts"] == [
        {"factory": FA, "sku": SKU1, "row": 2,
         "field": "net", "old": 15.0, "new": 99.0}]
    assert any("保留人工值未覆盖" in r.message for r in caplog.records)

    # 快照打 protected（记录值结转更新为人工值 99.0）
    row = _snapshot_row(batch, FA, SKU1, 1)
    assert row["protected"] is True
    assert row["protected_fields"] == ["net"]
    assert row["net"] == 99.0

    # 审计留痕：service._manual_write_traces → changes_json
    # （source="manual_override_kept"）经 _write_audit 落库
    values = {"factory_outputs": {FA: dict(cur)}}
    kept, overwritten, qty_changes, kept_rows = \
        service._manual_write_traces(values, None)
    assert len(kept) == 1 and kept[0]["source"] == "manual_override_kept"
    assert overwritten == [] and qty_changes == [] and kept_rows == []
    service._write_audit({
        "thread_id": batch, "factory_name": FA, "approved": True,
        "edited_count": 1, "changes": kept, "new_skus": [],
    }, "approved")
    with get_session() as session:
        audits = [a for a in session.query(ReviewAudit)
                  .filter(ReviewAudit.thread_id == batch).all()]
    assert len(audits) == 1
    changes = json.loads(audits[0].changes_json)
    assert changes[0]["source"] == "manual_override_kept"
    assert changes[0]["field"] == "net" and changes[0]["new"] == 99.0

    # 完成摘要透出：export_summary.json batch_summary 三字段
    export_node({
        "batch_id": batch,
        "downstream_file_path": "x.xlsx",
        "downstream_requirements": {FA: [SKU1]},
        "factory_outputs": {FA: dict(cur)},
        "final_output_path": str(copy_path),
    })
    summary = json.loads(
        (get_settings().batch_output_dir(batch) / "export_summary.json")
        .read_text(encoding="utf-8"))
    assert summary["manual_overrides_kept"] == 1
    assert summary["manual_values_overwritten"] == 0
    assert summary["qty_manual_changes"] == 0


# ---------------------------------------------------------------------------
# 3. 快照空 × 副本非空（人工抢先填）：通道 C 跳过保留 / 通道 A 覆写留痕
# ---------------------------------------------------------------------------

def test_snapshot_blank_copy_filled_channel_c_keeps():
    batch = "T-WDC-BLANK-C"
    # 基线时三列未写（快照空）
    copy_path = _make_copy(batch, [_row(name_cn=None, net=None, gross=None)])
    copy_snapshot.refresh_snapshot(batch)

    _set_cell(copy_path, 2, "净重", 88.0)   # 人工抢先填

    state = _write_state(batch)
    written, _ = writer._write_excel(state, copy_path, channel="first_write")
    assert _cell(copy_path, 2, "净重") == 88.0     # 保留人工值
    cur = state["current_factory_data"]
    # name_cn/gross 快照空副本空 → 正常写入；net 冲突保留
    assert cur["manual_override_kept"] == 1
    assert cur["manual_conflicts"][0]["field"] == "net"
    assert cur["manual_conflicts"][0]["old"] is None
    assert cur["manual_conflicts"][0]["new"] == 88.0
    assert _cell(copy_path, 2, "中文品名") == "钳子"
    assert _cell(copy_path, 2, "毛重") == 25.0
    assert written == 1   # 行内其余列已写，行计数仍算写入
    assert _snapshot_row(batch, FA, SKU1, 1)["protected_fields"] == ["net"]


def test_snapshot_blank_copy_filled_channel_a_overwrites():
    batch = "T-WDC-BLANK-A"
    copy_path = _make_copy(batch, [_row(name_cn=None, net=None, gross=None)])
    copy_snapshot.refresh_snapshot(batch)

    _set_cell(copy_path, 2, "净重", 88.0)   # 人工抢先填

    state = _write_state(batch)
    written, _ = writer._write_excel(state, copy_path, channel="rework")
    assert written == 1
    assert _cell(copy_path, 2, "净重") == 15.0     # 通道 A 直接覆盖
    cur = state["current_factory_data"]
    assert cur["manual_override_kept"] == 0
    assert cur["manual_values_overwritten"] == 1
    assert cur["manual_overwritten_cells"] == [
        {"factory": FA, "sku": SKU1, "row": 2,
         "field": "net", "old": 88.0, "new": 15.0}]
    # 无保护可结存（通道 A 不加保护）
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is False


# ---------------------------------------------------------------------------
# 4. MX1 回归：人工删行进墓碑，_sync_output_copy_rows 绝不复活
# ---------------------------------------------------------------------------

def test_sync_never_resurrects_manually_deleted_rows():
    batch = "T-WDC-MX1"
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    src_rows = [
        _row(),
        _row(factory=FMX, sku=SKU2, qty=7, po="272800", kanri="XC266877"),
    ]
    _make_xlsx(src, src_rows)
    copy_path = _make_copy(batch, src_rows, stem=stem)
    copy_snapshot.refresh_snapshot(batch)   # 快照含 FMX 行

    # 人工删掉副本中 FMX 全部行（MX1 事故现场）
    wb = load_workbook(copy_path)
    wb.active.delete_rows(3, 1)
    wb.save(copy_path)

    row_map = {FA: {SKU1: [2]}, FMX: {SKU2: [3]}}
    result = service._sync_output_copy_rows(batch, str(src), {FMX}, row_map)

    # 不复活、无「错位」报错（旧分支已删）、返回计数正确
    assert result == {"synced": 0, "skipped_manual_deleted": 1}
    assert _ws(copy_path).max_row == 2      # 仍只有 FA 行

    # 墓碑含 FMX 行键；快照行集不再含它
    snap = copy_snapshot.load_snapshot(batch)
    fmx_key = copy_snapshot.row_key_str((FMX, SKU2, 1))
    assert fmx_key in snap["tombstones"]
    assert all(r["key"]["factory"] != FMX for r in snap["rows"])

    # 再跑一次 sync：墓碑口径同样不复活（双口径印证）
    result2 = service._sync_output_copy_rows(batch, str(src), {FMX}, row_map)
    assert result2 == {"synced": 0, "skipped_manual_deleted": 1}
    assert _ws(copy_path).max_row == 2


# ---------------------------------------------------------------------------
# 5. 原件尾部真新增：追加副本尾部、三列留空、快照重建含新行
# ---------------------------------------------------------------------------

def test_sync_appends_genuinely_new_source_row():
    batch = "T-WDC-APPEND"
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    _make_xlsx(src, [_row()])
    copy_path = _make_copy(batch, [_row()], stem=stem)
    copy_snapshot.refresh_snapshot(batch)

    # 原件尾部追加 FB 新行（副本没有、快照没有、墓碑没有 → 真新增）
    wb = load_workbook(src)
    wb.active.append(["272900", "XC268453", FB, SKU2, "WRENCH", 5,
                      None, None, None])
    wb.save(src)

    result = service._sync_output_copy_rows(
        batch, str(src), {FB}, {FB: {SKU2: [3]}})
    assert result == {"synced": 1, "skipped_manual_deleted": 0}

    ws = _ws(copy_path)
    assert ws.max_row == 3                  # 追加在副本尾部
    assert ws.cell(row=3, column=COL["MAKER_MEI_KJ"]).value == FB
    assert ws.cell(row=3, column=COL["SHOHIN_CD"]).value == SKU2
    assert ws.cell(row=3, column=COL["KANRI_NO"]).value == "XC268453"
    assert ws.cell(row=3, column=COL["SOTOBAKO_D_HACCHU_SU"]).value == 5
    # 副本多出的三列留空待 writer 填
    for f in ("中文品名", "净重", "毛重"):
        assert ws.cell(row=3, column=COL[f]).value is None

    # 快照重建含新行
    new_row = _snapshot_row(batch, FB, SKU2, 1)
    assert new_row["row_no"] == 3
    assert new_row["net"] is None


# ---------------------------------------------------------------------------
# 6. clear_sku_rows 通道 A：人工非空值随 reopen 删除清空 + 留痕
# ---------------------------------------------------------------------------

def test_clear_sku_rows_channel_a_clears_manual_value():
    batch = "T-WDC-CLEAR-A"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)   # 快照：净重 15.0

    _set_cell(copy_path, 2, "净重", 77.0)   # 人工改非空

    state = _write_state(batch)
    cleared = writer.clear_sku_rows(state, copy_path, [SKU1], channel="rework")
    assert cleared == 1
    for f in ("中文品名", "净重", "毛重"):
        assert _cell(copy_path, 2, f) is None
    cur = state["current_factory_data"]
    # 只有人工改过的净重留痕（old=人工值，new=None）
    assert cur["manual_values_overwritten"] == 1
    assert cur["manual_overwritten_cells"] == [
        {"factory": FA, "sku": SKU1, "row": 2,
         "field": "net", "old": 77.0, "new": None}]


# ---------------------------------------------------------------------------
# 7. 补充工厂确认门（通道 B）：未确认 409 结构体未入图；确认后放行覆写
# ---------------------------------------------------------------------------

class _FakeSnap:
    def __init__(self, values=None, next_nodes=(), tasks=()):
        self.values = values if values is not None else {}
        self.next = tuple(next_nodes)
        self.tasks = tuple(tasks)


class _FakeGraph:
    """记录 update_state / stream 调用次数的假 graph（照 test_stream_fallback）。"""

    def __init__(self, snap: _FakeSnap):
        self._snap = snap
        self.updates: list[dict] = []
        self.stream_calls = 0

    def get_state(self, _cfg):
        return self._snap

    def update_state(self, _cfg, update, **_kw):
        self.updates.append(update)
        self._snap.values.update(update)

    def stream(self, *_a, **_kw):
        self.stream_calls += 1
        return iter([])     # 无事件 → completed 兜底分支


def _channel_b_setup(batch: str):
    """造通道 B 现场：FB 为待补充工厂，副本 FB 行净重被人工改非空。"""
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    src_rows = [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, name_cn=None, net=None, gross=None,
             po="272900", kanri="XC268453"),
    ]
    _make_xlsx(src, src_rows)
    copy_path = _make_copy(batch, src_rows, stem=stem)
    copy_snapshot.refresh_snapshot(batch)   # 基线：FB 行三列空
    _set_cell(copy_path, 3, "净重", 66.0)   # 人工抢先填 FB 行净重
    batch_store.upsert_batch(batch, status="completed")
    service._write_batch_config(batch, {
        "thread_id": batch,
        "downstream_file_path": str(src),
        "upstream_root": str(TMP / "upstream"),
    })
    return src, copy_path


def _channel_b_graph(batch: str) -> _FakeGraph:
    return _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {},
        "pending_factories": [],
        "downstream_requirements": {},
        "downstream_row_map": {},
    }))


def test_add_factories_conflict_gate(monkeypatch):
    batch = "T-WDC-GATE"
    _channel_b_setup(batch)
    fake = _channel_b_graph(batch)
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(
        "app.nodes.parse_downstream.parse_downstream_file",
        lambda _p: ({FB: {"total_qty": 5}}, {FB: {SKU2: [3]}}))
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)

    # 未确认：ManualConflictError，未入图（update_state/stream 均未发生）
    with pytest.raises(service.ManualConflictError) as ei:
        service.add_factories_to_batch(batch)
    assert ei.value.conflicts == [
        {"factory": FB, "sku": SKU2, "row": 3,
         "field": "net", "old": None, "new": 66.0}]
    assert fake.updates == [] and fake.stream_calls == 0

    # 路由层契约：409 + 结构体 detail（前端确认门弹窗数据源）
    resp = client.post(f"/api/v1/batches/{batch}/add-factories", json={})
    assert resp.status_code == 409
    detail = resp.json()["detail"]
    assert detail["code"] == "manual_conflicts"
    assert detail["count"] == 1
    assert detail["message"]
    assert detail["conflicts"][0]["field"] == "net"
    assert detail["conflicts"][0]["new"] == 66.0
    assert fake.stream_calls == 0          # 路由层同样未入图

    # 确认后：放行续跑，确认工厂写入 overwrite_manual_factories（通道 A 覆写）
    result = service.add_factories_to_batch(batch, confirm_overwrite=True)
    assert fake.stream_calls == 1
    assert fake.updates and \
        fake.updates[0]["overwrite_manual_factories"] == [FB]
    assert result["status"] == "completed"
    assert result["added"] == 1 and result["factories"] == [FB]
    assert result["manual_overrides_kept"] == 0     # 字段恒在（前端契约）


# ---------------------------------------------------------------------------
# 8. 快照结转（遗忘漏洞回归）：protected 跨重建不丢；人工清空解除保护
# ---------------------------------------------------------------------------

def test_snapshot_carry_over_protected_and_unprotect_on_clear():
    batch = "T-WDC-CARRY"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)

    _set_cell(copy_path, 2, "净重", 99.0)   # 人工改非空
    state = _write_state(batch)
    writer._write_excel(state, copy_path, channel="first_write")
    assert _snapshot_row(batch, FA, SKU1, 1)["protected_fields"] == ["net"]

    # 二次写入（快照重建后）：protected 结转生效，仍不覆盖人工值
    state2 = _write_state(batch)
    written, _ = writer._write_excel(state2, copy_path, channel="first_write")
    assert _cell(copy_path, 2, "净重") == 99.0
    assert state2["current_factory_data"]["manual_override_kept"] == 1
    # 中文品名/毛重无冲突正常写入
    assert written == 1
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is True

    # 人工清空 protected 单元格 → 重建解除该列保护 → 下次写入覆盖
    _set_cell(copy_path, 2, "净重", None)
    copy_snapshot.refresh_snapshot(batch)
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is False
    state3 = _write_state(batch)
    writer._write_excel(state3, copy_path, channel="first_write")
    assert _cell(copy_path, 2, "净重") == 15.0
    assert state3["current_factory_data"]["manual_override_kept"] == 0


# ---------------------------------------------------------------------------
# 9. 数量列告警：人工改数量 → 按新数量写入 + warning + 摘要透出
# ---------------------------------------------------------------------------

def test_qty_change_warns_and_writes_by_new_qty(caplog):
    batch = "T-WDC-QTY"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)   # 快照：qty 10

    _set_cell(copy_path, 2, "SOTOBAKO_D_HACCHU_SU", 20)   # 人工改数量

    state = _write_state(batch)
    with caplog.at_level(logging.WARNING, logger="app.nodes.writer"):
        written, _ = writer._write_excel(state, copy_path, channel="first_write")

    # 写入仍按副本当前数量吸收（1.5×20 / 2.5×20）
    assert written == 1
    assert _cell(copy_path, 2, "净重") == 30.0
    assert _cell(copy_path, 2, "毛重") == 50.0
    cur = state["current_factory_data"]
    assert cur["qty_manual_changes"] == [
        {"factory": FA, "sku": SKU1, "row": 2, "old": 10.0, "new": 20.0}]
    assert any("数量被人工改过" in r.message and "已按新数量写入" in r.message
               for r in caplog.records)

    # 摘要透出：traces 带 source + export_summary 计数字段
    _k, _o, qty_changes, _r = service._manual_write_traces(
        {"factory_outputs": {FA: dict(cur)}}, None)
    assert len(qty_changes) == 1
    assert qty_changes[0]["source"] == "manual_qty_changed"
    export_node({
        "batch_id": batch,
        "downstream_file_path": "x.xlsx",
        "downstream_requirements": {FA: [SKU1]},
        "factory_outputs": {FA: dict(cur)},
        "final_output_path": str(copy_path),
    })
    summary = json.loads(
        (get_settings().batch_output_dir(batch) / "export_summary.json")
        .read_text(encoding="utf-8"))
    assert summary["qty_manual_changes"] == 1


# ---------------------------------------------------------------------------
# 10. overwrite-protected 端点：通道 A 语义补写 + 审计 + 保护解除；400 分支
# ---------------------------------------------------------------------------

def test_overwrite_protected_endpoint(monkeypatch):
    batch = "T-WDC-OWP"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)

    _set_cell(copy_path, 2, "净重", 99.0)   # 人工改非空
    state = _write_state(batch)
    writer._write_excel(state, copy_path, channel="first_write")
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is True

    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {FA: {
            "calculated_items": [
                {"sku": SKU1, "name_cn": "钳子",
                 "calculation": {"calculated_unit_net": 1.5,
                                 "calculated_unit_gross": 2.5}}],
        }},
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)

    # service 层：按通道 A 语义补写系统值（qty=10 → 15.0）
    result = service.overwrite_protected_rows(batch)
    assert result == {"ok": True, "overwritten": 1}
    assert _cell(copy_path, 2, "净重") == 15.0
    # 保护解除（记录值=副本当前值=系统值）
    row = _snapshot_row(batch, FA, SKU1, 1)
    assert row["protected"] is False and row["protected_fields"] == []
    # 审计：被盖人工旧值留痕 manual_value_overwritten
    with get_session() as session:
        audits = [a for a in session.query(ReviewAudit)
                  .filter(ReviewAudit.thread_id == batch,
                          ReviewAudit.result_status == "overwrite_protected")
                  .all()]
    assert len(audits) == 1
    changes = json.loads(audits[0].changes_json)
    assert changes == [{"sku": SKU1, "field": "net", "old": 99.0, "new": 15.0,
                        "source": "manual_value_overwritten"}]

    # 路由层契约：200 {"ok": true, "overwritten": N}
    _set_cell(copy_path, 2, "毛重", 88.0)   # 再造一处人工保护
    state_b = _write_state(batch)
    writer._write_excel(state_b, copy_path, channel="first_write")
    resp = client.post(f"/api/v1/batches/{batch}/overwrite-protected")
    assert resp.status_code == 200
    assert resp.json() == {"ok": True, "overwritten": 1}
    assert _cell(copy_path, 2, "毛重") == 25.0

    # 无 protected 行 → 400
    resp2 = client.post(f"/api/v1/batches/{batch}/overwrite-protected")
    assert resp2.status_code == 400


def test_overwrite_protected_no_protected_rows_400(monkeypatch):
    batch = "T-WDC-OWP-NONE"
    _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)
    fake = _FakeGraph(_FakeSnap(values={"batch_id": batch,
                                        "factory_outputs": {}}))
    monkeypatch.setattr(service, "get_graph", lambda: fake)

    with pytest.raises(ValueError):
        service.overwrite_protected_rows(batch)
    resp = client.post(f"/api/v1/batches/{batch}/overwrite-protected")
    assert resp.status_code == 400


if __name__ == "__main__":
    import pytest as _pytest
    _pytest.main([__file__, "-v"])
