# -*- coding: utf-8 -*-
"""副本快照模块 + 行标识重定位测试（写入侧人工改动感知 · 阶段一 §1.4）。

覆盖：
1. 快照往返：同厂同 SKU 两行 seq 正确、空 KANRI_NO 行在快照内、
   工厂/SKU 空行不进索引不进快照、三列/qty 两位小数、
   qty/protected/protected_fields/tombstones 字段存在且默认（version=2）；
2. 行标识重定位（决策 3）：目标行前人工 insert_rows 插 2 行后
   _write_excel 仍写对行、clear_sku_rows 仍清对行（row_map 行号值不消费）；
3. 无快照迁移（决策 7）：副本存在无快照 → refresh 后 origin=="baseline"
   + warning 日志；
4. load_snapshot 容错：JSON 损坏 / 版本不符 → None + warning；
   无副本 refresh → None 不落盘。

运行：
    cd app && PYTHONPATH=. python3 -m pytest tests/test_copy_snapshot.py -v

隔离红线（2026-08-11 事故教训）：
- import app 模块前先设 YAMATO_TEST_MODE=1 + YAMATO_DOTENV_PATH=临时空 .env
  （llm_client 模块级 load_dotenv(override=True) 只可能读到空文件）；
- import 后再调 validation/_test_isolation.isolate_to_tmp 把
  checkpoint/master/output/sessions 全部指向临时目录（守卫断言不碰生产库）。
"""
from __future__ import annotations

import json
import logging
import os
import sys
import tempfile
from pathlib import Path

import pytest

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

# ---- 隔离门①：import 任何 app 模块前的双门（血泪红线）----
_TMP = Path(tempfile.mkdtemp(prefix="yamato_copy_snap_test_"))
(_TMP / ".env").write_text("# isolated .env\n", encoding="utf-8")
os.environ["YAMATO_TEST_MODE"] = "1"
os.environ["YAMATO_DOTENV_PATH"] = str(_TMP / ".env")
os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

from openpyxl import Workbook, load_workbook  # noqa: E402

from app.config import get_settings  # noqa: E402
from app.nodes import writer  # noqa: E402
from app.orchestrator import copy_snapshot  # noqa: E402

# ---- 隔离门②：import 后把 db/output/sessions 指向临时目录 ----
from _test_isolation import isolate_to_tmp  # noqa: E402

TMP = isolate_to_tmp("yamato_copy_snap_test_")


@pytest.fixture(autouse=True)
def _settings_pin():
    """settings 单例钉根：全量混跑时其他测试文件的 isolate_to_tmp 会把
    output_dir 重钉到它们自己的临时目录——每个用例前把 output_dir 钉回
    本模块的 TMP，结束后恢复原值（照 test_factory_watch 先例）。"""
    s = get_settings()
    original = s.output_dir
    s.output_dir = str(TMP / "output")
    try:
        yield
    finally:
        s.output_dir = original


FACTORY_A = "山東中地"
FACTORY_B = "Ｃ．正達工芸品"
SKU1 = "4550596211831"
SKU2 = "4900000009201"

HEADERS = ["PURCHASE_ORDER", "KANRI_NO", "MAKER_MEI_KJ", "SHOHIN_CD",
           "SHOHIN_MEI_E", "SOTOBAKO_D_HACCHU_SU", "中文品名", "净重", "毛重"]
COL = {h: HEADERS.index(h) + 1 for h in HEADERS}


def _make_xlsx(path: Path, rows: list[dict]) -> None:
    """按 HEADERS 造副本 xlsx；rows 为 {列名: 值}，缺列留空。"""
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    wb.save(path)


def _make_copy(batch_id: str, rows: list[dict]) -> Path:
    """在批次 containers/ 下造 *_filled.xlsx 副本（快照发现口径同 XD 透视）。"""
    cdir = get_settings().batch_containers_dir(batch_id)
    cdir.mkdir(parents=True, exist_ok=True)
    path = cdir / f"ContentsOfTheContainer_{batch_id}_filled.xlsx"
    _make_xlsx(path, rows)
    return path


def _ws(path: Path):
    return load_workbook(path).active


# ---------------------------------------------------------------------------
# 1. 快照往返（build → save → load 字段级断言）
# ---------------------------------------------------------------------------

def test_build_row_index_seq_and_blank_rows():
    """row_key 索引：同 (工厂, SKU) 组内按行号升序编 seq；
    工厂/SKU 任一空值的行不进索引。"""
    path = TMP / "row_index.xlsx"
    _make_xlsx(path, [
        {"MAKER_MEI_KJ": FACTORY_A, "SHOHIN_CD": SKU1, "KANRI_NO": "XC268452"},
        {"MAKER_MEI_KJ": FACTORY_A, "SHOHIN_CD": SKU1},   # 同厂同 SKU 第二行
        {"MAKER_MEI_KJ": "", "SHOHIN_CD": SKU2},          # 工厂空 → 不进索引
        {"MAKER_MEI_KJ": FACTORY_B, "SHOHIN_CD": None},   # SKU 空 → 不进索引
        {"MAKER_MEI_KJ": "  ", "SHOHIN_CD": SKU1},        # 空白工厂 → 不进索引
        {"MAKER_MEI_KJ": FACTORY_B, "SHOHIN_CD": SKU2},
    ])
    idx = copy_snapshot.build_row_index(_ws(path))
    assert idx == {
        (FACTORY_A, SKU1, 1): 2,
        (FACTORY_A, SKU1, 2): 3,
        (FACTORY_B, SKU2, 1): 7,
    }


def test_snapshot_roundtrip():
    """build/save/load 往返：seq、空 KANRI_NO 行、空键行排除、数值规整、
    version=2 默认字段（qty/protected/protected_fields/tombstones）。"""
    batch = "T-SNAP-RT"
    copy_path = _make_copy(batch, [
        {"PURCHASE_ORDER": "272752", "KANRI_NO": "XC268452",
         "MAKER_MEI_KJ": FACTORY_A, "SHOHIN_CD": SKU1,
         "SOTOBAKO_D_HACCHU_SU": 120,
         "中文品名": "钳子", "净重": 12.3, "毛重": 13.567},
        # 同厂同 SKU 第二行 + 空 KANRI_NO（仍在快照内，kanri_no=None）
        {"PURCHASE_ORDER": "272753", "MAKER_MEI_KJ": FACTORY_A,
         "SHOHIN_CD": SKU1, "SOTOBAKO_D_HACCHU_SU": 63},
        # 工厂空行 / SKU 空行：不进快照
        {"MAKER_MEI_KJ": None, "SHOHIN_CD": SKU2, "KANRI_NO": "XC000000"},
        {"MAKER_MEI_KJ": FACTORY_B, "SHOHIN_CD": None},
        # MX 人工行
        {"PURCHASE_ORDER": "MX2-268510-001", "KANRI_NO": "XC266877",
         "MAKER_MEI_KJ": FACTORY_B, "SHOHIN_CD": SKU2,
         "SOTOBAKO_D_HACCHU_SU": 5,
         "中文品名": "人工品名", "净重": 12.34, "毛重": 23.45},
    ])

    snap = copy_snapshot.build_snapshot(batch)
    assert snap is not None
    assert snap["version"] == 2
    assert snap["batch_id"] == batch
    assert snap["copy_file"] == copy_path.name
    assert snap["tombstones"] == []          # 阶段一恒为空名单
    assert snap["data_row_count"] == 5       # 数据行总数（含空键填充行）
    assert snap["copy_mtime_ns"] > 0 and snap["copy_size"] > 0

    rows = snap["rows"]
    assert len(rows) == 3                    # 空键两行被排除
    assert [r["key"] for r in rows] == [
        {"factory": FACTORY_A, "sku": SKU1, "seq": 1},
        {"factory": FACTORY_A, "sku": SKU1, "seq": 2},
        {"factory": FACTORY_B, "sku": SKU2, "seq": 1},
    ]
    r0, r1, r2 = rows
    assert r0["row_no"] == 2 and r1["row_no"] == 3 and r2["row_no"] == 6
    # 数值规整：三列/qty 统一两位小数
    assert r0["net"] == 12.3 and r0["gross"] == 13.57 and r0["qty"] == 120.0
    assert r1["kanri_no"] is None and r1["net"] is None and r1["qty"] == 63.0
    assert r0["kanri_no"] == "XC268452" and r0["name_cn"] == "钳子"
    assert r2["is_mx"] is True and r0["is_mx"] is False
    for r in rows:
        assert r["origin"] == "system"
        assert r["protected"] is False       # 阶段一恒默认
        assert r["protected_fields"] == []

    copy_snapshot.save_snapshot(batch, snap)
    assert copy_snapshot.snapshot_path(batch).exists()
    loaded = copy_snapshot.load_snapshot(batch)
    assert loaded == snap


def test_row_key_str_literal_sep():
    """JSON 字符串键：字面反斜杠零两个字符（非 NUL 控制字符）。"""
    assert len(copy_snapshot.KEY_SEP) == 2
    s = copy_snapshot.row_key_str((FACTORY_A, SKU1, 2))
    assert s == f"{FACTORY_A}\\0{SKU1}\\02"
    assert "\x00" not in s
    assert copy_snapshot.parse_row_key(s) == (FACTORY_A, SKU1, 2)


# ---------------------------------------------------------------------------
# 2. 行标识重定位：人工插行后写对/清对（row_map 行号值不消费）
# ---------------------------------------------------------------------------

def _write_state() -> dict:
    return {
        "current_factory_data": {
            "factory_name": FACTORY_A,
            "calculated_items": [
                {"sku": SKU1, "name_cn": "钳子",
                 "calculation": {"calculated_unit_net": 1.5,
                                 "calculated_unit_gross": 2.5}},
            ],
        },
        # row_map 只提供「该 SKU 行数需求」（len=1）；行号值 2 是插行前的
        # 旧值，重定位后不再被消费（人工插行后实际行号已移位到 4）
        "downstream_row_map": {FACTORY_A: {SKU1: [2]}},
    }


def test_write_excel_relocates_after_manual_insert():
    """目标行前人工插入 2 行：_write_excel 按 (工厂, SKU, seq) 现扫定位，
    写在移位后的真实行（row4），插入的空白行不被误写。"""
    path = TMP / "reloc_write.xlsx"
    _make_xlsx(path, [
        {"PURCHASE_ORDER": "272752", "KANRI_NO": "XC268452",
         "MAKER_MEI_KJ": FACTORY_A, "SHOHIN_CD": SKU1,
         "SOTOBAKO_D_HACCHU_SU": 10},
    ])
    wb = load_workbook(path)
    wb.active.insert_rows(2, 2)   # 目标行 row2 → row4
    wb.save(path)

    written, skipped_placeholder = writer._write_excel(_write_state(), path)
    assert (written, skipped_placeholder) == (1, 0)

    ws = _ws(path)
    # 净重 = 1.5 × 10，毛重 = 2.5 × 10
    assert ws.cell(row=4, column=COL["中文品名"]).value == "钳子"
    assert ws.cell(row=4, column=COL["净重"]).value == 15.0
    assert ws.cell(row=4, column=COL["毛重"]).value == 25.0
    for r in (2, 3):  # 人工插入的空白行未被误写
        assert ws.cell(row=r, column=COL["中文品名"]).value is None
        assert ws.cell(row=r, column=COL["净重"]).value is None


def test_clear_sku_rows_relocates_after_manual_insert():
    """同口径：插行后 clear_sku_rows 清移位后的真实行，其他 SKU 行不动。"""
    path = TMP / "reloc_clear.xlsx"
    _make_xlsx(path, [
        {"PURCHASE_ORDER": "272752", "KANRI_NO": "XC268452",
         "MAKER_MEI_KJ": FACTORY_A, "SHOHIN_CD": SKU1,
         "SOTOBAKO_D_HACCHU_SU": 10, "中文品名": "钳子",
         "净重": 15.0, "毛重": 25.0},
        {"PURCHASE_ORDER": "272753", "KANRI_NO": "XC266877",
         "MAKER_MEI_KJ": FACTORY_A, "SHOHIN_CD": SKU2,
         "SOTOBAKO_D_HACCHU_SU": 5, "中文品名": "扳手",
         "净重": 50.0, "毛重": 55.0},
    ])
    wb = load_workbook(path)
    wb.active.insert_rows(2, 2)   # SKU1 行 row2 → row4；SKU2 行 row3 → row5
    wb.save(path)

    state = {
        "current_factory_data": {"factory_name": FACTORY_A},
        "downstream_row_map": {FACTORY_A: {SKU1: [2], SKU2: [3]}},
    }
    cleared = writer.clear_sku_rows(state, path, [SKU1])
    assert cleared == 1

    ws = _ws(path)
    # 被清 SKU 行（移位后 row4）：三列全空，其余列不动
    assert [ws.cell(row=4, column=c).value
            for c in (COL["中文品名"], COL["净重"], COL["毛重"])] == [None] * 3
    assert ws.cell(row=4, column=COL["SHOHIN_CD"]).value == SKU1
    # 未删 SKU 行（移位后 row5）不受影响
    assert [ws.cell(row=5, column=c).value
            for c in (COL["中文品名"], COL["净重"], COL["毛重"])] == ["扳手", 50.0, 55.0]
    for r in (2, 3):  # 插入的空白行未被误清出脏值
        assert ws.cell(row=r, column=COL["净重"]).value is None


def test_write_excel_missing_row_warns_and_skips(caplog):
    """row_key 在副本中查不到（人工删行）：记 warning 跳过，不报错不复活
    （阶段一不加冲突语义）。"""
    path = TMP / "reloc_missing.xlsx"
    _make_xlsx(path, [
        {"PURCHASE_ORDER": "272752", "KANRI_NO": "XC268452",
         "MAKER_MEI_KJ": FACTORY_A, "SHOHIN_CD": SKU2,   # 只剩别的 SKU
         "SOTOBAKO_D_HACCHU_SU": 10},
    ])
    with caplog.at_level(logging.WARNING, logger="app.nodes.writer"):
        written, _ = writer._write_excel(_write_state(), path)
    assert written == 0
    assert any("未找到" in r.message for r in caplog.records)


# ---------------------------------------------------------------------------
# 3. 无快照迁移（决策 7）：副本存在无快照 → baseline + warning
# ---------------------------------------------------------------------------

def test_refresh_baseline_when_no_snapshot(caplog):
    batch = "T-SNAP-BASELINE"
    _make_copy(batch, [
        {"MAKER_MEI_KJ": FACTORY_A, "SHOHIN_CD": SKU1,
         "SOTOBAKO_D_HACCHU_SU": 10, "净重": 15.0},
    ])
    assert not copy_snapshot.snapshot_path(batch).exists()

    with caplog.at_level(logging.WARNING,
                         logger="app.orchestrator.copy_snapshot"):
        snap = copy_snapshot.refresh_snapshot(batch)
    assert snap is not None
    assert all(r["origin"] == "baseline" for r in snap["rows"])
    assert any("基线" in r.message and "不可追溯" in r.message
               for r in caplog.records if r.levelno >= logging.WARNING)
    assert copy_snapshot.snapshot_path(batch).exists()

    # 第二次 refresh：快照已存在 → 常规重建 origin 回 system
    snap2 = copy_snapshot.refresh_snapshot(batch)
    assert all(r["origin"] == "system" for r in snap2["rows"])


# ---------------------------------------------------------------------------
# 4. 容错：load 损坏/版本不符 → None + warning；无副本 refresh → None
# ---------------------------------------------------------------------------

def test_load_snapshot_corrupt_and_version_mismatch(caplog):
    batch = "T-SNAP-BAD"
    p = copy_snapshot.snapshot_path(batch)
    p.parent.mkdir(parents=True, exist_ok=True)

    with caplog.at_level(logging.WARNING,
                         logger="app.orchestrator.copy_snapshot"):
        p.write_text("{not json", encoding="utf-8")
        assert copy_snapshot.load_snapshot(batch) is None
        p.write_text(json.dumps({"version": 1, "rows": []}), encoding="utf-8")
        assert copy_snapshot.load_snapshot(batch) is None
    assert sum("快照" in r.message for r in caplog.records
               if r.levelno >= logging.WARNING) == 2

    # 文件不存在：静默 None（无 warning——无快照是常态）
    assert copy_snapshot.load_snapshot("T-SNAP-MISSING") is None


def test_refresh_no_copy_returns_none():
    """无副本：build_snapshot 返回 None，refresh 不落盘不报错。"""
    batch = "T-SNAP-NOCOPY"
    assert copy_snapshot.build_snapshot(batch) is None
    assert copy_snapshot.refresh_snapshot(batch) is None
    assert not copy_snapshot.snapshot_path(batch).exists()


def main():
    test_build_row_index_seq_and_blank_rows()
    test_snapshot_roundtrip()
    test_row_key_str_literal_sep()
    test_write_excel_relocates_after_manual_insert()
    test_clear_sku_rows_relocates_after_manual_insert()
    test_refresh_no_copy_returns_none()
    print("\ntest_copy_snapshot（caplog 用例需走 pytest）: PASS")


if __name__ == "__main__":
    main()
