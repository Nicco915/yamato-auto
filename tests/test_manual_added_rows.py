# -*- coding: utf-8 -*-
"""人工新增行纳入（写入侧人工改动感知 · 阶段五 §5.6，决策 4）测试。

覆盖（对应任务六项）：
1. A 态：原件已有厂 + 副本新增 SKU 行 → 写前 diff 三态分类 → 并入计划
   （requirements 追加 SKU + row_map 等长占位）→ writer 写入时提取结果
   含此 SKU → 按通道规则写入；写后快照该行 origin="manual_added"；
2. A 态无提取结果支撑：人工行原样保留 + warning + cur 留痕 →
   service._manual_write_traces 带 source="manual_row_kept_no_extract"；
3. B 态：副本新增全新厂行 → add_factories_to_batch（假 graph + 假解析）
   pending/requirements/row_map 合并入 state + sync 快照打 manual_added；
   （无文件夹自动跳转过 awaiting 属既有 Node4C 链路，e2e 归阶段六）
4. C 态：缺 KANRI_NO/工厂/SKU 行 → 不纳入 + scan_incomplete_rows 活扫
   透出 + 批次详情顶层 manual_rows_incomplete（前端契约形状）+
   补充工厂早退摘要提示；
5. loader 一致性：新增行缺 KANRI_NO → split loader 计数 warning +
   skipped_kanri_rows 行号收集；
6. 墓碑衔接与对称：加回已删行 → 出墓碑 + 走三态；剔除厂加回 → readded
   审计；纳入后人工再删 → 回删除分支进墓碑不复活。

运行：
    cd app && PYTHONPATH=. python3 -m pytest tests/test_manual_added_rows.py -v

隔离红线（2026-08-11 事故教训，骨架照抄 test_factory_removal.py）：
- import app 模块前先设 YAMATO_TEST_MODE=1 + YAMATO_DOTENV_PATH=临时空 .env；
- import 后调 validation/_test_isolation.isolate_to_tmp 把
  checkpoint/master/output/sessions 全部指向临时目录。
"""
from __future__ import annotations

import json
import logging
import os
import sys
import tempfile
from pathlib import Path

import pytest

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

# ---- 隔离门①：import 任何 app 模块前的双门（血泪红线）----
_TMP = Path(tempfile.mkdtemp(prefix="yamato_manual_added_rows_test_"))
(_TMP / ".env").write_text("# isolated .env\n", encoding="utf-8")
os.environ["YAMATO_TEST_MODE"] = "1"
os.environ["YAMATO_DOTENV_PATH"] = str(_TMP / ".env")
os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.api import service  # noqa: E402
from app.api.main import app  # noqa: E402,F401  确保全部 app 模块先于隔离 import
from app.config import get_settings  # noqa: E402
from app.db import batch_store  # noqa: E402
from app.db.models import ReviewAudit  # noqa: E402
from app.db.session import get_session  # noqa: E402
from app.nodes import writer  # noqa: E402
from app.orchestrator import copy_snapshot  # noqa: E402
from app.split.loader import load_filled_excel  # noqa: E402

# ---- 隔离门②：import 后把 db/output/sessions 指向临时目录 ----
from _test_isolation import isolate_to_tmp  # noqa: E402

TMP = isolate_to_tmp("yamato_manual_added_rows_test_")


@pytest.fixture(autouse=True)
def _settings_pin():
    """settings 单例钉根（照 test_write_diff_conflicts 先例）：全量混跑时
    其他测试文件的 isolate_to_tmp 会重钉 output_dir——每个用例前钉回本模块
    的 TMP，结束后恢复原值。"""
    s = get_settings()
    original = s.output_dir
    s.output_dir = str(TMP / "output")
    try:
        yield
    finally:
        s.output_dir = original


@pytest.fixture(autouse=True)
def _no_auto_backup(monkeypatch):
    """结案自动备份打空（照 test_awaiting_factories 先例，防全量混跑撞锁）。"""
    from app.orchestrator import backup
    monkeypatch.setattr(backup, "trigger_async", lambda reason, **kw: True)
    yield


@pytest.fixture(autouse=True)
def _cleanup_batches():
    """结束后删本模块造的批次行（共享临时库）。"""
    try:
        yield
    finally:
        for tid in _CREATED_TIDS:
            batch_store.delete_batch(tid)
        _CREATED_TIDS.clear()


_CREATED_TIDS: list[str] = []

FA = "山東中地"          # 原件已有工厂
FB = "青島新厂"          # 副本人工新增的全新工厂
SKU1 = "4550596211831"
SKU2 = "4900000009201"

HEADERS = ["PURCHASE_ORDER", "KANRI_NO", "MAKER_MEI_KJ", "SHOHIN_CD",
           "SHOHIN_MEI_E", "SOTOBAKO_D_HACCHU_SU", "中文品名", "净重", "毛重"]
COL = {h: HEADERS.index(h) + 1 for h in HEADERS}


def _make_xlsx(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    wb.save(path)


def _make_copy(batch_id: str, rows: list[dict],
               stem: str | None = None) -> Path:
    """在批次 containers/ 下造 *_filled.xlsx 副本（口径同
    test_write_diff_conflicts._make_copy）。"""
    cdir = get_settings().batch_containers_dir(batch_id)
    cdir.mkdir(parents=True, exist_ok=True)
    stem = stem or f"ContentsOfTheContainer_{batch_id}"
    path = cdir / f"{stem}_filled.xlsx"
    _make_xlsx(path, rows)
    return path


def _row(factory=FA, sku=SKU1, qty=10, name_cn="钳子", net=15.0, gross=25.0,
         po="272752", kanri="XC268452"):
    return {"PURCHASE_ORDER": po, "KANRI_NO": kanri, "MAKER_MEI_KJ": factory,
            "SHOHIN_CD": sku, "SHOHIN_MEI_E": "PLIERS",
            "SOTOBAKO_D_HACCHU_SU": qty, "中文品名": name_cn,
            "净重": net, "毛重": gross}


def _append_row(copy_path: Path, row: dict) -> None:
    """模拟人工在副本尾部新增一行。"""
    wb = load_workbook(copy_path)
    ws = wb.active
    ws.append([row.get(h) for h in HEADERS])
    wb.save(copy_path)


def _delete_row(path: Path, row_no: int) -> None:
    """模拟人工删副本整行。"""
    wb = load_workbook(path)
    wb.active.delete_rows(row_no, 1)
    wb.save(path)


def _cell(path: Path, row: int, field: str):
    wb = load_workbook(path)
    try:
        return wb.active.cell(row=row, column=COL[field]).value
    finally:
        wb.close()


def _snapshot_row(batch_id: str, factory: str, sku: str, seq: int) -> dict:
    snap = copy_snapshot.load_snapshot(batch_id)
    for r in snap["rows"]:
        k = r["key"]
        if (k["factory"], k["sku"], k["seq"]) == (factory, sku, seq):
            return r
    raise AssertionError(f"快照中无行 {(factory, sku, seq)}")


def _write_state(batch_id: str, items: list[dict],
                 requirements=None, row_map=None) -> dict:
    """Node6 写入 state（口径照 test_write_diff_conflicts._write_state，
    外加阶段五消费的 downstream_requirements）。"""
    return {
        "batch_id": batch_id,
        "current_factory_data": {
            "factory_name": FA,
            "calculated_items": items,
        },
        "downstream_requirements": requirements or {FA: [SKU1]},
        # row_map 只承载「行数需求」（len），行号值不消费
        "downstream_row_map": row_map or {FA: {SKU1: [2]}},
    }


def _item(sku: str, unit_net=1.5, unit_gross=2.5, name_cn="钳子") -> dict:
    return {"sku": sku, "name_cn": name_cn,
            "calculation": {"calculated_unit_net": unit_net,
                            "calculated_unit_gross": unit_gross}}


def _audits(batch: str) -> list[ReviewAudit]:
    with get_session() as session:
        return list(session.scalars(
            select(ReviewAudit)
            .where(ReviewAudit.thread_id == batch)
            .order_by(ReviewAudit.audit_id)
        ).all())


# ---------------------------------------------------------------------------
# 1. A 态：提取结果含新增 SKU → 并入需求 + 按通道规则写入 + 快照打标
# ---------------------------------------------------------------------------

def test_a_state_merge_and_write_with_extract():
    batch = "T-MA-A-WRITE"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)   # 快照只有 FA/SKU1

    # 人工新增：FA 的新 SKU 行（管理番号齐备 → A 态）
    _append_row(copy_path, _row(sku=SKU2, qty=6, name_cn=None, net=None,
                                gross=None, po="272753", kanri="XC268499"))

    state = _write_state(batch, [_item(SKU1), _item(SKU2, name_cn="剪刀")])
    written, _ = writer._write_excel(state, copy_path, channel="first_write")

    # 两行都写入（SKU1 原有行 + SKU2 人工新增行）
    assert written == 2
    assert _cell(copy_path, 3, "净重") == 9.0      # 1.5 × 副本数量 6
    assert _cell(copy_path, 3, "中文品名") == "剪刀"

    # 并入计划：A 态需求追加 + row_map 等长占位 + 打标键集
    cur = state["current_factory_data"]
    merge = cur["_manual_added_merge"]
    assert merge["req_updates"] == {FA: [SKU1, SKU2]}
    assert merge["row_map_updates"][FA][SKU2] == [3]   # 占位列表 len=行数需求
    assert merge["b_factories"] == []
    assert merge["manual_keys"] == {
        copy_snapshot.row_key_str((FA, SKU2, 1))}
    # 提取结果有支撑 → 不留「无提取」痕
    assert cur["manual_rows_kept_no_extract"] == []

    # 写后快照：新增行入库且 origin="manual_added"；原有行是建基线时入的
    # 库（baseline，来源不可考——首次 refresh 时副本已有人工/系统值）
    row2 = _snapshot_row(batch, FA, SKU2, 1)
    assert row2["origin"] == "manual_added"
    assert _snapshot_row(batch, FA, SKU1, 1)["origin"] == "baseline"

    # 幂等：快照已纳入 → 再次 diff 不再报 added_rows
    diff = copy_snapshot.diff_against_copy(batch)
    assert diff is not None and diff.added_rows == ()


# ---------------------------------------------------------------------------
# 2. A 态无提取结果支撑：人工行保留 + warning + 留痕（审计 source）
# ---------------------------------------------------------------------------

def test_a_state_no_extract_row_kept(caplog):
    batch = "T-MA-A-KEPT"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)

    # 人工新增行并手工填了值（保留人工行不动 = 这些值不被清）
    _append_row(copy_path, _row(sku=SKU2, qty=6, name_cn="手工品名",
                                net=77.7, gross=88.8,
                                po="272753", kanri="XC268499"))

    # 提取结果只有 SKU1（无 SKU2）
    state = _write_state(batch, [_item(SKU1)])
    with caplog.at_level(logging.WARNING, logger="app.nodes.writer"):
        written, _ = writer._write_excel(state, copy_path, channel="first_write")

    assert written == 1                              # 只写 SKU1 行
    assert _cell(copy_path, 3, "净重") == 77.7       # 人工行原样保留
    assert _cell(copy_path, 3, "中文品名") == "手工品名"
    cur = state["current_factory_data"]
    assert cur["manual_rows_kept_no_extract"] == [
        {"factory": FA, "sku": SKU2, "row": 3}]
    assert any("无提取结果支撑" in r.message and SKU2 in r.message
               for r in caplog.records)

    # 审计留痕：traces 第四通道带 source="manual_row_kept_no_extract"
    kept, overwritten, qty_changes, kept_rows = service._manual_write_traces(
        {"factory_outputs": {FA: dict(cur)}}, None)
    assert kept == [] and overwritten == [] and qty_changes == []
    assert len(kept_rows) == 1
    assert kept_rows[0]["source"] == "manual_row_kept_no_extract"
    assert kept_rows[0]["sku"] == SKU2


# ---------------------------------------------------------------------------
# 3. B 态：全新厂入批（add_factories_to_batch 假 graph）+ 快照打标
# ---------------------------------------------------------------------------

class _FakeSnap:
    def __init__(self, values=None, next_nodes=(), tasks=()):
        self.values = values if values is not None else {}
        self.next = tuple(next_nodes)
        self.tasks = tuple(tasks)
        self.created_at = None  # get_batch_detail 的 _summarize_snapshot 消费


class _FakeGraph:
    """记录 update_state / stream 调用次数的假 graph（照 test_write_diff_conflicts）。"""

    def __init__(self, snap: _FakeSnap):
        self._snap = snap
        self.updates: list[dict] = []
        self.stream_calls = 0

    def get_state(self, _cfg):
        return self._snap

    def update_state(self, _cfg, update, **_kw):
        self.updates.append(update)
        self._snap.values.update(update)

    def stream(self, *_a, **_kw):
        self.stream_calls += 1
        return iter([])     # 无事件 → completed 兜底分支


def test_b_state_new_factory_added_to_batch(monkeypatch):
    batch = "T-MA-B-ADD"
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    _make_xlsx(src, [_row()])                    # 原件只有 FA
    copy_path = _make_copy(batch, [_row()], stem=stem)
    copy_snapshot.refresh_snapshot(batch)
    # 人工新增全新厂 FB 的行（原件没有 FB → 走 B 态，不走原件解析）
    _append_row(copy_path, _row(factory=FB, sku=SKU2, qty=5, name_cn=None,
                                net=None, gross=None,
                                po="272900", kanri="XC266877"))

    batch_store.upsert_batch(batch, status="completed")
    _CREATED_TIDS.append(batch)
    service._write_batch_config(batch, {
        "thread_id": batch,
        "downstream_file_path": str(src),
        "upstream_root": str(TMP / "upstream"),
    })
    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {FA: {"calculated_items": []}},
        "pending_factories": [],
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
        "factory_filter": [FA],
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    # 原件解析不含 FB（人工新增行不走原件）——只回 FA
    monkeypatch.setattr(
        "app.nodes.parse_downstream.parse_downstream_file",
        lambda _p: ({FA: [SKU1]}, {FA: {SKU1: [2]}}))
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)

    result = service.add_factories_to_batch(batch)

    # B 态全新厂视同待补充入批
    assert result["added"] == 1 and result["factories"] == [FB]
    assert fake.stream_calls == 1
    update = fake.updates[0]
    assert update["pending_factories"] == [FB]
    # 新厂需求明细来自副本新增行（SKU 去重）
    assert update["downstream_requirements"][FB] == [SKU2]
    assert update["downstream_row_map"][FB] == {SKU2: [3]}
    # 既有厂原样保留
    assert update["downstream_requirements"][FA] == [SKU1]
    # factory_filter 同步扩充（否则新厂不进进度统计口径）
    assert update["factory_filter"] == [FA, FB]

    # sync 后快照重建：FB 行入库且 origin="manual_added"（5.5 打标）
    fb_row = _snapshot_row(batch, FB, SKU2, 1)
    assert fb_row["origin"] == "manual_added"


def test_merge_factories_into_state_unit():
    """公共 helper 单测：整体替换语义 + pending 去重保序 + filter 扩充。"""
    values = {
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
        "pending_factories": ["待办厂"],
        "factory_filter": [FA],
    }
    update = service._merge_factories_into_state(
        values, {FB: [SKU2]}, {FB: {SKU2: [3]}},
        pending_add=[FB, "待办厂"])   # 待办厂已在 pending → 不重复加
    assert update["downstream_requirements"] == {FA: [SKU1], FB: [SKU2]}
    assert update["downstream_row_map"][FB] == {SKU2: [3]}
    assert update["pending_factories"] == ["待办厂", FB]
    # pending_add 全集扩充进 filter（待办厂此前不在 filter → 本次补入）
    assert update["factory_filter"] == [FA, FB, "待办厂"]
    # 无 factory_filter 的批次不带该键（不缺省造出）
    update2 = service._merge_factories_into_state(
        {"downstream_requirements": {}}, {}, {})
    assert "factory_filter" not in update2


# ---------------------------------------------------------------------------
# 4. C 态：缺关键列行不纳入 + 活扫透出 + 详情契约 + 摘要提示
# ---------------------------------------------------------------------------

def test_c_state_scan_incomplete_rows():
    batch = "T-MA-C-SCAN"
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    _make_xlsx(src, [_row()])
    copy_path = _make_copy(batch, [_row()], stem=stem)
    copy_snapshot.refresh_snapshot(batch)

    # 缺管理番号（工厂/SKU 齐）→ C 态
    _append_row(copy_path, _row(sku=SKU2, qty=6, kanri=None))
    # 缺工厂（SKU 有值）→ 进不了快照 diff，只能活扫透出
    _append_row(copy_path, {"PURCHASE_ORDER": "272754", "KANRI_NO": "XC268500",
                            "MAKER_MEI_KJ": None, "SHOHIN_CD": "4900000009302",
                            "SHOHIN_MEI_E": "SAW", "SOTOBAKO_D_HACCHU_SU": 3,
                            "中文品名": None, "净重": None, "毛重": None})
    # 全空排版行不参与
    _append_row(copy_path, {})

    rows = copy_snapshot.scan_incomplete_rows(batch, str(src))
    assert rows == [
        {"row": 3, "factory": FA, "sku": SKU2, "missing": ["KANRI_NO"]},
        {"row": 4, "factory": "", "sku": "4900000009302",
         "missing": ["工厂"]},
    ]

    # 原件本身缺列的同签名行不冤报：把缺番号行也写进原件 → 活扫归零
    _make_xlsx(src, [_row(), _row(sku=SKU2, qty=6, kanri=None)])
    rows2 = copy_snapshot.scan_incomplete_rows(batch, str(src))
    assert [r["row"] for r in rows2] == [4]


def test_c_state_not_merged_and_summary_hint(monkeypatch):
    """缺管理番号的新增行不进 requirements；早退摘要带补全提示。"""
    batch = "T-MA-C-EARLY"
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    _make_xlsx(src, [_row()])
    copy_path = _make_copy(batch, [_row()], stem=stem)
    copy_snapshot.refresh_snapshot(batch)
    _append_row(copy_path, _row(sku=SKU2, qty=6, kanri=None))   # C 态

    batch_store.upsert_batch(batch, status="completed")
    _CREATED_TIDS.append(batch)
    service._write_batch_config(batch, {
        "thread_id": batch,
        "downstream_file_path": str(src),
        "upstream_root": str(TMP / "upstream"),
    })
    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {FA: {"calculated_items": []}},
        "pending_factories": [],
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(
        "app.nodes.parse_downstream.parse_downstream_file",
        lambda _p: ({FA: [SKU1]}, {FA: {SKU1: [2]}}))

    result = service.add_factories_to_batch(batch)

    # C 态不纳入：无新厂、无 state 合并、不跑图
    assert result["added"] == 0 and result["factories"] == []
    assert fake.updates == [] and fake.stream_calls == 0
    # 摘要提示补全（不阻断）
    assert "没有待补充的工厂" in result["message"]
    assert "1 行人工新增的行缺管理番号" in result["message"]
    assert "请补全" in result["message"]


def test_batch_detail_manual_rows_incomplete_contract(monkeypatch):
    """前端契约：GET /api/v1/batches/{tid} 顶层 manual_rows_incomplete。"""
    batch = "T-MA-C-DETAIL"
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    _make_xlsx(src, [_row()])
    copy_path = _make_copy(batch, [_row()], stem=stem)
    copy_snapshot.refresh_snapshot(batch)
    _append_row(copy_path, _row(sku=SKU2, qty=6, kanri=None))

    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "downstream_file_path": str(src),
        "factory_outputs": {FA: {"calculated_items": []}},
        "pending_factories": [],
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)

    detail = service.get_batch_detail(batch)
    assert detail["manual_rows_incomplete"] == [
        {"row": 3, "factory": FA, "sku": SKU2, "missing": ["KANRI_NO"]}]

    # 无缺列行时为空数组（契约：恒在）
    batch2 = "T-MA-C-DETAIL-OK"
    _make_copy(batch2, [_row()])
    copy_snapshot.refresh_snapshot(batch2)
    fake2 = _FakeGraph(_FakeSnap(values={
        "batch_id": batch2,
        "downstream_file_path": str(src),
        "factory_outputs": {FA: {"calculated_items": []}},
        "pending_factories": [],
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake2)
    assert service.get_batch_detail(batch2)["manual_rows_incomplete"] == []


# ---------------------------------------------------------------------------
# 5. loader 一致性：缺 KANRI_NO 行 → 计数 warning + 行号收集
# ---------------------------------------------------------------------------

SPLIT_HEADERS = ["KANRI_NO", "MINATO_MEI_KJ", "CONTAINER_MEI", "MAKER_MEI_KJ",
                 "SHOHIN_CD", "净重", "毛重", "SOTOBAKO_D_HACCHU_SU",
                 "中文品名", "D_HACCHU_SU", "KAKAKUKEI", "TSUKA_MEI", "M3",
                 "PURCHASE_ORDER"]


def _make_filled(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(SPLIT_HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in SPLIT_HEADERS])
    wb.save(path)


def _split_row(kanri="XC268452", maker=FA, sku=SKU1, qty=10):
    return {"KANRI_NO": kanri, "MINATO_MEI_KJ": "青島", "CONTAINER_MEI": "40HC",
            "MAKER_MEI_KJ": maker, "SHOHIN_CD": sku, "净重": 15.0, "毛重": 25.0,
            "SOTOBAKO_D_HACCHU_SU": qty, "中文品名": "钳子",
            "D_HACCHU_SU": qty, "KAKAKUKEI": 100.0, "TSUKA_MEI": "USD",
            "M3": 0.05, "PURCHASE_ORDER": "272752"}


def test_loader_counts_missing_kanri(tmp_path, caplog):
    filled = tmp_path / "filled.xlsx"
    _make_filled(filled, [
        _split_row(),
        _split_row(kanri=None, sku=SKU2, qty=6),   # 人工新增缺番号行 → 跳过计数
        {},                                        # 纯空行 → 不计
    ])

    skipped: list[int] = []
    with caplog.at_level(logging.WARNING, logger="app.split.loader"):
        items = load_filled_excel(filled, skipped_kanri_rows=skipped)

    # 缺番号行不参与分票；纯排版空行不计入 warning
    assert [i.kanri_no for i in items] == ["XC268452"]
    assert skipped == [3]
    assert any("1 行缺管理番号" in r.message and "未参与分票" in r.message
               for r in caplog.records)

    # 无缺番号行时零 warning 零收集
    _make_filled(filled, [_split_row()])
    skipped2: list[int] = []
    caplog.clear()
    with caplog.at_level(logging.WARNING, logger="app.split.loader"):
        items2 = load_filled_excel(filled, skipped_kanri_rows=skipped2)
    assert len(items2) == 1 and skipped2 == []
    assert not any("缺管理番号" in r.message for r in caplog.records)


# ---------------------------------------------------------------------------
# 6. 墓碑衔接与对称：加回出碑走三态 / 剔除厂加回 readded / 纳入后再删入碑
# ---------------------------------------------------------------------------

def test_tombstone_readd_exits_and_goes_three_states():
    batch = "T-MA-TOMB-READD"
    copy_path = _make_copy(batch, [_row(), _row(sku=SKU2, qty=5, po="272753",
                                                kanri="XC268499")])
    copy_snapshot.refresh_snapshot(batch)

    # 人工删 SKU2 行 → 入碑
    _delete_row(copy_path, 3)
    copy_snapshot.refresh_snapshot(batch)
    key = copy_snapshot.row_key_str((FA, SKU2, 1))
    assert key in copy_snapshot.load_snapshot(batch)["tombstones"]

    # 人工把该行加回（命中墓碑）
    _append_row(copy_path, _row(sku=SKU2, qty=5, po="272753", kanri="XC268499"))
    diff = copy_snapshot.diff_against_copy(batch)
    assert diff is not None and len(diff.added_rows) == 1
    assert diff.added_rows[0]["in_tombstones"] is True

    # 出碑后照常走三态：厂在原件需求集 → A 态
    classified = copy_snapshot.classify_added_rows(diff.added_rows, {FA})
    assert len(classified["a"]) == 1
    assert classified["b"] == [] and classified["c"] == []

    # refresh 结转：键出墓碑、行回快照
    copy_snapshot.refresh_snapshot(batch)
    snap = copy_snapshot.load_snapshot(batch)
    assert key not in snap["tombstones"]
    assert _snapshot_row(batch, FA, SKU2, 1)


def test_removed_factory_readd_writes_readded_audit():
    """剔除厂加回（决策 2 对称恢复经决策 4 检出）：readded 审计落库。"""
    batch = "T-MA-TOMB-READD-FAC"
    copy_path = _make_copy(batch, [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, po="272900", kanri="XC266877"),
    ])
    copy_snapshot.refresh_snapshot(batch)

    # 剔除：删光 FB 行 → removed 审计 + 入碑
    _delete_row(copy_path, 3)
    diff = copy_snapshot.diff_against_copy(batch)
    assert copy_snapshot._detect_factory_removals(batch, diff, {FA, FB}) == [FB]
    copy_snapshot.refresh_snapshot(batch)
    assert service._manually_removed_factories(batch) == {FB}

    # 人工加回 FB 行 → 三态分类照走（FB 仍在需求集 → A 态）+ readded 审计
    _append_row(copy_path, _row(factory=FB, sku=SKU2, qty=5, po="272900",
                                kanri="XC266877"))
    diff2 = copy_snapshot.diff_against_copy(batch)
    fb_added = [a for a in diff2.added_rows if a["factory"] == FB]
    assert len(fb_added) == 1 and fb_added[0]["in_tombstones"] is True
    classified = copy_snapshot.classify_added_rows(diff2.added_rows, {FA, FB})
    assert len(classified["a"]) == 1
    assert copy_snapshot._detect_factory_removals(batch, diff2, {FA, FB}) == []
    assert [a.result_status for a in _audits(batch)] == [
        "factory_manually_removed", "factory_manually_readded"]
    assert service._manually_removed_factories(batch) == set()


def test_manual_added_row_then_deleted_goes_tombstone():
    """对称：纳入管理（manual_added）的行人工再删 → 回删除分支进墓碑。"""
    batch = "T-MA-SYM-DELETE"
    copy_path = _make_copy(batch, [_row()])
    copy_snapshot.refresh_snapshot(batch)

    # 纳入：人工新增 SKU2 行 → writer 写入 → 快照 origin=manual_added
    _append_row(copy_path, _row(sku=SKU2, qty=6, po="272753", kanri="XC268499"))
    state = _write_state(batch, [_item(SKU1), _item(SKU2)])
    writer._write_excel(state, copy_path, channel="first_write")
    assert _snapshot_row(batch, FA, SKU2, 1)["origin"] == "manual_added"

    # 人工再删该行 → diff 删除分支 → 入碑不复活
    _delete_row(copy_path, 3)
    diff = copy_snapshot.diff_against_copy(batch)
    key = copy_snapshot.row_key_str((FA, SKU2, 1))
    assert diff is not None and key in diff.deleted_rows
    copy_snapshot.refresh_snapshot(batch)
    snap = copy_snapshot.load_snapshot(batch)
    assert key in snap["tombstones"]
    assert all(r["key"]["sku"] != SKU2 for r in snap["rows"])


if __name__ == "__main__":
    import pytest as _pytest
    _pytest.main([__file__, "-v"])
