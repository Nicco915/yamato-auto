# -*- coding: utf-8 -*-
"""整厂剔除（写入侧人工改动感知 · 阶段四 §4.6，决策 2）测试。

覆盖（对应任务五项）：
1. 剔除检出：副本删光某厂行 → 写前 diff + _detect_factory_removals →
   factory_manually_removed 审计落库（approved=False + changes_json 行数）
   + 幂等不重复写 + 不在 requirements 不判定 + refresh 后墓碑含该厂行键
   （审计与墓碑双轨印证）；
2. 剔除后完成判定：_classify_incomplete_batch 对「唯一缺口=剔除厂」返回
   completed（无审计 → awaiting 对照；混合缺口 → 仍 awaiting）；
3. factory_watch：剔除厂 gap 减除后不触发（资料到达也不补充），对照批
   正常触发；
4. add_factories_to_batch：剔除厂不进 added（假 graph + 假解析，照
   test_write_diff_conflicts 通道 B 模式）+ info 日志「不再补回」+
   对照批正常补入；
5. 对称恢复：人工加回剔除厂行 → factory_manually_readded 审计落库 →
   _manually_removed_factories 自动出名单 + 出墓碑。
另：get_batch_detail role 透出 manually_removed（前端徽章数据源）。

运行：
    cd app && PYTHONPATH=. python3 -m pytest tests/test_factory_removal.py -v

隔离红线（2026-08-11 事故教训，骨架照抄 test_write_diff_conflicts.py）：
- import app 模块前先设 YAMATO_TEST_MODE=1 + YAMATO_DOTENV_PATH=临时空 .env；
- import 后调 validation/_test_isolation.isolate_to_tmp 把
  checkpoint/master/output/sessions 全部指向临时目录。
"""
from __future__ import annotations

import json
import logging
import os
import sys
import tempfile
from pathlib import Path

import pytest

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

# ---- 隔离门①：import 任何 app 模块前的双门（血泪红线）----
_TMP = Path(tempfile.mkdtemp(prefix="yamato_factory_removal_test_"))
(_TMP / ".env").write_text("# isolated .env\n", encoding="utf-8")
os.environ["YAMATO_TEST_MODE"] = "1"
os.environ["YAMATO_DOTENV_PATH"] = str(_TMP / ".env")
os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.api import service  # noqa: E402
from app.api.main import app  # noqa: E402,F401  确保全部 app 模块先于隔离 import
from app.config import get_settings  # noqa: E402
from app.db import batch_store  # noqa: E402
from app.db.models import ReviewAudit  # noqa: E402
from app.db.session import get_session  # noqa: E402
from app.graph import NODE7, get_graph  # noqa: E402
from app.orchestrator import copy_snapshot, factory_watch, pipeline_state  # noqa: E402

# ---- 隔离门②：import 后把 db/output/sessions 指向临时目录 ----
from _test_isolation import isolate_to_tmp  # noqa: E402

TMP = isolate_to_tmp("yamato_factory_removal_test_")


@pytest.fixture(autouse=True)
def _settings_pin():
    """settings 单例钉根（照 test_write_diff_conflicts 先例）：全量混跑时
    其他测试文件的 isolate_to_tmp 会重钉 output_dir——每个用例前钉回本模块
    的 TMP，结束后恢复原值。"""
    s = get_settings()
    original = s.output_dir
    s.output_dir = str(TMP / "output")
    try:
        yield
    finally:
        s.output_dir = original


@pytest.fixture(autouse=True)
def _no_auto_backup(monkeypatch):
    """结案自动备份打空（照 test_awaiting_factories 先例，防全量混跑撞锁）。"""
    from app.orchestrator import backup
    monkeypatch.setattr(backup, "trigger_async", lambda reason, **kw: True)
    yield


@pytest.fixture(autouse=True)
def _clean_watch_state():
    """每个用例前清监控线程模块级状态；结束后删本模块造的批次行
    （共享临时库，残留 awaiting 批次会被其他用例的 _scan_once 捞到）。"""
    factory_watch.stop_factory_watch()
    with factory_watch._state_lock:
        factory_watch._snapshots.clear()
        factory_watch._cooldowns.clear()
        factory_watch._last_triggered.clear()
    try:
        yield
    finally:
        factory_watch.stop_factory_watch()
        with factory_watch._state_lock:
            factory_watch._snapshots.clear()
            factory_watch._cooldowns.clear()
            factory_watch._last_triggered.clear()
        for tid in _CREATED_TIDS:
            batch_store.delete_batch(tid)
        _CREATED_TIDS.clear()


_CREATED_TIDS: list[str] = []

FA = "山東中地"          # 已完成/在批工厂
FB = "Ｃ．正達工芸品"     # 被整厂剔除的工厂
SKU1 = "4550596211831"
SKU2 = "4900000009201"

HEADERS = ["PURCHASE_ORDER", "KANRI_NO", "MAKER_MEI_KJ", "SHOHIN_CD",
           "SHOHIN_MEI_E", "SOTOBAKO_D_HACCHU_SU", "中文品名", "净重", "毛重"]
COL = {h: HEADERS.index(h) + 1 for h in HEADERS}


def _make_xlsx(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    wb.save(path)


def _make_copy(batch_id: str, rows: list[dict],
               stem: str | None = None) -> Path:
    """在批次 containers/ 下造 *_filled.xlsx 副本（口径同
    test_write_diff_conflicts._make_copy）。"""
    cdir = get_settings().batch_containers_dir(batch_id)
    cdir.mkdir(parents=True, exist_ok=True)
    stem = stem or f"ContentsOfTheContainer_{batch_id}"
    path = cdir / f"{stem}_filled.xlsx"
    _make_xlsx(path, rows)
    return path


def _row(factory=FA, sku=SKU1, qty=10, name_cn="钳子", net=15.0, gross=25.0,
         po="272752", kanri="XC268452"):
    return {"PURCHASE_ORDER": po, "KANRI_NO": kanri, "MAKER_MEI_KJ": factory,
            "SHOHIN_CD": sku, "SHOHIN_MEI_E": "PLIERS",
            "SOTOBAKO_D_HACCHU_SU": qty, "中文品名": name_cn,
            "净重": net, "毛重": gross}


def _delete_row(path: Path, row_no: int) -> None:
    """模拟人工删副本整行。"""
    wb = load_workbook(path)
    wb.active.delete_rows(row_no, 1)
    wb.save(path)


def _audits(batch: str) -> list[ReviewAudit]:
    with get_session() as session:
        return list(session.scalars(
            select(ReviewAudit)
            .where(ReviewAudit.thread_id == batch)
            .order_by(ReviewAudit.audit_id)
        ).all())


def _write_removed_audit(batch: str, factory: str, rows: int = 1) -> None:
    """直接插一条剔除审计（不经过检出链路，供排除点/名单测试造现场）。"""
    with get_session() as session:
        session.add(ReviewAudit(
            thread_id=batch, factory_name=factory, approved=False,
            edited_count=0,
            changes_json=json.dumps([{"factory_removed": True, "rows": rows}],
                                    ensure_ascii=False),
            new_skus_json="[]",
            result_status="factory_manually_removed"))
        session.commit()


# ---------------------------------------------------------------------------
# 1. 剔除检出：审计落库 + 幂等 + requirements 门 + 墓碑双轨印证
# ---------------------------------------------------------------------------

def test_removal_detection_audit_idempotent_and_tombstones():
    batch = "T-FR-DETECT"
    copy_path = _make_copy(batch, [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, po="272900", kanri="XC266877"),
    ])
    copy_snapshot.refresh_snapshot(batch)   # 快照含 FA + FB 行

    _delete_row(copy_path, 3)               # 人工删光 FB 全部行

    diff = copy_snapshot.diff_against_copy(batch)
    assert diff is not None
    assert diff.deleted_rows == (copy_snapshot.row_key_str((FB, SKU2, 1)),)

    # requirements 门：FB 不在需求集 → 不判定剔除、不落审计
    assert copy_snapshot._detect_factory_removals(batch, diff, {FA}) == []
    assert _audits(batch) == []

    # 正式检出：FB 在 requirements 且快照行全灭 → removed 审计落库
    newly = copy_snapshot._detect_factory_removals(batch, diff, {FA, FB})
    assert newly == [FB]
    audits = _audits(batch)
    assert len(audits) == 1
    a = audits[0]
    assert a.result_status == "factory_manually_removed"
    assert a.factory_name == FB
    assert a.approved is False
    assert json.loads(a.changes_json) == [{"factory_removed": True, "rows": 1}]

    # 幂等：最新审计已是 removed → 重复检出不重复写
    assert copy_snapshot._detect_factory_removals(batch, diff, {FA, FB}) == []
    assert len(_audits(batch)) == 1

    # 名单 helper：FB 在剔除名单
    assert service._manually_removed_factories(batch) == {FB}

    # 墓碑双轨印证：refresh 结转后该厂行键入碑、快照行集不再含 FB
    copy_snapshot.refresh_snapshot(batch)
    snap = copy_snapshot.load_snapshot(batch)
    fb_key = copy_snapshot.row_key_str((FB, SKU2, 1))
    assert fb_key in snap["tombstones"]
    assert all(r["key"]["factory"] != FB for r in snap["rows"])


def test_removal_not_detected_when_factory_partially_deleted():
    """部分删行不算整厂剔除（快照行未全灭）。"""
    batch = "T-FR-PARTIAL"
    copy_path = _make_copy(batch, [
        _row(factory=FB, sku=SKU2, qty=5, po="272900", kanri="XC266877"),
        _row(factory=FB, sku=SKU2, qty=3, po="272901", kanri="XC266877"),
    ])
    copy_snapshot.refresh_snapshot(batch)   # FB 同 SKU 两行（seq 1/2）

    _delete_row(copy_path, 2)               # 只删一行

    diff = copy_snapshot.diff_against_copy(batch)
    assert diff is not None and len(diff.deleted_rows) == 1
    assert copy_snapshot._detect_factory_removals(batch, diff, {FB}) == []
    assert _audits(batch) == []
    assert service._manually_removed_factories(batch) == set()


# ---------------------------------------------------------------------------
# 2. 剔除后完成判定：唯一缺口=剔除厂 → completed
# ---------------------------------------------------------------------------

def test_classify_incomplete_batch_excludes_removed():
    batch = "T-FR-CLASSIFY"
    # 无审计 → 缺口未定性，待补充
    assert pipeline_state._classify_incomplete_batch(batch, [FB], None) \
        == "awaiting_factories"

    _write_removed_audit(batch, FB)
    # 唯一缺口=剔除厂 → completed（剔除厂不再算缺口）
    assert pipeline_state._classify_incomplete_batch(batch, [FB], None) \
        == "completed"
    # 混合缺口（剔除厂 + 真缺口）→ 仍待补充
    assert pipeline_state._classify_incomplete_batch(
        batch, [FB, "另一缺口厂"], None) == "awaiting_factories"


# ---------------------------------------------------------------------------
# 3. factory_watch：剔除厂 gap 减除后不触发（资料到达也不监控）
# ---------------------------------------------------------------------------

def _make_awaiting_batch(tid: str, upstream: Path, *,
                         done: list[str], gap: list[str]) -> None:
    """照 test_factory_watch 先例：awaiting 批次行 + 图已跑完 checkpoint
    （as_node=NODE7，next 为空）。"""
    batch_store.upsert_batch(tid, status="awaiting_factories")
    _CREATED_TIDS.append(tid)
    graph = get_graph()
    cfg = {"configurable": {"thread_id": tid}}
    graph.update_state(cfg, {
        "batch_id": tid,
        "downstream_requirements": {f: ["4900000000001"] for f in done + gap},
        "factory_outputs": {f: {"calculated_items": []} for f in done},
        "upstream_root": str(upstream),
        "final_output_path": "/tmp/out.xlsx",
    }, as_node=NODE7)


def test_factory_watch_removed_factory_not_triggered(tmp_path, monkeypatch):
    calls: list[str] = []

    def fake_add(tid: str):
        calls.append(tid)
        return {"added": 1, "factories": [], "status": "pending_human_review"}

    monkeypatch.setattr(service, "add_factories_to_batch", fake_add)

    upstream = tmp_path / "upstream"
    upstream.mkdir()
    # 对照批：正常缺口厂，资料到达应触发
    (upstream / "缺口厂C").mkdir()
    (upstream / "缺口厂C" / "packing.pdf").write_bytes(b"x" * 128)
    _make_awaiting_batch("T-FR-WATCH-CTL", upstream,
                         done=["完成厂C"], gap=["缺口厂C"])
    # 剔除批：缺口厂已被人工剔除，资料到达也不触发
    (upstream / "剔除厂R").mkdir()
    (upstream / "剔除厂R" / "packing.pdf").write_bytes(b"x" * 128)
    _make_awaiting_batch("T-FR-WATCH-RM", upstream,
                         done=["完成厂R"], gap=["剔除厂R"])
    _write_removed_audit("T-FR-WATCH-RM", "剔除厂R")

    factory_watch._scan_once()   # 首轮仅记录快照
    factory_watch._scan_once()   # 第二轮稳定 → 触发
    factory_watch._scan_once()   # 再扫一轮防迟发

    assert calls.count("T-FR-WATCH-CTL") == 1     # 对照批正常触发
    assert "T-FR-WATCH-RM" not in calls           # 剔除厂不再监控


# ---------------------------------------------------------------------------
# 4. add_factories_to_batch：剔除厂不进 added（假 graph + 假解析）
# ---------------------------------------------------------------------------

class _FakeSnap:
    def __init__(self, values=None, next_nodes=(), tasks=()):
        self.values = values if values is not None else {}
        self.next = tuple(next_nodes)
        self.tasks = tuple(tasks)
        self.created_at = None  # get_batch_detail 的 _summarize_snapshot 消费


class _FakeGraph:
    """记录 update_state / stream 调用次数的假 graph（照 test_write_diff_conflicts）。"""

    def __init__(self, snap: _FakeSnap):
        self._snap = snap
        self.updates: list[dict] = []
        self.stream_calls = 0

    def get_state(self, _cfg):
        return self._snap

    def update_state(self, _cfg, update, **_kw):
        self.updates.append(update)
        self._snap.values.update(update)

    def stream(self, *_a, **_kw):
        self.stream_calls += 1
        return iter([])     # 无事件 → completed 兜底分支


def _add_factories_setup(batch: str) -> Path:
    """造补充工厂现场：原件含 FA+FB，副本同内容且有快照；FA 已完成、
    FB 待补充。返回原件路径。"""
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    src_rows = [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, name_cn=None, net=None, gross=None,
             po="272900", kanri="XC266877"),
    ]
    _make_xlsx(src, src_rows)
    _make_copy(batch, src_rows, stem=stem)
    copy_snapshot.refresh_snapshot(batch)
    batch_store.upsert_batch(batch, status="completed")
    _CREATED_TIDS.append(batch)
    service._write_batch_config(batch, {
        "thread_id": batch,
        "downstream_file_path": str(src),
        "upstream_root": str(TMP / "upstream"),
    })
    return src


def _add_factories_graph(batch: str) -> _FakeGraph:
    return _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {FA: {"calculated_items": []}},
        "pending_factories": [],
        "downstream_requirements": {FA: {"total_qty": 10}},
        "downstream_row_map": {FA: {SKU1: [2]}},
    }))


def _patch_parse(monkeypatch):
    monkeypatch.setattr(
        "app.nodes.parse_downstream.parse_downstream_file",
        lambda _p: ({FA: {"total_qty": 10}, FB: {"total_qty": 5}},
                    {FA: {SKU1: [2]}, FB: {SKU2: [3]}}))


def test_add_factories_removed_factory_not_readded(monkeypatch, caplog):
    batch = "T-FR-ADD-RM"
    _add_factories_setup(batch)
    _write_removed_audit(batch, FB)     # FB 已被人工剔除
    fake = _add_factories_graph(batch)
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    _patch_parse(monkeypatch)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)

    with caplog.at_level(logging.INFO, logger="app.api.service"):
        result = service.add_factories_to_batch(batch)

    # 剔除厂不进 added，不入图；info 日志明示「不再补回」
    assert result["added"] == 0
    assert result["factories"] == []
    assert result["message"] == "没有待补充的工厂"
    assert fake.updates == [] and fake.stream_calls == 0
    assert any("已被人工剔除，不再补回" in r.message and FB in r.message
               for r in caplog.records)


def test_add_factories_normal_factory_still_added(monkeypatch):
    """对照：无剔除审计时待补充厂正常补入（不新增回归）。"""
    batch = "T-FR-ADD-CTL"
    _add_factories_setup(batch)
    fake = _add_factories_graph(batch)
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    _patch_parse(monkeypatch)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)

    result = service.add_factories_to_batch(batch)
    assert result["status"] == "completed"
    assert result["added"] == 1 and result["factories"] == [FB]
    assert fake.stream_calls == 1
    # 剔除厂不进缺口文案；此处无剔除，FB 仍是缺口（假图未真写入）
    assert result["awaiting_factories"] == [FB]


# ---------------------------------------------------------------------------
# 5. 对称恢复：人工加回 → readded 审计 → 自动出名单 + 出墓碑
# ---------------------------------------------------------------------------

def test_symmetric_readd_exits_removed_list():
    batch = "T-FR-READD"
    copy_path = _make_copy(batch, [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, po="272900", kanri="XC266877"),
    ])
    copy_snapshot.refresh_snapshot(batch)

    # 剔除：删光 FB 行 → 检出 → 墓碑入碑
    _delete_row(copy_path, 3)
    diff = copy_snapshot.diff_against_copy(batch)
    assert copy_snapshot._detect_factory_removals(batch, diff, {FA, FB}) == [FB]
    copy_snapshot.refresh_snapshot(batch)
    fb_key = copy_snapshot.row_key_str((FB, SKU2, 1))
    assert fb_key in copy_snapshot.load_snapshot(batch)["tombstones"]
    assert service._manually_removed_factories(batch) == {FB}

    # 人工把 FB 行加回副本尾部（墓碑键重新出现）
    wb = load_workbook(copy_path)
    wb.active.append(["272900", "XC266877", FB, SKU2, "WRENCH", 5,
                      None, None, None])
    wb.save(copy_path)

    diff2 = copy_snapshot.diff_against_copy(batch)
    assert diff2 is not None
    fb_added = [a for a in diff2.added_rows if a["factory"] == FB]
    assert len(fb_added) == 1 and fb_added[0]["in_tombstones"] is True

    # 对称恢复：不写剔除、写 readded 审计
    assert copy_snapshot._detect_factory_removals(batch, diff2, {FA, FB}) == []
    audits = _audits(batch)
    assert [a.result_status for a in audits] == [
        "factory_manually_removed", "factory_manually_readded"]
    assert json.loads(audits[1].changes_json) == [
        {"factory_readded": True, "rows": 1}]

    # 最新一条不再是 removed → 自动出名单；refresh 后出墓碑
    assert service._manually_removed_factories(batch) == set()
    copy_snapshot.refresh_snapshot(batch)
    snap = copy_snapshot.load_snapshot(batch)
    assert fb_key not in snap["tombstones"]


# ---------------------------------------------------------------------------
# 附：get_batch_detail role 透出 manually_removed（前端徽章数据源）
# ---------------------------------------------------------------------------

def test_batch_detail_role_manually_removed(monkeypatch):
    batch = "T-FR-ROLE"
    _write_removed_audit(batch, FB)
    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {FA: {"calculated_items": []}},
        "pending_factories": [],
        "downstream_requirements": {FA: {"total_qty": 10},
                                    FB: {"total_qty": 5}},
        "downstream_row_map": {FA: {SKU1: [2]}, FB: {SKU2: [3]}},
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)

    detail = service.get_batch_detail(batch)
    roles = {f["factory"]: f["role"] for f in detail["factories"]}
    assert roles[FB] == "manually_removed"
    assert roles[FA] == "done"


if __name__ == "__main__":
    import pytest as _pytest
    _pytest.main([__file__, "-v"])
