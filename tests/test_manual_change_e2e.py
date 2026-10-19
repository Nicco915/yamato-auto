# -*- coding: utf-8 -*-
"""写入侧人工改动感知 · 阶段六 e2e 场景矩阵（计划 §6.1，14 场景）。

与五个阶段单测的差别：单测钉单元口径，本套件串完整链路——临时目录造全套
（原件 xlsx + 上游根 + output 副本 + 批次配置 + 指纹基线），走服务层真实
调用链（add_factories_to_batch / retry_factory_extraction /
reopen_factory_for_edit + apply_reopen_payload / rerun_batch /
_sync_output_copy_rows / writer() 节点 / export_node），图执行部分用
假 graph（update_state/stream 记录）模拟挂起与续跑，不调真实 LLM。

场景矩阵（每场景独立 test，共用 fixture 工厂函数）：
 1 改空覆盖；2 改非空通道C跳过（副本=人工值+审计 manual_override_kept+
   摘要+快照 protected）；3 改非空通道A覆盖（reopen 全链路，审计
   manual_value_overwritten）；4 通道B确认门（未确认 409 未入图/不确认
   保留/确认后 state 标记+Node6 覆写）；5 删行不复活（墓碑含键）；
 6 整厂剔除（审计 removed+完成判定 completed+watch 不触发+补充不加回）；
 7 剔除后加回（出碑+readded 审计+重新纳入 pending）；8 移位重定位
   （插 3 行后 reopen 写对清对）；9 新增行三态（A/B/C 各一断言）；
 10 无快照迁移（baseline+warning，基线前改动不保护）；11 快照结转
   （二次写入不覆盖 protected+墓碑不复活）；12 原件差异三入口各一断言；
 13 完全重跑警告（未确认拒绝+确认后归档可查+副本重建）；14 数量列告警
   （按新数量写入+warning+摘要）。

运行（worktree 根）：
    python3 -m pytest tests/test_manual_change_e2e.py -v

隔离红线（2026-08-11 事故教训，骨架照抄 test_factory_removal.py）：
- import app 模块前先设 YAMATO_TEST_MODE=1 + YAMATO_DOTENV_PATH=临时空 .env；
- import 后调 validation/_test_isolation.isolate_to_tmp 把
  checkpoint/master/output/sessions 全部指向临时目录。
"""
from __future__ import annotations

import json
import logging
import os
import sys
import tempfile
from pathlib import Path

import pytest

APP_ROOT = Path(__file__).resolve().parent.parent
if str(APP_ROOT) not in sys.path:
    sys.path.insert(0, str(APP_ROOT))
sys.path.insert(0, str(APP_ROOT / "validation"))

# ---- 隔离门①：import 任何 app 模块前的双门（血泪红线）----
_TMP = Path(tempfile.mkdtemp(prefix="yamato_manual_change_e2e_"))
(_TMP / ".env").write_text("# isolated .env\n", encoding="utf-8")
os.environ["YAMATO_TEST_MODE"] = "1"
os.environ["YAMATO_DOTENV_PATH"] = str(_TMP / ".env")
os.environ["EXTRACTION_MOCK"] = "1"
os.environ["DISPATCHER_MOCK"] = "1"

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.api import service  # noqa: E402
from app.api.main import app  # noqa: E402,F401  确保全部 app 模块先于隔离 import
from app.config import get_settings  # noqa: E402
from app.db import batch_store  # noqa: E402
from app.db.models import ReviewAudit  # noqa: E402
from app.db.session import get_session  # noqa: E402
from app.graph import NODE7, get_graph  # noqa: E402
from app.nodes import writer  # noqa: E402
from app.nodes.export_node import export_node  # noqa: E402
from app.orchestrator import (copy_snapshot, factory_watch,  # noqa: E402
                              pipeline_state, source_diff)

# ---- 隔离门②：import 后把 db/output/sessions 指向临时目录 ----
from _test_isolation import isolate_to_tmp  # noqa: E402

TMP = isolate_to_tmp("yamato_manual_change_e2e_")

client = TestClient(app)


@pytest.fixture(autouse=True)
def _settings_pin():
    """settings 单例钉根（照 test_factory_removal 先例）：全量混跑时其他
    测试文件的 isolate_to_tmp 会重钉 output_dir——每个用例前钉回本模块
    的 TMP，结束后恢复原值。"""
    s = get_settings()
    original = s.output_dir
    s.output_dir = str(TMP / "output")
    try:
        yield
    finally:
        s.output_dir = original


@pytest.fixture(autouse=True)
def _no_auto_backup(monkeypatch):
    """结案自动备份打空（照 test_awaiting_factories 先例，防全量混跑撞锁）。"""
    from app.orchestrator import backup
    monkeypatch.setattr(backup, "trigger_async", lambda reason, **kw: True)
    yield


@pytest.fixture(autouse=True)
def _clean_watch_state():
    """每个用例前清监控线程模块级状态；结束后删本模块造的批次行
    （共享临时库，残留 awaiting 批次会被其他用例的 _scan_once 捞到）。"""
    factory_watch.stop_factory_watch()
    with factory_watch._state_lock:
        factory_watch._snapshots.clear()
        factory_watch._cooldowns.clear()
        factory_watch._last_triggered.clear()
    try:
        yield
    finally:
        factory_watch.stop_factory_watch()
        with factory_watch._state_lock:
            factory_watch._snapshots.clear()
            factory_watch._cooldowns.clear()
            factory_watch._last_triggered.clear()
        for tid in _CREATED_TIDS:
            batch_store.delete_batch(tid)
        _CREATED_TIDS.clear()


_CREATED_TIDS: list[str] = []

FA = "山東中地"          # 已完成/在批工厂
FB = "Ｃ．正達工芸品"     # 待补充/被剔除工厂
FC = "青島新厂"          # 人工新增的全新工厂（B 态）
SKU1 = "4550596211831"
SKU2 = "4900000009201"

HEADERS = ["PURCHASE_ORDER", "KANRI_NO", "MAKER_MEI_KJ", "SHOHIN_CD",
           "SHOHIN_MEI_E", "SOTOBAKO_D_HACCHU_SU", "中文品名", "净重", "毛重"]
COL = {h: HEADERS.index(h) + 1 for h in HEADERS}


def _make_xlsx(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    wb.save(path)


def _make_copy(batch_id: str, rows: list[dict],
               stem: str | None = None) -> Path:
    """在批次 containers/ 下造 *_filled.xlsx 副本（口径同阶段二~五测试）。"""
    cdir = get_settings().batch_containers_dir(batch_id)
    cdir.mkdir(parents=True, exist_ok=True)
    stem = stem or f"ContentsOfTheContainer_{batch_id}"
    path = cdir / f"{stem}_filled.xlsx"
    _make_xlsx(path, rows)
    return path


def _row(factory=FA, sku=SKU1, qty=10, name_cn="钳子", net=15.0, gross=25.0,
         po="272752", kanri="XC268452"):
    """标准数据行（缺省为「系统已写入」状态：三列有值）。"""
    return {"PURCHASE_ORDER": po, "KANRI_NO": kanri, "MAKER_MEI_KJ": factory,
            "SHOHIN_CD": sku, "SHOHIN_MEI_E": "PLIERS",
            "SOTOBAKO_D_HACCHU_SU": qty, "中文品名": name_cn,
            "净重": net, "毛重": gross}


def _ws(path: Path):
    return load_workbook(path).active


def _cell(path: Path, row: int, field: str):
    return _ws(path).cell(row=row, column=COL[field]).value


def _set_cell(path: Path, row: int, field: str, value) -> None:
    """模拟人工改副本/原件单元格。"""
    wb = load_workbook(path)
    wb.active.cell(row=row, column=COL[field]).value = value
    wb.save(path)


def _append_row(path: Path, row: dict) -> None:
    """模拟人工在文件尾部新增一行。"""
    wb = load_workbook(path)
    wb.active.append([row.get(h) for h in HEADERS])
    wb.save(path)


def _delete_row(path: Path, row_no: int) -> None:
    """模拟人工删副本整行。"""
    wb = load_workbook(path)
    wb.active.delete_rows(row_no, 1)
    wb.save(path)


def _item(sku: str, unit_net=1.5, unit_gross=2.5, name_cn="钳子") -> dict:
    return {"sku": sku, "name_cn": name_cn,
            "extracted_data": {"total_quantity": 10,
                               "total_net_weight": 15.0,
                               "total_gross_weight": 25.0},
            "calculation": {"calculated_unit_net": unit_net,
                            "calculated_unit_gross": unit_gross}}


def _write_state(batch_id: str, factory: str = FA, items=None,
                 requirements=None, row_map=None) -> dict:
    """Node6 写入 state（口径照 test_write_diff_conflicts._write_state）。"""
    items = items if items is not None else [_item(SKU1)]
    return {
        "batch_id": batch_id,
        "current_factory_data": {"factory_name": factory,
                                 "calculated_items": items},
        "downstream_requirements": requirements or {factory: [SKU1]},
        # row_map 只承载「行数需求」（len），行号值不消费
        "downstream_row_map": row_map or {factory: {SKU1: [2]}},
    }


def _snapshot_row(batch_id: str, factory: str, sku: str, seq: int) -> dict:
    snap = copy_snapshot.load_snapshot(batch_id)
    for r in snap["rows"]:
        k = r["key"]
        if (k["factory"], k["sku"], k["seq"]) == (factory, sku, seq):
            return r
    raise AssertionError(f"快照中无行 {(factory, sku, seq)}")


def _audits(batch: str) -> list[ReviewAudit]:
    with get_session() as session:
        return list(session.scalars(
            select(ReviewAudit)
            .where(ReviewAudit.thread_id == batch)
            .order_by(ReviewAudit.audit_id)
        ).all())


def _setup_batch(batch: str, src_rows: list[dict], *, snapshot: bool = True):
    """造批次全套现场：原件 + 副本 + 快照（可选）+ 批次行 + 批次配置 +
    原件指纹基线。返回 (src, copy_path)。原件 stem 与副本同 stem
    （sync 定位口径）。"""
    stem = f"ContentsOfTheContainer_{batch}"
    src = TMP / "src" / f"{stem}.xlsx"
    _make_xlsx(src, src_rows)
    copy_path = _make_copy(batch, src_rows, stem=stem)
    if snapshot:
        copy_snapshot.refresh_snapshot(batch)
    batch_store.upsert_batch(batch, status="completed")
    _CREATED_TIDS.append(batch)
    service._write_batch_config(batch, {
        "thread_id": batch,
        "downstream_file_path": str(src),
        "upstream_root": str(TMP / "upstream"),
    })
    source_diff.record_fingerprint(batch, str(src))
    return src, copy_path


class _FakeSnap:
    def __init__(self, values=None, next_nodes=(), tasks=()):
        self.values = values if values is not None else {}
        self.next = tuple(next_nodes)
        self.tasks = tuple(tasks)
        self.created_at = None  # get_batch_detail 的 _summarize_snapshot 消费


class _FakeGraph:
    """记录 update_state / stream 调用次数的假 graph（照阶段二~五先例）。"""

    def __init__(self, snap: _FakeSnap):
        self._snap = snap
        self.updates: list[dict] = []
        self.stream_calls = 0

    def get_state(self, _cfg):
        return self._snap

    def update_state(self, _cfg, update, **_kw):
        self.updates.append(update)
        self._snap.values.update(update)

    def stream(self, *_a, **_kw):
        self.stream_calls += 1
        return iter([])     # 无事件 → completed 兜底分支


def _add_factories_graph(batch: str, *, done=(FA,), requirements=None,
                         row_map=None) -> _FakeGraph:
    reqs = requirements if requirements is not None else {FA: [SKU1]}
    rm = row_map if row_map is not None else {FA: {SKU1: [2]}}
    return _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "factory_outputs": {f: {"calculated_items": []} for f in done},
        "pending_factories": [],
        "downstream_requirements": reqs,
        "downstream_row_map": rm,
        "upstream_root": str(TMP / "upstream"),
    }))


def _export_summary(batch: str, cur: dict, copy_path: Path) -> dict:
    """走真实 Node7 导出，返回批次级 export_summary.json。"""
    export_node({
        "batch_id": batch,
        "downstream_file_path": "x.xlsx",
        "downstream_requirements": {FA: [SKU1]},
        "factory_outputs": {FA: dict(cur)},
        "final_output_path": str(copy_path),
    })
    return json.loads(
        (get_settings().batch_output_dir(batch) / "export_summary.json")
        .read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# 场景 1：改空覆盖（快照非空 × 副本空 = 等重新录入，通道 C 正常覆盖）
# ---------------------------------------------------------------------------

def test_e2e_01_cleared_cell_overwritten_channel_c():
    batch = "T-E2E-01-CLEAR"
    _, copy_path = _setup_batch(batch, [_row()])

    _set_cell(copy_path, 2, "净重", None)   # 人工改空（等重新录入）

    state = _write_state(batch)
    written, _ = writer._write_excel(state, copy_path, channel="first_write")

    # 副本=系统新值；改空不算冲突，无留痕无保护
    assert written == 1
    assert _cell(copy_path, 2, "净重") == 15.0
    cur = state["current_factory_data"]
    assert cur["manual_override_kept"] == 0
    assert cur["manual_conflicts"] == []
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is False

    # 链路末端：审计 traces 全空 + 完成摘要三字段为零
    kept, overwritten, qty_changes, kept_rows = service._manual_write_traces(
        {"factory_outputs": {FA: dict(cur)}}, None)
    assert kept == [] and overwritten == [] and kept_rows == []
    summary = _export_summary(batch, cur, copy_path)
    assert summary["manual_overrides_kept"] == 0
    assert summary["manual_values_overwritten"] == 0
    assert summary["qty_manual_changes"] == 0


# ---------------------------------------------------------------------------
# 场景 2：改非空 · 通道 C 跳过（保留人工值 + 审计 + 摘要 + protected）
# ---------------------------------------------------------------------------

def test_e2e_02_channel_c_keeps_manual_value_full_chain():
    batch = "T-E2E-02-CHC"
    _, copy_path = _setup_batch(batch, [_row()])

    _set_cell(copy_path, 2, "净重", 99.0)   # 人工改非空 W

    state = _write_state(batch)
    written, _ = writer._write_excel(state, copy_path, channel="first_write")

    # 副本=人工值（默认跳过保留）
    assert written == 1
    assert _cell(copy_path, 2, "净重") == 99.0
    cur = state["current_factory_data"]
    assert cur["manual_override_kept"] == 1
    assert cur["manual_conflicts"] == [
        {"factory": FA, "sku": SKU1, "row": 2,
         "field": "net", "old": 15.0, "new": 99.0}]

    # 审计留痕：traces → _write_audit 落库（source=manual_override_kept）
    kept, overwritten, _, _ = service._manual_write_traces(
        {"factory_outputs": {FA: dict(cur)}}, None)
    assert len(kept) == 1 and kept[0]["source"] == "manual_override_kept"
    assert overwritten == []
    service._write_audit({
        "thread_id": batch, "factory_name": FA, "approved": True,
        "edited_count": 1, "changes": kept, "new_skus": [],
    }, "approved")
    audits = _audits(batch)
    assert len(audits) == 1
    changes = json.loads(audits[0].changes_json)
    assert changes[0]["source"] == "manual_override_kept"
    assert changes[0]["field"] == "net" and changes[0]["new"] == 99.0

    # 完成摘要透出：export_summary.json manual_overrides_kept=1
    summary = _export_summary(batch, cur, copy_path)
    assert summary["manual_overrides_kept"] == 1
    assert summary["manual_values_overwritten"] == 0

    # 快照 protected 打标（记录值结转更新为人工值）
    row = _snapshot_row(batch, FA, SKU1, 1)
    assert row["protected"] is True
    assert row["protected_fields"] == ["net"]
    assert row["net"] == 99.0


# ---------------------------------------------------------------------------
# 场景 3：改非空 · 通道 A 覆盖（reopen 全链路：重开 → 提交 → 覆写留痕）
# ---------------------------------------------------------------------------

def test_e2e_03_channel_a_reopen_overwrites(monkeypatch):
    batch = "T-E2E-03-CHA"
    src, copy_path = _setup_batch(batch, [_row()])

    _set_cell(copy_path, 2, "净重", 99.0)   # 人工改非空 W

    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "downstream_file_path": str(src),
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
        # 旧格式列表快照（reopen 数据源优先级：快照 > Excel 兜底）
        "factory_outputs": {FA: [_item(SKU1)]},
        "current_factory_data": {},
        "upstream_root": str(TMP / "upstream"),
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(service, "_latest_approved_audit",
                        lambda tid, fn: None)

    # ① 重开（真实读取链）：payload 正常构建、无原件变更警告
    payload = service.reopen_factory_for_edit(batch, FA)
    assert payload is not None and payload["reopen_mode"] is True
    assert payload["source_changes_warning"] is None

    # ② 提交（通道 A）：人工非空值直接覆盖，旧值入审计
    service.apply_reopen_payload(
        batch, FA, {"approved": True, "items": [_item(SKU1)]})

    assert _cell(copy_path, 2, "净重") == 15.0     # 副本=系统值
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is False

    audits = [a for a in _audits(batch) if a.result_status == "reopen"]
    assert len(audits) == 1
    changes = json.loads(audits[0].changes_json)
    overwritten = [c for c in changes
                   if c.get("source") == "manual_value_overwritten"]
    assert len(overwritten) == 1
    assert {k: overwritten[0][k] for k in ("sku", "field", "old", "new")} == {
        "sku": SKU1, "field": "net", "old": 99.0, "new": 15.0}
    # 批次已结束 → factory_outputs 快照回写（锚 NODE7）
    assert fake.updates and "factory_outputs" in fake.updates[0]


# ---------------------------------------------------------------------------
# 场景 4：改非空 · 通道 B 确认门（补充工厂：409 未入图 → 确认覆写/不确认保留）
# ---------------------------------------------------------------------------

def test_e2e_04_channel_b_confirm_gate(monkeypatch):
    batch = "T-E2E-04-CHB"
    src_rows = [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, name_cn=None, net=None, gross=None,
             po="272900", kanri="XC268453"),
    ]
    src, copy_path = _setup_batch(batch, src_rows)
    _set_cell(copy_path, 3, "净重", 66.0)   # 人工抢先填待补充厂 FB 行净重

    fake = _add_factories_graph(batch)      # FA 已完成，FB 待补充
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)
    # 原件解析走真实 parse_downstream_file（不 mock）

    # ① 未确认：ManualConflictError，未入图（update_state/stream 均未发生）
    with pytest.raises(service.ManualConflictError) as ei:
        service.add_factories_to_batch(batch)
    assert ei.value.conflicts == [
        {"factory": FB, "sku": SKU2, "row": 3,
         "field": "net", "old": None, "new": 66.0}]
    assert fake.updates == [] and fake.stream_calls == 0
    assert _cell(copy_path, 3, "净重") == 66.0      # 不确认保留人工值

    # ② 路由层契约：409 + 结构体 detail（前端确认门弹窗数据源）
    resp = client.post(f"/api/v1/batches/{batch}/add-factories", json={})
    assert resp.status_code == 409
    detail = resp.json()["detail"]
    assert detail["code"] == "manual_conflicts"
    assert detail["count"] == 1
    assert fake.stream_calls == 0                   # 路由层同样未入图
    assert _cell(copy_path, 3, "净重") == 66.0      # 仍保留

    # ③ 确认后：放行续跑，确认工厂写入 overwrite_manual_factories
    result = service.add_factories_to_batch(batch, confirm_overwrite=True)
    assert result["status"] == "completed"
    assert result["added"] == 1 and result["factories"] == [FB]
    assert result["manual_overrides_kept"] == 0     # 字段恒在（前端契约）
    assert fake.stream_calls == 1
    assert fake.updates[0]["overwrite_manual_factories"] == [FB]

    # ④ Node6 真实写入（假 graph 不跑节点，此处直接调 writer() 节点模拟）：
    # state 带 overwrite_manual_factories → 通道 A 覆写人工值
    node_state = {
        "batch_id": batch,
        "downstream_file_path": str(src),
        "validation_status": "Approved",
        "current_factory_data": {
            "factory_name": FB,
            "calculated_items": [_item(SKU2, unit_net=2.0, unit_gross=3.0,
                                       name_cn="扳手")],
        },
        "downstream_requirements": fake.updates[0]["downstream_requirements"],
        "downstream_row_map": {FB: {SKU2: [3]}},
        "overwrite_manual_factories":
            fake._snap.values["overwrite_manual_factories"],
        "factory_outputs": {},
        "upstream_root": str(TMP / "upstream"),
    }
    node_result = writer.writer(node_state)
    assert _cell(copy_path, 3, "净重") == 10.0      # 2.0 × 副本数量 5
    assert _cell(copy_path, 3, "中文品名") == "扳手"
    assert node_result["overwrite_manual_factories"] == []   # 消费后自清
    # 已知行为边界（阶段六偏差说明记录）：确认后 add_factories_to_batch 的
    # sync 在跑图前 refresh_snapshot 重建快照，把人工值 66 并入基线 → Node6
    # 写前 diff 干净，通道 A 覆写成立但 manual_value_overwritten 留痕为空
    # （service.py「旧值留痕」注释与实际链路存在落差，留待热修复决策）。
    # 此处钉死现状：覆写生效 + 确认标记消费自清。
    assert node_state["current_factory_data"]["manual_values_overwritten"] == 0


# ---------------------------------------------------------------------------
# 场景 5：删行不复活（MX1 回归 e2e：补充工厂全链路 + 墓碑含键）
# ---------------------------------------------------------------------------

def test_e2e_05_deleted_rows_never_resurrect(monkeypatch):
    batch = "T-E2E-05-MX1"
    src_rows = [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, name_cn=None, net=None, gross=None,
             po="272900", kanri="XC266877"),
        _row(factory=FB, sku=SKU2, qty=5, name_cn=None, net=None, gross=None,
             po="272901", kanri="XC266877"),
    ]
    src, copy_path = _setup_batch(batch, src_rows)

    _delete_row(copy_path, 3)       # 人工删 FB 第一行（po 272900；部分删行
                                    # 不触发整厂剔除）

    fake = _add_factories_graph(batch)      # FA 已完成，FB 待补充
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)

    # 补充工厂全链路（真实解析原件 + 真实 sync + 真实快照）
    result = service.add_factories_to_batch(batch)
    assert result["status"] == "completed"
    assert result["added"] == 1 and result["factories"] == [FB]

    # 被删行不复活：副本仍 3 行（表头 + FA + FB 剩余行），po 272900 不回来
    pos = [_cell(copy_path, r, "PURCHASE_ORDER") for r in (2, 3)]
    assert _ws(copy_path).max_row == 3
    assert "272900" not in pos and "272901" in pos

    # 墓碑含被删行键（seq2：副本重扫后剩余行占 seq1，消失的是快照 seq2）
    snap = copy_snapshot.load_snapshot(batch)
    assert copy_snapshot.row_key_str((FB, SKU2, 2)) in snap["tombstones"]
    # 未整厂剔除（FB 还有一行存活）：无 removed 审计
    assert [a for a in _audits(batch)
            if a.result_status == "factory_manually_removed"] == []


# ---------------------------------------------------------------------------
# 场景 6：整厂剔除（审计 removed + 完成判定 + watch 不触发 + 补充不加回）
# ---------------------------------------------------------------------------

def _make_awaiting_batch(tid: str, upstream: Path, *,
                         done: list[str], gap: list[str]) -> None:
    """照 test_factory_removal 先例：awaiting 批次行 + 真实 graph checkpoint
    （as_node=NODE7，next 为空）。"""
    batch_store.upsert_batch(tid, status="awaiting_factories")
    _CREATED_TIDS.append(tid)
    graph = get_graph()
    cfg = {"configurable": {"thread_id": tid}}
    graph.update_state(cfg, {
        "batch_id": tid,
        "downstream_requirements": {f: ["4900000000001"] for f in done + gap},
        "factory_outputs": {f: {"calculated_items": []} for f in done},
        "upstream_root": str(upstream),
        "final_output_path": "/tmp/out.xlsx",
    }, as_node=NODE7)


def test_e2e_06_factory_removal_full_chain(monkeypatch, tmp_path, caplog):
    batch = "T-E2E-06-RM"
    src_rows = [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, name_cn=None, net=None, gross=None,
             po="272900", kanri="XC266877"),
    ]
    src, copy_path = _setup_batch(batch, src_rows)

    # 剔除前：FB 是缺口（完成文案口径对照）
    final_values = {"batch_id": batch,
                    "downstream_requirements": {FA: [SKU1], FB: [SKU2]},
                    "factory_outputs": {FA: {"calculated_items": []}}}
    assert service._gap_factories(final_values) == [FB]
    assert pipeline_state._classify_incomplete_batch(batch, [FB], None) \
        == "awaiting_factories"

    # 人工删光 FB 全部行 → 走真实写前 diff 检出（sync 是写入路径之一；
    # 真实解析原件提供 row_map/requirements 全集）
    _delete_row(copy_path, 3)
    from app.nodes.parse_downstream import parse_downstream_file
    _reqs, rm = parse_downstream_file(str(src))
    sync_result = service._sync_output_copy_rows(batch, str(src), {FB}, rm)
    assert sync_result == {"synced": 0, "skipped_manual_deleted": 1}

    # ① 审计 removed 落库（approved=False + 行数留痕）
    audits = [a for a in _audits(batch)
              if a.result_status == "factory_manually_removed"]
    assert len(audits) == 1
    assert audits[0].factory_name == FB and audits[0].approved is False
    assert json.loads(audits[0].changes_json) == [
        {"factory_removed": True, "rows": 1}]
    assert service._manually_removed_factories(batch) == {FB}
    # 墓碑双轨印证
    snap = copy_snapshot.load_snapshot(batch)
    assert copy_snapshot.row_key_str((FB, SKU2, 1)) in snap["tombstones"]

    # ② 完成判定：唯一缺口=剔除厂 → completed；缺口文案减除剔除厂
    assert pipeline_state._classify_incomplete_batch(batch, [FB], None) \
        == "completed"
    assert service._gap_factories(final_values) == []

    # ③ 补充不加回：add_factories_to_batch 早退 + info 日志，未入图
    fake = _add_factories_graph(
        batch, requirements={FA: [SKU1], FB: [SKU2]},
        row_map={FA: {SKU1: [2]}, FB: {SKU2: [3]}})
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)
    with caplog.at_level(logging.INFO, logger="app.api.service"):
        result = service.add_factories_to_batch(batch)
    assert result["added"] == 0 and result["factories"] == []
    assert result["message"] == "没有待补充的工厂"
    assert fake.updates == [] and fake.stream_calls == 0
    assert any("已被人工剔除，不再补回" in r.message and FB in r.message
               for r in caplog.records)

    # ④ factory_watch：剔除厂资料到达也不触发（对照批正常触发）。
    # watch 扫描经 service.get_graph 读 checkpoint——先把上面钉的假 graph
    # 还原成真 graph（watch 批次走真实 checkpoint，上游根在 state 里）。
    monkeypatch.setattr(service, "get_graph", get_graph)
    calls: list[str] = []

    def fake_add(tid: str):
        calls.append(tid)
        return {"added": 1, "factories": [], "status": "pending_human_review"}

    monkeypatch.setattr(service, "add_factories_to_batch", fake_add)
    upstream = tmp_path / "upstream"
    upstream.mkdir()
    (upstream / "缺口厂C").mkdir()
    (upstream / "缺口厂C" / "packing.pdf").write_bytes(b"x" * 128)
    _make_awaiting_batch("T-E2E-06W-CTL", upstream,
                         done=["完成厂C"], gap=["缺口厂C"])
    (upstream / "剔除厂R").mkdir()
    (upstream / "剔除厂R" / "packing.pdf").write_bytes(b"x" * 128)
    _make_awaiting_batch("T-E2E-06W-RM", upstream,
                         done=["完成厂R"], gap=["剔除厂R"])
    # 剔除审计（直接插一条，供 watch 名单判定）
    with get_session() as session:
        session.add(ReviewAudit(
            thread_id="T-E2E-06W-RM", factory_name="剔除厂R", approved=False,
            edited_count=0,
            changes_json=json.dumps([{"factory_removed": True, "rows": 1}],
                                    ensure_ascii=False),
            new_skus_json="[]",
            result_status="factory_manually_removed"))
        session.commit()

    factory_watch._scan_once()   # 首轮仅记录快照
    factory_watch._scan_once()   # 第二轮稳定 → 触发
    factory_watch._scan_once()   # 再扫一轮防迟发
    assert calls.count("T-E2E-06W-CTL") == 1     # 对照批正常触发
    assert "T-E2E-06W-RM" not in calls           # 剔除厂不再监控


# ---------------------------------------------------------------------------
# 场景 7：剔除后加回（出碑 + readded 审计 + 重新纳入 pending）
# ---------------------------------------------------------------------------

def test_e2e_07_readd_after_removal(monkeypatch):
    batch = "T-E2E-07-READD"
    src_rows = [
        _row(),
        _row(factory=FB, sku=SKU2, qty=5, name_cn=None, net=None, gross=None,
             po="272900", kanri="XC266877"),
    ]
    src, copy_path = _setup_batch(batch, src_rows)

    # ① 剔除：删光 FB 行 → sync 写前 diff 检出 → removed 审计 + 入碑
    _delete_row(copy_path, 3)
    from app.nodes.parse_downstream import parse_downstream_file
    _reqs, rm = parse_downstream_file(str(src))
    service._sync_output_copy_rows(batch, str(src), {FB}, rm)
    assert service._manually_removed_factories(batch) == {FB}
    fb_key = copy_snapshot.row_key_str((FB, SKU2, 1))
    assert fb_key in copy_snapshot.load_snapshot(batch)["tombstones"]

    # ② 人工把 FB 行加回副本尾部
    _append_row(copy_path, _row(factory=FB, sku=SKU2, qty=5, name_cn=None,
                                net=None, gross=None,
                                po="272900", kanri="XC266877"))

    # ③ 任一写入路径的写前 diff 检出对称恢复（此处走真实 Node6 写入链）：
    # readded 审计落库 → 自动出名单；写后 refresh 结转 → 出碑 + 入快照
    state = _write_state(batch, requirements={FA: [SKU1], FB: [SKU2]})
    writer._write_excel(state, copy_path, channel="first_write")

    audits = _audits(batch)
    assert [a.result_status for a in audits] == [
        "factory_manually_removed", "factory_manually_readded"]
    assert json.loads(audits[1].changes_json) == [
        {"factory_readded": True, "rows": 1}]
    assert service._manually_removed_factories(batch) == set()
    snap = copy_snapshot.load_snapshot(batch)
    assert fb_key not in snap["tombstones"]
    fb_row = _snapshot_row(batch, FB, SKU2, 1)
    assert fb_row["origin"] == "manual_added"

    # ④ 重新纳入：补充工厂把 FB 补回 pending（出名单后恢复待补充口径）
    fake = _add_factories_graph(
        batch, requirements={FA: [SKU1], FB: [SKU2]},
        row_map={FA: {SKU1: [2]}, FB: {SKU2: [3]}})
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)
    result = service.add_factories_to_batch(batch)
    assert result["status"] == "completed"
    assert result["added"] == 1 and result["factories"] == [FB]
    assert fake.updates[0]["pending_factories"] == [FB]
    assert fake.stream_calls == 1


# ---------------------------------------------------------------------------
# 场景 8：移位重定位（插 3 行后 reopen 提交写对行、清对行）
# ---------------------------------------------------------------------------

def test_e2e_08_relocation_after_manual_insert(monkeypatch):
    batch = "T-E2E-08-RELOC"
    src_rows = [
        _row(),                                                  # row2 SKU1
        _row(sku=SKU2, qty=5, name_cn="扳手", net=50.0, gross=55.0,
             po="272753", kanri="XC266877"),                     # row3 SKU2
    ]
    src, copy_path = _setup_batch(batch, src_rows)

    # 人工在目标行前插入 3 行：SKU1 row2→row5、SKU2 row3→row6
    wb = load_workbook(copy_path)
    wb.active.insert_rows(2, 3)
    wb.save(copy_path)

    fake = _FakeGraph(_FakeSnap(values={
        "batch_id": batch,
        "downstream_file_path": str(src),
        "downstream_requirements": {FA: [SKU1, SKU2]},
        # row_map 残留插行前的旧行号——重定位后不再被消费
        "downstream_row_map": {FA: {SKU1: [2], SKU2: [3]}},
        "factory_outputs": {FA: [_item(SKU1),
                                 _item(SKU2, name_cn="扳手")]},
        "current_factory_data": {},
        "upstream_root": str(TMP / "upstream"),
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    monkeypatch.setattr(service, "_latest_approved_audit",
                        lambda tid, fn: None)

    # reopen 提交：SKU1 保留（重写三列），SKU2 人工删除（清空三列）
    service.apply_reopen_payload(batch, FA, {
        "approved": True,
        "items": [_item(SKU1),
                  {"sku": SKU2, "orig_sku": SKU2, "deleted": True}],
    })

    # 写对行：移位后 row5 = SKU1 系统值；插入的空白行未被误写
    assert _cell(copy_path, 5, "SHOHIN_CD") == SKU1
    assert _cell(copy_path, 5, "中文品名") == "钳子"
    assert _cell(copy_path, 5, "净重") == 15.0
    assert _cell(copy_path, 5, "毛重") == 25.0
    # 清对行：移位后 row6 = SKU2 三列全空，其余列不动
    assert _cell(copy_path, 6, "SHOHIN_CD") == SKU2
    for f in ("中文品名", "净重", "毛重"):
        assert _cell(copy_path, 6, f) is None
    for r in (2, 3, 4):
        assert _cell(copy_path, r, "净重") is None
    # reopen 审计落库（删除条目留痕）
    assert any(a.result_status == "reopen" for a in _audits(batch))


# ---------------------------------------------------------------------------
# 场景 9：新增行三态（A 并入写入 / B 全新厂入批 / C 缺番号不纳入）
# ---------------------------------------------------------------------------

def test_e2e_09_manual_added_rows_three_states(monkeypatch):
    # ---- A 态：原件已有厂 + 副本新增 SKU 行，提取结果含此 SKU → 写入 ----
    batch_a = "T-E2E-09A"
    src_a, copy_a = _setup_batch(batch_a, [_row()])
    _append_row(copy_a, _row(sku=SKU2, qty=6, name_cn=None, net=None,
                             gross=None, po="272753", kanri="XC268499"))

    node_state = {
        "batch_id": batch_a,
        "downstream_file_path": str(src_a),
        "validation_status": "Approved",
        "current_factory_data": {
            "factory_name": FA,
            "calculated_items": [_item(SKU1), _item(SKU2, name_cn="剪刀")],
        },
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
        "factory_outputs": {},
        "upstream_root": str(TMP / "upstream"),
    }
    result_a = writer.writer(node_state)    # 真实 Node6 节点

    assert _cell(copy_a, 3, "中文品名") == "剪刀"
    assert _cell(copy_a, 3, "净重") == 9.0          # 1.5 × 副本数量 6
    # 并入计划经节点返回值合并：A 态需求追加 + row_map 占位
    assert result_a["downstream_requirements"][FA] == [SKU1, SKU2]
    assert result_a["downstream_row_map"][FA][SKU2] == [3]
    # 写后快照：新增行入库且 origin="manual_added"
    assert _snapshot_row(batch_a, FA, SKU2, 1)["origin"] == "manual_added"

    # ---- B 态：副本新增全新厂行 → 补充工厂入批 + 快照打标 ----
    batch_b = "T-E2E-09B"
    _src_b, copy_b = _setup_batch(batch_b, [_row()])   # 原件只有 FA
    _append_row(copy_b, _row(factory=FC, sku=SKU2, qty=5, name_cn=None,
                             net=None, gross=None,
                             po="272900", kanri="XC266877"))

    fake_b = _add_factories_graph(batch_b)
    monkeypatch.setattr(service, "get_graph", lambda: fake_b)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)
    result_b = service.add_factories_to_batch(batch_b)  # 真实解析（无 FC）

    assert result_b["added"] == 1 and result_b["factories"] == [FC]
    assert fake_b.stream_calls == 1
    update_b = fake_b.updates[0]
    assert update_b["pending_factories"] == [FC]
    # 新厂需求明细来自副本新增行（SKU 去重）
    assert update_b["downstream_requirements"][FC] == [SKU2]
    assert _snapshot_row(batch_b, FC, SKU2, 1)["origin"] == "manual_added"

    # ---- C 态：缺 KANRI_NO 行 → 不纳入 + 活扫透出 + 摘要提示 ----
    batch_c = "T-E2E-09C"
    src_c, copy_c = _setup_batch(batch_c, [_row()])
    _append_row(copy_c, _row(sku=SKU2, qty=6, kanri=None))

    rows = copy_snapshot.scan_incomplete_rows(batch_c, str(src_c))
    assert rows == [{"row": 3, "factory": FA, "sku": SKU2,
                     "missing": ["KANRI_NO"]}]

    fake_c = _add_factories_graph(batch_c)
    monkeypatch.setattr(service, "get_graph", lambda: fake_c)
    result_c = service.add_factories_to_batch(batch_c)
    assert result_c["added"] == 0 and result_c["factories"] == []
    assert fake_c.updates == [] and fake_c.stream_calls == 0
    assert "1 行人工新增的行缺管理番号" in result_c["message"]
    assert "请补全" in result_c["message"]


# ---------------------------------------------------------------------------
# 场景 10：无快照迁移（老批次首写建 baseline + warning；基线前改动不保护）
# ---------------------------------------------------------------------------

def test_e2e_10_legacy_batch_baseline_migration(caplog):
    batch = "T-E2E-10-BASELINE"
    _, copy_path = _setup_batch(batch, [_row()], snapshot=False)  # 无快照
    assert not copy_snapshot.snapshot_path(batch).exists()

    # 基线前的人工改动（历史批次副本上的既有手改值）
    _set_cell(copy_path, 2, "净重", 55.0)

    # 首次写入操作：diff 无快照保守放行 → 系统值直接覆盖（决策 7：
    # 基线前改动不享受保护）；落盘后 refresh 以副本现状建 baseline
    state = _write_state(batch)
    with caplog.at_level(logging.WARNING,
                         logger="app.orchestrator.copy_snapshot"):
        written, _ = writer._write_excel(state, copy_path,
                                         channel="first_write")
    assert written == 1
    assert _cell(copy_path, 2, "净重") == 15.0      # 不保护，覆盖
    assert state["current_factory_data"]["manual_override_kept"] == 0

    # baseline 快照已建 + warning 透出「不可追溯」
    assert copy_snapshot.snapshot_path(batch).exists()
    snap = copy_snapshot.load_snapshot(batch)
    assert all(r["origin"] == "baseline" for r in snap["rows"])
    assert snap["tombstones"] == []
    assert any("基线" in r.message and "不可追溯" in r.message
               for r in caplog.records if r.levelno >= logging.WARNING)


# ---------------------------------------------------------------------------
# 场景 11：快照结转（二次写入不覆盖 protected + 墓碑不复活）
# ---------------------------------------------------------------------------

def test_e2e_11_snapshot_carry_over():
    batch = "T-E2E-11-CARRY"
    src_rows = [
        _row(),
        _row(sku=SKU2, qty=5, name_cn="扳手", net=50.0, gross=55.0,
             po="272753", kanri="XC266877"),
    ]
    src, copy_path = _setup_batch(batch, src_rows)

    # 通道 C 第一次：人工改非空 → 跳过保留 + protected 打标
    _set_cell(copy_path, 2, "净重", 99.0)
    writer._write_excel(_write_state(batch), copy_path, channel="first_write")
    assert _snapshot_row(batch, FA, SKU1, 1)["protected_fields"] == ["net"]

    # 二次写入（快照已重建）：protected 结转生效，仍不覆盖人工值
    state2 = _write_state(batch)
    writer._write_excel(state2, copy_path, channel="first_write")
    assert _cell(copy_path, 2, "净重") == 99.0
    assert state2["current_factory_data"]["manual_override_kept"] == 1
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is True

    # 人工删 SKU2 行 → sync（真实服务链路）→ 墓碑结转不复活
    _delete_row(copy_path, 3)
    from app.nodes.parse_downstream import parse_downstream_file
    _reqs, rm = parse_downstream_file(str(src))
    result = service._sync_output_copy_rows(batch, str(src), {FA}, rm)
    assert result == {"synced": 0, "skipped_manual_deleted": 1}
    assert _ws(copy_path).max_row == 2              # SKU2 行未复活
    snap = copy_snapshot.load_snapshot(batch)
    assert copy_snapshot.row_key_str((FA, SKU2, 1)) in snap["tombstones"]
    # protected 行不受删行/重建影响，保护仍在
    assert _snapshot_row(batch, FA, SKU1, 1)["protected"] is True
    # 再 sync 一次：墓碑口径同样不复活（双口径印证）
    result2 = service._sync_output_copy_rows(batch, str(src), {FA}, rm)
    assert result2 == {"synced": 0, "skipped_manual_deleted": 1}
    assert _ws(copy_path).max_row == 2


# ---------------------------------------------------------------------------
# 场景 12：原件差异三入口（补充工厂确认吸收 / 单厂重试确认 / reopen 警告条）
# ---------------------------------------------------------------------------

def test_e2e_12_source_diff_three_entries(monkeypatch):
    # ---- 入口一·补充工厂：已完成工厂原件加行 → 确认门 → 吸收重写 ----
    batch_a = "T-E2E-12A"
    src_a, copy_a = _setup_batch(batch_a, [_row()])
    fake_a = _FakeGraph(_FakeSnap(values={
        "batch_id": batch_a,
        "factory_outputs": {FA: {
            "factory_name": FA,
            "calculated_items": [_item(SKU1),
                                 _item(SKU2, unit_net=2.0, unit_gross=3.0,
                                       name_cn="扳手")],
        }},
        "pending_factories": [],
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
        "upstream_root": str(TMP / "upstream"),
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake_a)
    monkeypatch.setattr(
        service, "_auto_create_factory_folders", lambda *a, **k: None)

    # 批次完成后人工改原件：给已完成工厂 FA 加一行（SKU2，qty=3）
    _append_row(src_a, _row(sku=SKU2, qty=3, name_cn=None, net=None,
                            gross=None, po="272900", kanri="XC268453"))

    # 未确认：SourceChangesError 未入图，副本不动
    with pytest.raises(service.SourceChangesError) as ei:
        service.add_factories_to_batch(batch_a)
    assert ei.value.changes == [{"factory": FA, "added": 1, "removed": 0,
                                 "qty_changed": 0}]
    assert "是否吸收重新写入" in str(ei.value)
    assert fake_a.updates == [] and fake_a.stream_calls == 0
    assert _ws(copy_a).max_row == 2

    # 确认后：state 吸收 + 副本重写（新行追加尾部 + 通道 A 写三列）
    result_a = service.add_factories_to_batch(
        batch_a, confirm_source_changes=True)
    assert result_a["status"] == "completed"
    assert result_a["source_changes_absorbed"] == [FA]
    assert "已吸收原件变更" in result_a["message"]
    merged_rm = fake_a.updates[0]["downstream_row_map"]
    assert merged_rm[FA][SKU2] == [3]
    assert _cell(copy_a, 3, "中文品名") == "扳手"
    assert _cell(copy_a, 3, "净重") == 6.0          # 2.0 × qty 3
    assert _cell(copy_a, 3, "毛重") == 9.0
    assert _cell(copy_a, 2, "净重") == 15.0         # 既有行不受影响
    assert any(a.result_status == "source_absorbed" for a in _audits(batch_a))
    # 指纹基线已重记：再次调用不再报 source_changes
    assert source_diff.source_fingerprint_changed(batch_a) is False

    # ---- 入口二·单厂重试：指纹变 + 该厂数量改 → 确认门 → 吸收续跑 ----
    batch_b = "T-E2E-12B"
    src_b, _ = _setup_batch(batch_b, [_row()])
    fake_b = _FakeGraph(_FakeSnap(
        values={
            "batch_id": batch_b,
            "current_factory_data": {"factory_name": FA},
            "downstream_requirements": {FA: [SKU1]},
            "downstream_row_map": {FA: {SKU1: [2]}},
            "downstream_file_path": str(src_b),
            "upstream_root": str(TMP / "upstream"),
        },
        next_nodes=(service.NODE5,),                # 挂起待审核
    ))
    monkeypatch.setattr(service, "get_graph", lambda: fake_b)

    _set_cell(src_b, 2, "SOTOBAKO_D_HACCHU_SU", 20)  # 人工改原件数量

    with pytest.raises(service.SourceChangesError) as ei2:
        service.retry_factory_extraction(batch_b)
    assert ei2.value.changes == [{"factory": FA, "added": 0, "removed": 0,
                                  "qty_changed": 1}]
    assert fake_b.stream_calls == 0                 # 未确认不重跑

    result_b = service.retry_factory_extraction(
        batch_b, confirm_source_changes=True)
    assert result_b["status"] == "completed"
    assert result_b["source_changes_absorbed"] == [
        {"factory": FA, "added": 0, "removed": 0, "qty_changed": 1}]
    assert fake_b.stream_calls == 1
    assert fake_b.updates and fake_b.updates[0]["force_reextract"] is True
    # state 吸收：该厂 row_map 按原件重解析并入
    assert fake_b.updates[0]["downstream_row_map"][FA] == {SKU1: [2]}
    # 指纹基线不在此更新（其余工厂差异仍应由入口一检出）
    assert source_diff.source_fingerprint_changed(batch_b) is True

    # ---- 入口三·reopen：指纹变 → 审核页警告条数据透出且不阻断 ----
    batch_c = "T-E2E-12C"
    src_c, _ = _setup_batch(batch_c, [_row()])
    fake_c = _FakeGraph(_FakeSnap(values={
        "batch_id": batch_c,
        "factory_outputs": {FA: [_item(SKU1)]},
        "downstream_requirements": {FA: [SKU1]},
        "downstream_row_map": {FA: {SKU1: [2]}},
        "upstream_root": str(TMP / "upstream"),
    }))
    monkeypatch.setattr(service, "get_graph", lambda: fake_c)
    monkeypatch.setattr(service, "_latest_approved_audit",
                        lambda tid, fn: None)

    # 指纹未变：无警告条
    payload = service.reopen_factory_for_edit(batch_c, FA)
    assert payload is not None and payload["source_changes_warning"] is None

    # 人工改原件数量 → 警告条透出，不阻断编辑
    _set_cell(src_c, 2, "SOTOBAKO_D_HACCHU_SU", 20)
    payload2 = service.reopen_factory_for_edit(batch_c, FA)
    assert payload2 is not None and payload2["reopen_mode"] is True
    warning = payload2["source_changes_warning"]
    assert warning is not None
    assert "原件已在批次运行后被修改" in warning["message"]
    assert warning["changes"] == [{"factory": FA, "added": 0, "removed": 0,
                                   "qty_changed": 1}]


# ---------------------------------------------------------------------------
# 场景 13：完全重跑警告（未确认拒绝未归档；确认后旧档归档可查 + 副本重建）
# ---------------------------------------------------------------------------

def test_e2e_13_rerun_discard_manual_gate(monkeypatch):
    batch = "T-E2E-13-RERUN"
    src, copy_path = _setup_batch(batch, [_row()])
    _set_cell(copy_path, 2, "净重", 99.0)   # 人工改副本 → 痕迹

    settings = get_settings()
    out_dir = settings.batch_output_dir(batch)

    # 痕迹检测：diff value_conflicts 非空
    traces = service.detect_manual_traces(batch)
    assert traces is not None and traces["value_conflicts"] == 1

    # ① 未确认：ManualDiscardError 拒绝，无任何副作用（未归档）
    with pytest.raises(service.ManualDiscardError) as ei:
        service.rerun_batch(batch)
    assert "副本上的全部人工改动将丢失" in str(ei.value)
    assert ei.value.traces["value_conflicts"] == 1
    assert out_dir.exists() and (out_dir / "batch_config.json").exists()
    assert not settings.history_output_dir(batch).exists()

    # 路由层契约：409 结构体（code=discard_manual_changes，带 traces）
    resp = client.post(f"/api/v1/batches/{batch}/rerun")
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "discard_manual_changes"
    assert out_dir.exists()                         # 路由层同样未归档

    # ② 确认后：放行——output 归档 _history/、checkpoint 重置、指纹重记
    fake = _FakeGraph(_FakeSnap(values={"batch_id": batch}))
    monkeypatch.setattr(service, "get_graph", lambda: fake)
    result = service.rerun_batch(batch, confirm_discard_manual=True)
    assert result["status"] == "completed"
    assert fake.updates, "checkpoint 应被重置（as_node=START 合并）"
    assert settings.history_output_dir(batch).exists()
    assert out_dir.exists()                         # 新 run 目录重建
    cfg = json.loads((out_dir / "batch_config.json").read_text(
        encoding="utf-8"))
    assert cfg["run_count"] == 2

    # ③ 旧档归档可查回：归档副本里仍是人工值 99
    archived = list(settings.history_output_dir(batch).glob(
        "r1_*/run/containers/*_filled.xlsx"))
    assert len(archived) == 1
    assert _cell(archived[0], 2, "净重") == 99.0
    # 旧快照随 run 归档（不残留 in 在产目录）
    assert not copy_snapshot.snapshot_path(batch).exists()

    # ④ 副本重建：下一写入链路 _ensure_output_copy 从原件重新复制，
    # 人工改动不带入新副本；基线快照重建
    new_copy = writer._ensure_output_copy(
        {"batch_id": batch, "downstream_file_path": str(src)})
    assert new_copy.exists()
    assert _cell(new_copy, 2, "净重") == 15.0       # 原件原值
    assert copy_snapshot.snapshot_path(batch).exists()


# ---------------------------------------------------------------------------
# 场景 14：数量列告警（人工改数量 → 按新数量写入 + warning + 摘要透出）
# ---------------------------------------------------------------------------

def test_e2e_14_qty_change_warns_and_writes_by_new_qty(caplog):
    batch = "T-E2E-14-QTY"
    _, copy_path = _setup_batch(batch, [_row()])

    _set_cell(copy_path, 2, "SOTOBAKO_D_HACCHU_SU", 20)   # 人工改数量

    state = _write_state(batch)
    with caplog.at_level(logging.WARNING, logger="app.nodes.writer"):
        written, _ = writer._write_excel(state, copy_path,
                                         channel="first_write")

    # 写入仍按副本当前数量吸收（1.5×20 / 2.5×20）
    assert written == 1
    assert _cell(copy_path, 2, "净重") == 30.0
    assert _cell(copy_path, 2, "毛重") == 50.0
    cur = state["current_factory_data"]
    assert cur["qty_manual_changes"] == [
        {"factory": FA, "sku": SKU1, "row": 2, "old": 10.0, "new": 20.0}]
    assert any("数量被人工改过" in r.message and "已按新数量写入" in r.message
               for r in caplog.records)

    # 摘要透出：traces 带 source + export_summary 计数字段
    _k, _o, qty_changes, _r = service._manual_write_traces(
        {"factory_outputs": {FA: dict(cur)}}, None)
    assert len(qty_changes) == 1
    assert qty_changes[0]["source"] == "manual_qty_changed"
    summary = _export_summary(batch, cur, copy_path)
    assert summary["qty_manual_changes"] == 1


if __name__ == "__main__":
    import pytest as _pytest
    _pytest.main([__file__, "-v"])
